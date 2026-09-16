# -*- coding: utf-8 -*-
"""Idempotent migration of legacy Contract Excel exports into TaskHub v7.

Default mode is a read-only dry run. Nothing is inserted unless ``--apply``
is passed. Orphans are retained in MigrationQuarantine. Historical statement
and extension money mappings use the rules validated against the legacy DB.
"""
from __future__ import annotations

import argparse
import datetime
import decimal
import json
import os
import sys

from openpyxl import load_workbook

# Application modules live in ../src.
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'src'))
from config import connection_string


PERSIAN_DIGITS = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
FINAL_SNAPSHOT_EXTENSION_IDS = {3029, 3030}


def clean(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    return value


def json_value(value):
    value = clean(value)
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    if isinstance(value, decimal.Decimal):
        return str(value)
    return value


def payload(row):
    return json.dumps({k: json_value(v) for k, v in row.items()}, ensure_ascii=False, default=str)


def as_int(value):
    try:
        return int(str(value).translate(PERSIAN_DIGITS)) if value not in (None, '') else None
    except Exception:
        return None


def as_decimal(value):
    try:
        return decimal.Decimal(str(value).translate(PERSIAN_DIGITS).replace(',', '')) if value not in (None, '') else None
    except Exception:
        return None


def as_bool(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ('1', 'true', 'yes', 'بله')


def fa_period(value, fallback_year=None):
    if value:
        try:
            parts = str(value).translate(PERSIAN_DIGITS).replace('-', '/').split('/')
            if len(parts) >= 2:
                year, month = int(parts[0]), int(parts[1])
                if 1300 <= year <= 1600 and 1 <= month <= 12:
                    return year, month
        except Exception:
            pass
    return as_int(fallback_year), None


def read_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    headers = [str(x).strip() if x is not None else '' for x in next(iterator)]
    for values in iterator:
        row = {headers[i]: clean(values[i]) for i in range(min(len(headers), len(values)))}
        if any(v not in (None, '') for v in row.values()):
            yield row


class Migrator:
    def __init__(self, conn, apply=False):
        self.conn = conn
        self.cur = conn.cursor() if conn is not None else None
        self.apply = apply
        self.stats = {
            'contracts_seen': 0, 'contracts_inserted': 0, 'contracts_existing': 0,
            'sentinels': 0, 'extensions_seen': 0, 'extensions_inserted': 0,
            'extensions_orphan': 0, 'extensions_ambiguous': 0,
            'statements_seen': 0, 'statements_inserted': 0,
            'statements_orphan': 0, 'quarantine': 0,
        }
        self.legacy_contracts = {}

    def quarantine(self, source, legacy_id, reason, data):
        self.stats['quarantine'] += 1
        if self.apply:
            self.cur.execute("""IF NOT EXISTS(SELECT 1 FROM MigrationQuarantine
                WHERE source_table=? AND legacy_id=? AND reason=? AND resolution_status='pending')
                INSERT INTO MigrationQuarantine(source_table,legacy_id,reason,payload_json)
                VALUES(?,?,?,?)""", source, str(legacy_id) if legacy_id is not None else None,
                reason, source, str(legacy_id) if legacy_id is not None else None, reason, payload(data))

    def load_existing_contracts(self):
        if self.cur is None:
            self.legacy_contracts = {}
            return
        self.cur.execute("SELECT id,legacy_contract_id FROM Contracts WHERE legacy_contract_id IS NOT NULL")
        self.legacy_contracts = {int(r[1]): int(r[0]) for r in self.cur.fetchall()}

    def contracts(self, path):
        for row in read_rows(path):
            self.stats['contracts_seen'] += 1
            legacy_id = as_int(row.get('ContractId'))
            if legacy_id is None or legacy_id <= 0:
                self.stats['sentinels'] += 1
                self.quarantine('t_Contract', legacy_id, 'sentinel_or_invalid_contract', row)
                continue
            if legacy_id in self.legacy_contracts:
                self.stats['contracts_existing'] += 1
                continue
            title = str(row.get('ContractTitle') or '').strip()
            if not title:
                self.quarantine('t_Contract', legacy_id, 'missing_required_title', row)
                continue
            if self.apply:
                self.cur.execute("""INSERT INTO Contracts(
                    project_id,legacy_contract_id,legacy_company_id,title,legacy_employer_id,
                    contract_number,base_price,notification_date,notification_date_fa,
                    notification_letter_number,start_date,start_date_fa,end_date,end_date_fa,
                    contract_type_id,contract_status_id,work_order_code,financial_code,
                    register_in_sajat,register_date_sajat,register_date_sajat_fa,
                    list_price_version_id,category_id,insurance_coefficient,guarantee_coefficient,
                    commercial_code,scope_id,is_new,tax_info_crn,source_json)
                    OUTPUT INSERTED.id VALUES(NULL,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    legacy_id, as_int(row.get('CompanyId')), title, as_int(row.get('EmployerId')),
                    row.get('ContractNumber'), as_decimal(row.get('ContractPrice')),
                    clean(row.get('ContractNotificationDate')), row.get('ContractNotificationDateFa'),
                    row.get('ContractNotificationLetterNumber'), clean(row.get('ContractStartDate')),
                    row.get('ContractStartDateFa'), clean(row.get('ContractEndDate')),
                    row.get('ContractEndDateFa'), as_int(row.get('ContractTypeId')),
                    as_int(row.get('ContractStatusId')), row.get('WorkOrderCode'),
                    row.get('ContractFinancialCode'), as_bool(row.get('RegisterInSajat')),
                    clean(row.get('RegisterDateSajat')), row.get('RegisterDateSajatFa'),
                    as_int(row.get('ListPriceVersionId')), as_int(row.get('ContractCategoryId')),
                    as_decimal(row.get('ContractInsuranceCoefficient')),
                    as_decimal(row.get('ContractGuaranteeCoefficient')), row.get('ContractCommercialCode'),
                    as_int(row.get('ContractScopeId')), as_bool(row.get('IsNew')),
                    row.get('ContractStatementTaxInfoCRN'), payload(row))
                new_id = int(self.cur.fetchone()[0])
                self.legacy_contracts[legacy_id] = new_id
            else:
                self.legacy_contracts[legacy_id] = -legacy_id  # dry-run marker
            self.stats['contracts_inserted'] += 1

            # The five flattened legacy extension slots are intentionally not
            # converted into active addenda: child-table rows are authoritative.
            # Preserve a review marker only when a slot actually contains data.
            slots = []
            for i in range(1, 6):
                if row.get('ContractExtensionDate%d' % i) is not None or row.get('ContractExtensionPrice%d' % i) is not None:
                    slots.append({'slot': i, 'date': json_value(row.get('ContractExtensionDate%d' % i)),
                                  'date_fa': row.get('ContractExtensionDateFa%d' % i),
                                  'price': json_value(row.get('ContractExtensionPrice%d' % i))})
            if slots:
                self.quarantine('t_Contract', legacy_id, 'flattened_extension_slots_for_review', {'slots': slots})

    def extensions(self, path):
        for row in read_rows(path):
            self.stats['extensions_seen'] += 1
            legacy_id = as_int(row.get('ContractExtensionId'))
            legacy_contract = as_int(row.get('ContractId'))
            contract_id = self.legacy_contracts.get(legacy_contract)
            if not contract_id:
                self.stats['extensions_orphan'] += 1
                self.quarantine('t_ContractExtension', legacy_id, 'orphan_contract_%s' % legacy_contract, row)
                continue
            if self.apply:
                self.cur.execute("SELECT 1 FROM ContractExtensions WHERE legacy_extension_id=?", legacy_id)
                if self.cur.fetchone():
                    continue
            typ = as_int(row.get('ContractExtensionTypeId'))
            price = as_decimal(row.get('ContractExtensionPrice'))
            has_dates = bool(row.get('ContractExtensionStartDate') or row.get('ContractExtensionEndDate'))
            if typ not in (1, 2, 3):
                typ = 2 if (has_dates and price is not None) else (1 if has_dates else 3)
                self.quarantine('t_ContractExtension', legacy_id, 'inferred_extension_type_%s' % typ, row)
            if price is None:
                price_mode, price_delta, resulting_price = 'none', None, None
            elif legacy_id in FINAL_SNAPSHOT_EXTENSION_IDS:
                price_mode, price_delta, resulting_price = 'final_snapshot', decimal.Decimal('0'), price
            else:
                price_mode, price_delta, resulting_price = 'delta', price, None
            if self.apply:
                self.cur.execute("""INSERT INTO ContractExtensions(
                    contract_id,legacy_extension_id,legacy_contract_id,extension_number,letter_number,
                    extension_date,extension_date_fa,title,start_date,start_date_fa,end_date,end_date_fa,
                    extension_type_id,legacy_price,price_mode,price_delta,resulting_price,internal_status,
                    review_note,legacy_attach_id,source_json)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'approved',N'تأیید داده تاریخی پس از کنترل کانورت R6',?,?)""",
                    contract_id, legacy_id, legacy_contract, row.get('ContractExtensionNumber'),
                    row.get('ContractExtensionLetterNumber'), clean(row.get('ContractExtensionDate')),
                    row.get('ContractExtensionDateFa'), str(row.get('ContractExtensionTitle') or 'الحاقیه مهاجرت‌شده'),
                    clean(row.get('ContractExtensionStartDate')), row.get('ContractExtensionStartDateFa'),
                    clean(row.get('ContractExtensionEndDate')), row.get('ContractExtensionEndDateFa'),
                    typ, price, price_mode, price_delta, resulting_price,
                    as_int(row.get('AttachId')), payload(row))
            self.stats['extensions_inserted'] += 1

    def statements(self, path):
        for row in read_rows(path):
            self.stats['statements_seen'] += 1
            legacy_id = as_int(row.get('ContractStatementId'))
            legacy_contract = as_int(row.get('ContractId'))
            contract_id = self.legacy_contracts.get(legacy_contract)
            if not contract_id:
                self.stats['statements_orphan'] += 1
                self.quarantine('t_ContractStatement', legacy_id, 'orphan_contract_%s' % legacy_contract, row)
                continue
            if self.apply:
                self.cur.execute("SELECT 1 FROM ContractStatements WHERE legacy_statement_id=?", legacy_id)
                if self.cur.fetchone():
                    continue
            period_year, period_month = fa_period(row.get('ContractStatementDateFa'), row.get('ContractStatementYear'))
            description = str(row.get('ContractStatementDescription') or '').strip()
            requested_base = as_decimal(row.get('ContractStatementPrice'))
            requested_vat = as_decimal(row.get('ContractStatementVat'))
            requested_total = None if requested_base is None else requested_base + (requested_vat or decimal.Decimal('0'))
            confirmed_base = as_decimal(row.get('ContractStatementConfirmedPrice'))
            confirmed_vat = as_decimal(row.get('ContractStatementConfirmedVat'))
            is_legacy_approved = confirmed_base is not None and confirmed_base > decimal.Decimal('0')
            business_status = 'employer_approved' if is_legacy_approved else (
                'sent' if row.get('ContractStatementDate') or row.get('ContractStatementDateFa') else 'draft'
            )
            confirmed_total = (confirmed_base + (confirmed_vat or decimal.Decimal('0'))
                               if is_legacy_approved else None)
            if is_legacy_approved:
                # The official approved value is the canonical full statement
                # amount in the new workflow; original requested values remain
                # available in source_json for audit.
                requested_base = confirmed_base
                requested_vat = confirmed_vat or decimal.Decimal('0')
                requested_total = confirmed_total
            else:
                confirmed_base = confirmed_vat = None
            migration_note = ('مهاجرت از سامانه قدیم؛ تأیید کارفرما از مبلغ تأییدشده تشخیص داده شد.'
                              if is_legacy_approved else
                              'مهاجرت از سامانه قدیم؛ این صورت‌وضعیت مبلغ تأییدشده ندارد و از مانده قرارداد کم نمی‌شود.')
            description = (description + '\n' + migration_note).strip()
            without_vat = as_bool(row.get('ContractStatementWithoutVat'))
            if (requested_vat or decimal.Decimal('0')) != 0 or (confirmed_vat or decimal.Decimal('0')) != 0:
                without_vat = False
            if self.apply:
                self.cur.execute("""INSERT INTO ContractStatements(
                    contract_id,legacy_statement_id,legacy_contract_id,statement_type_id,statement_number,
                    statement_date,statement_date_fa,letter_number,title,start_date,start_date_fa,end_date,end_date_fa,
                    requested_price,requested_without_vat,requested_vat,vat_factor_number,
                    confirmed_price,confirmed_without_vat,confirmed_vat,confirmed_vat_factor_number,
                    progress_percentage,legacy_contract_price,legacy_remain_price,legacy_attach_id,payment_method,
                    vat_percentage_id,period_year,period_month,vat_percentage_value,without_vat,
                    list_price_item_base_id,taxpayer_status_id,reference_number,tax_json,tax_sent_at,
                    tax_sent_at_fa,tax_error,description,region_id,legacy_user_id,business_status,source_json)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    contract_id, legacy_id, legacy_contract, as_int(row.get('ContractStatementTypeId')) or 1,
                    str(row.get('ContractStatementNumber') or legacy_id), clean(row.get('ContractStatementDate')),
                    row.get('ContractStatementDateFa'), row.get('ContractStatementLetterNumber'),
                    str(row.get('ContractStatementTitle') or 'صورت وضعیت مهاجرت‌شده'),
                    clean(row.get('ContractStatementStartDate')), row.get('ContractStatementStartDateFa'),
                    clean(row.get('ContractStatementEndDate')), row.get('ContractStatementEndDateFa'),
                    requested_total, requested_base, requested_vat,
                    row.get('ContractStatementVatFactorNumber'),
                    confirmed_total, confirmed_base, confirmed_vat,
                    row.get('ContractStatementConfirmedVatFactorNumber'),
                    as_int(row.get('ContractStatementProgressPercentage')),
                    as_decimal(row.get('ContractStatementContractPrice')),
                    as_decimal(row.get('ContractStatementRemainPrice')), as_int(row.get('AttachId')),
                    as_bool(row.get('ContractStatementPaymentMethod')), as_int(row.get('VatPercentageId')),
                    period_year, period_month, as_decimal(row.get('VatPercentageValue')),
                    without_vat,
                    as_int(row.get('ListPriceItemBaseId')), as_int(row.get('TaxPayerStatusId')),
                    row.get('ContractStatementReferenceNumber'), row.get('ContractStatementJson'),
                    clean(row.get('ContractStatementSendTaxDateTime')),
                    row.get('ContractStatementSendTaxDateTimeFa'), row.get('ContractStatementTaxError'),
                    description, as_int(row.get('RegionId')), as_int(row.get('UserId')),
                    business_status, payload(row))
            self.stats['statements_inserted'] += 1

    def finish(self):
        if self.apply and self.conn is not None:
            self.conn.commit()
        elif self.conn is not None:
            self.conn.rollback()
        return self.stats


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    parser = argparse.ArgumentParser(description='TaskHub v7 legacy contract migration')
    parser.add_argument('--contracts', default=os.path.join(here, 'migration_data', 't_Contract Data.xlsx'))
    parser.add_argument('--extensions', default=os.path.join(here, 'migration_data', 't_ContractExtension Data.xlsx'))
    parser.add_argument('--statements', default=os.path.join(here, 'migration_data', 't_ContractStatement Data.xlsx'))
    parser.add_argument('--apply', action='store_true', help='write changes (default is dry-run)')
    args = parser.parse_args()
    for path in (args.contracts, args.extensions, args.statements):
        if not os.path.isfile(path):
            parser.error('file not found: ' + path)
    conn = None
    try:
        if args.apply:
            import pyodbc
            from v7_features import init_v7_tables
            conn = pyodbc.connect(connection_string(), timeout=30)
            init_v7_tables(conn)
        m = Migrator(conn, apply=args.apply)
        m.load_existing_contracts()
        m.contracts(args.contracts)
        m.extensions(args.extensions)
        m.statements(args.statements)
        stats = m.finish()
        print(json.dumps({'mode': 'APPLY' if args.apply else 'DRY-RUN', **stats}, ensure_ascii=False, indent=2))
    except Exception:
        if conn is not None:
            try:conn.rollback()
            except Exception:pass
        raise
    finally:
        if conn is not None:conn.close()
    return 0


if __name__ == '__main__':
    sys.exit(main())
