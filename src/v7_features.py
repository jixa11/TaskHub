# -*- coding: utf-8 -*-
"""TaskHub v7 domain module.

This module deliberately contains no database connection settings.  It is
registered by :mod:`taskhub` and receives the already configured connection
factory/lock, so upgrading to v7 cannot accidentally change how SQL Server is
reached.
"""
from __future__ import annotations

import datetime as _dt
import decimal as _decimal
import hashlib
import io
import json
import mimetypes
import os
import tempfile

from flask import after_this_request, g, jsonify, request, send_file
import fa_font
from v7_domain import contract_remaining, expiry_level, validate_monthly_allocation
from rbac import user_has_permission
from team_scope import group_project_ids, has_company_scope, lead_group_ids

# Visibility is permission-driven; data scope is team-driven.


READ_FINANCE_ROLES = ('admin', 'manager', 'planner', 'finance', 'reporter')
WRITE_FINANCE_ROLES = ('admin', 'manager', 'planner', 'finance')
PLAN_WRITE_ROLES = ('admin', 'manager', 'planner')
MANAGER_ROLES = ('admin', 'manager')

JALALI_MONTH_NAMES = ('فروردین','اردیبهشت','خرداد','تیر','مرداد','شهریور',
                       'مهر','آبان','آذر','دی','بهمن','اسفند')


def _statement_is_employer_approved(row):
    """Return True only for a persisted employer approval decision.

    Current records use ``business_status=employer_approved``. Legacy records
    migrated into R9 may retain ``business_status=sent`` while already having
    ``employer_decision_at`` and a confirmed amount. The decision date is the
    authoritative migration signal; a positive amount alone is deliberately
    not enough because draft/pending records may also carry calculated values.
    """
    row = row or {}
    status = row.get('business_status')
    confirmed = row.get('confirmed_without_vat')
    if confirmed is None:
        confirmed = row.get('confirmed_price')
    try:
        has_confirmed_amount = _decimal.Decimal(str(confirmed or 0)) > 0
    except (_decimal.InvalidOperation, ValueError, TypeError):
        has_confirmed_amount = False
    return (
        status == 'employer_approved'
        or (
            row.get('employer_decision_at') is not None
            and status not in ('employer_rejected', 'revised', 'void')
            and has_confirmed_amount
        )
    )


def _jalali_to_gregorian(jy, jm, jd):
    """Convert a Jalali date to Gregorian without external packages."""
    jy = int(jy) + 1595
    days = -355668 + (365 * jy) + ((jy // 33) * 8) + (((jy % 33) + 3) // 4) + int(jd)
    if int(jm) < 7:
        days += (int(jm) - 1) * 31
    else:
        days += ((int(jm) - 7) * 30) + 186
    gy = 400 * (days // 146097)
    days %= 146097
    if days > 36524:
        days -= 1
        gy += 100 * (days // 36524)
        days %= 36524
        if days >= 365:
            days += 1
    gy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        gy += (days - 1) // 365
        days = (days - 1) % 365
    gd = days + 1
    month_days = [31, 29 if (gy % 4 == 0 and (gy % 100 != 0 or gy % 400 == 0)) else 28,
                  31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    gm = 0
    while gm < 12 and gd > month_days[gm]:
        gd -= month_days[gm]
        gm += 1
    return gy, gm + 1, gd


def _gregorian_to_jalali(gy, gm, gd):
    """Convert a Gregorian date to Jalali without external packages."""
    g_days = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    j_days = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    gy2, gm2, gd2 = int(gy) - 1600, int(gm) - 1, int(gd) - 1
    day_no = 365 * gy2 + (gy2 + 3) // 4 - (gy2 + 99) // 100 + (gy2 + 399) // 400
    day_no += sum(g_days[:gm2])
    if gm2 > 1 and ((int(gy) % 4 == 0 and int(gy) % 100 != 0) or int(gy) % 400 == 0):
        day_no += 1
    day_no += gd2
    j_day_no = day_no - 79
    j_np = j_day_no // 12053
    j_day_no %= 12053
    jy = 979 + 33 * j_np + 4 * (j_day_no // 1461)
    j_day_no %= 1461
    if j_day_no >= 366:
        jy += (j_day_no - 1) // 365
        j_day_no = (j_day_no - 1) % 365
    jm = 0
    while jm < 11 and j_day_no >= j_days[jm]:
        j_day_no -= j_days[jm]
        jm += 1
    return jy, jm + 1, j_day_no + 1


def _jalali_month_window(year, month):
    """Return [start, end) Gregorian dates and the previous Jalali month."""
    year, month = int(year), int(month)
    if not 1300 <= year <= 1600 or not 1 <= month <= 12:
        raise ValueError('ماه شمسی انتخاب‌شده نامعتبر است')
    next_year, next_month = (year + 1, 1) if month == 12 else (year, month + 1)
    previous_year, previous_month = (year - 1, 12) if month == 1 else (year, month - 1)
    start = _dt.date(*_jalali_to_gregorian(year, month, 1))
    end = _dt.date(*_jalali_to_gregorian(next_year, next_month, 1))
    previous_start = _dt.date(*_jalali_to_gregorian(previous_year, previous_month, 1))
    return start, end, previous_start, previous_year, previous_month


def _apply_r6_legacy_finance_repair(cur):
    """Repair the confirmed legacy mappings exactly once.

    The old system stores ``ContractStatementPrice`` as the amount before VAT;
    the ``...PriceWithoutVat`` columns in the exported workbook are mostly
    zero/NULL and must not be used as the source amount. Historical extensions
    are also activated after their delta/snapshot meaning has been resolved.
    """
    migration_key = 'r6_legacy_finance_repair_v1'
    cur.execute("SELECT 1 FROM AppSettings WHERE setting_key=?", migration_key)
    if cur.fetchone():
        return

    cur.execute(r"""
        UPDATE s SET
            requested_without_vat = raw.requested_base,
            requested_vat = raw.requested_vat,
            requested_price = CASE WHEN raw.requested_base IS NULL THEN NULL
                ELSE raw.requested_base + ISNULL(raw.requested_vat,0) END,
            confirmed_without_vat = raw.confirmed_base,
            confirmed_vat = raw.confirmed_vat,
            confirmed_price = CASE WHEN raw.confirmed_base IS NULL THEN NULL
                ELSE raw.confirmed_base + ISNULL(raw.confirmed_vat,0) END,
            without_vat = CASE
                WHEN ISNULL(raw.requested_vat,0)<>0 OR ISNULL(raw.confirmed_vat,0)<>0 THEN 0
                ELSE s.without_vat END,
            updated_at = GETDATE()
        FROM ContractStatements s
        CROSS APPLY (SELECT
            TRY_CONVERT(DECIMAL(18,0),JSON_VALUE(s.source_json,'$.ContractStatementPrice')) AS requested_base,
            TRY_CONVERT(DECIMAL(18,0),JSON_VALUE(s.source_json,'$.ContractStatementVat')) AS requested_vat,
            TRY_CONVERT(DECIMAL(18,0),JSON_VALUE(s.source_json,'$.ContractStatementConfirmedPrice')) AS confirmed_base,
            TRY_CONVERT(DECIMAL(18,0),JSON_VALUE(s.source_json,'$.ContractStatementConfirmedVat')) AS confirmed_vat
        ) raw
        WHERE s.legacy_statement_id IS NOT NULL
          AND ISJSON(s.source_json)=1
          AND raw.requested_base IS NOT NULL
    """)

    # Dataset-confirmed final snapshots. Every other monetary historical
    # extension is a delta; non-monetary extensions use price_mode=none.
    cur.execute(r"""
        UPDATE e SET
            price_mode = CASE
                WHEN e.legacy_price IS NULL THEN 'none'
                WHEN e.legacy_extension_id IN (3029,3030) THEN 'final_snapshot'
                ELSE 'delta' END,
            price_delta = CASE
                WHEN e.legacy_price IS NULL THEN NULL
                WHEN e.legacy_extension_id IN (3029,3030) THEN 0
                ELSE e.legacy_price END,
            resulting_price = CASE
                WHEN e.legacy_extension_id IN (3029,3030) THEN e.legacy_price
                ELSE NULL END,
            internal_status = 'approved',
            review_note = COALESCE(e.review_note,N'تأیید خودکار داده تاریخی پس از کنترل کانورت R6'),
            updated_at = GETDATE()
        FROM ContractExtensions e
        WHERE e.legacy_extension_id IS NOT NULL
    """)
    cur.execute(r"""
        UPDATE MigrationQuarantine SET resolution_status='resolved',
            resolved_at=GETDATE()
        WHERE source_table=N't_ContractExtension'
          AND reason=N'ambiguous_price_delta_or_final_snapshot'
          AND resolution_status='pending'
    """)
    cur.execute("""INSERT INTO AppSettings(setting_key,setting_value,updated_at)
        VALUES(?,N'legacy statements and extensions repaired',GETDATE())""",
                migration_key)


def _apply_r7_legacy_statement_status_repair(cur):
    """Map historical employer approval from the legacy confirmed columns.

    TaxPayerStatusId belongs to the tax workflow and is not a reliable employer
    decision flag.  A positive legacy confirmed base amount is therefore the
    only evidence used to mark a migrated statement as employer-approved.
    Historical rows without a positive confirmed amount remain sent/pending and
    do not reduce the contract balance.
    """
    migration_key = 'r7_legacy_statement_status_repair_v1'
    cur.execute("SELECT 1 FROM AppSettings WHERE setting_key=?", migration_key)
    if cur.fetchone():
        return

    cur.execute(r"""
        UPDATE s SET
            business_status = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0
                    THEN 'employer_approved'
                ELSE 'sent'
            END,
            requested_without_vat = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0 THEN raw.confirmed_base
                ELSE s.requested_without_vat END,
            requested_vat = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0 THEN ISNULL(raw.confirmed_vat,0)
                ELSE s.requested_vat END,
            requested_price = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0
                    THEN raw.confirmed_base + ISNULL(raw.confirmed_vat,0)
                ELSE s.requested_price END,
            confirmed_without_vat = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0 THEN raw.confirmed_base
                ELSE NULL END,
            confirmed_vat = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0
                    THEN ISNULL(raw.confirmed_vat,0)
                ELSE NULL END,
            confirmed_price = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0
                    THEN raw.confirmed_base + ISNULL(raw.confirmed_vat,0)
                ELSE NULL END,
            employer_decision_at = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0
                    THEN COALESCE(s.employer_decision_at,s.statement_date,s.updated_at,s.created_at,GETDATE())
                ELSE NULL END,
            employer_decision_note = CASE
                WHEN ISNULL(raw.confirmed_base,0) > 0
                    THEN COALESCE(s.employer_decision_note,N'تأیید کارفرما از مبلغ تأییدشده سامانه قدیم تشخیص داده شد')
                ELSE s.employer_decision_note END,
            updated_at = GETDATE()
        FROM ContractStatements s
        CROSS APPLY (SELECT
            TRY_CONVERT(DECIMAL(18,0),JSON_VALUE(s.source_json,'$.ContractStatementConfirmedPrice')) AS confirmed_base,
            TRY_CONVERT(DECIMAL(18,0),JSON_VALUE(s.source_json,'$.ContractStatementConfirmedVat')) AS confirmed_vat
        ) raw
        WHERE s.legacy_statement_id IS NOT NULL
          AND ISJSON(s.source_json)=1
    """)
    cur.execute("""INSERT INTO AppSettings(setting_key,setting_value,updated_at)
        VALUES(?,N'legacy employer approval mapped from confirmed amount',GETDATE())""",
                migration_key)


def init_v7_tables(conn):
    """Idempotently add the v7 schema to an existing v6 database."""
    cur = conn.cursor()
    sqls = [
        # Session presence: one user can legitimately have several devices;
        # the UI groups these rows and uses last_seen to remove stale sessions.
        "IF COL_LENGTH('Sessions','last_seen') IS NULL ALTER TABLE Sessions ADD last_seen DATETIME NULL",
        "IF COL_LENGTH('Sessions','ip_address') IS NULL ALTER TABLE Sessions ADD ip_address NVARCHAR(45) NULL",
        "IF COL_LENGTH('Sessions','user_agent') IS NULL ALTER TABLE Sessions ADD user_agent NVARCHAR(300) NULL",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Sessions_user_seen' AND object_id=OBJECT_ID('Sessions')) CREATE INDEX IX_Sessions_user_seen ON Sessions(user_id,last_seen)",

        # Existing tasks can optionally point at a contract and/or a planned
        # statement.  Monetary values intentionally never live on Tasks.
        "IF COL_LENGTH('Tasks','contract_id') IS NULL ALTER TABLE Tasks ADD contract_id INT NULL",
        "IF COL_LENGTH('Tasks','planned_statement_id') IS NULL ALTER TABLE Tasks ADD planned_statement_id INT NULL",

        # Every project has exactly one type.  These statements are intentionally
        # separate execute() calls. SQL Server compiles a whole batch before
        # running ALTER TABLE, so referencing project_type_id later in the same
        # batch raises error 207 on an existing v6 database.
        """IF OBJECT_ID('ProjectTypes','U') IS NULL CREATE TABLE ProjectTypes(
            id INT NOT NULL PRIMARY KEY,name NVARCHAR(120) NOT NULL,
            is_active BIT NOT NULL CONSTRAINT DF_ProjectTypes_active DEFAULT 1)""",
        """IF NOT EXISTS(SELECT 1 FROM ProjectTypes WHERE id=0)
              INSERT INTO ProjectTypes(id,name,is_active) VALUES(0,N'نوع تعیین نشده',1)""",
        """IF COL_LENGTH('Projects','project_type_id') IS NULL
              ALTER TABLE Projects ADD project_type_id INT NOT NULL
                  CONSTRAINT DF_Projects_project_type DEFAULT 0 WITH VALUES""",
        "UPDATE Projects SET project_type_id=0 WHERE project_type_id IS NULL",
        """IF EXISTS(SELECT 1 FROM sys.columns WHERE object_id=OBJECT_ID('Projects')
                     AND name='project_type_id' AND is_nullable=1)
              ALTER TABLE Projects ALTER COLUMN project_type_id INT NOT NULL""",
        """IF NOT EXISTS(SELECT 1 FROM sys.foreign_keys
                          WHERE name='FK_Projects_ProjectTypes'
                            AND parent_object_id=OBJECT_ID('Projects'))
              ALTER TABLE Projects ADD CONSTRAINT FK_Projects_ProjectTypes
                  FOREIGN KEY(project_type_id) REFERENCES ProjectTypes(id)""",
        """IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Projects_project_type' AND object_id=OBJECT_ID('Projects'))
              CREATE INDEX IX_Projects_project_type ON Projects(project_type_id,city_id) INCLUDE(name)""",
        """IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='UX_ProjectTypes_name' AND object_id=OBJECT_ID('ProjectTypes'))
              CREATE UNIQUE INDEX UX_ProjectTypes_name ON ProjectTypes(name)""",

        """IF OBJECT_ID('ContractTypes','U') IS NULL CREATE TABLE ContractTypes(
            id INT NOT NULL PRIMARY KEY,name NVARCHAR(120) NOT NULL,is_active BIT NOT NULL DEFAULT 1)""",
        """IF OBJECT_ID('ContractStatuses','U') IS NULL CREATE TABLE ContractStatuses(
            id INT NOT NULL PRIMARY KEY,name NVARCHAR(120) NOT NULL,is_active BIT NOT NULL DEFAULT 1)""",
        """IF OBJECT_ID('ContractCategories','U') IS NULL CREATE TABLE ContractCategories(
            id INT NOT NULL PRIMARY KEY,name NVARCHAR(120) NOT NULL,description NVARCHAR(500) NULL,is_active BIT NOT NULL DEFAULT 1)""",
        """IF OBJECT_ID('Contracts','U') IS NULL CREATE TABLE Contracts(
            id INT IDENTITY PRIMARY KEY,
            project_id INT NULL REFERENCES Projects(id),
            legacy_contract_id INT NULL,
            legacy_company_id INT NULL,
            title NVARCHAR(350) NOT NULL,
            legacy_employer_id INT NULL,
            employer_name NVARCHAR(250) NULL,
            contract_number NVARCHAR(50) NULL,
            base_price DECIMAL(18,0) NULL,
            notification_date DATE NULL,
            notification_date_fa NVARCHAR(50) NULL,
            notification_letter_number NVARCHAR(50) NULL,
            start_date DATE NULL,start_date_fa NVARCHAR(50) NULL,
            end_date DATE NULL,end_date_fa NVARCHAR(50) NULL,
            contract_type_id INT NULL,
            contract_status_id INT NULL,
            work_order_code NVARCHAR(50) NULL,
            financial_code NVARCHAR(50) NULL,
            register_in_sajat BIT NULL,
            register_date_sajat DATE NULL,register_date_sajat_fa NVARCHAR(50) NULL,
            list_price_version_id INT NULL,
            category_id INT NULL,
            insurance_coefficient DECIMAL(18,3) NULL,
            guarantee_coefficient DECIMAL(18,3) NULL,
            commercial_code NVARCHAR(50) NULL,
            scope_id INT NULL,is_new BIT NULL,tax_info_crn NVARCHAR(50) NULL,
            source_json NVARCHAR(MAX) NULL,
            payer_rating TINYINT NOT NULL DEFAULT 0,
            payer_note NVARCHAR(500) NULL,
            payer_locked BIT NOT NULL DEFAULT 0,
            payer_rated_by INT NULL,payer_rated_at DATETIME NULL,
            is_active BIT NOT NULL DEFAULT 1,
            created_by INT NULL,created_at DATETIME NOT NULL DEFAULT GETDATE(),
            updated_by INT NULL,updated_at DATETIME NULL,
            row_version ROWVERSION)""",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='UX_Contracts_legacy' AND object_id=OBJECT_ID('Contracts')) CREATE UNIQUE INDEX UX_Contracts_legacy ON Contracts(legacy_contract_id) WHERE legacy_contract_id IS NOT NULL",
        "IF COL_LENGTH('Contracts','source_json') IS NULL ALTER TABLE Contracts ADD source_json NVARCHAR(MAX) NULL",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Contracts_project_end' AND object_id=OBJECT_ID('Contracts')) CREATE INDEX IX_Contracts_project_end ON Contracts(project_id,end_date) INCLUDE(base_price,payer_rating,is_active)",

        """IF OBJECT_ID('ContractExtensionTypes','U') IS NULL CREATE TABLE ContractExtensionTypes(
            id INT NOT NULL PRIMARY KEY,name NVARCHAR(160) NOT NULL)""",
        """IF OBJECT_ID('ContractExtensions','U') IS NULL CREATE TABLE ContractExtensions(
            id INT IDENTITY PRIMARY KEY,
            contract_id INT NOT NULL REFERENCES Contracts(id),
            legacy_extension_id INT NULL,legacy_contract_id INT NULL,
            extension_number NVARCHAR(50) NULL,letter_number NVARCHAR(50) NULL,
            extension_date DATE NULL,extension_date_fa NVARCHAR(50) NULL,
            title NVARCHAR(200) NOT NULL,
            start_date DATE NULL,start_date_fa NVARCHAR(50) NULL,
            end_date DATE NULL,end_date_fa NVARCHAR(50) NULL,
            extension_type_id INT NULL,
            legacy_price DECIMAL(18,0) NULL,
            price_mode NVARCHAR(20) NOT NULL DEFAULT 'delta',
            price_delta DECIMAL(18,0) NULL,
            resulting_price DECIMAL(18,0) NULL,
            internal_status NVARCHAR(24) NOT NULL DEFAULT 'draft',
            review_note NVARCHAR(1000) NULL,
            approved_by INT NULL,approved_at DATETIME NULL,
            legacy_attach_id INT NULL,
            source_json NVARCHAR(MAX) NULL,
            is_active BIT NOT NULL DEFAULT 1,
            created_by INT NULL,created_at DATETIME NOT NULL DEFAULT GETDATE(),
            updated_by INT NULL,updated_at DATETIME NULL,row_version ROWVERSION)""",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='UX_Extensions_legacy' AND object_id=OBJECT_ID('ContractExtensions')) CREATE UNIQUE INDEX UX_Extensions_legacy ON ContractExtensions(legacy_extension_id) WHERE legacy_extension_id IS NOT NULL",
        "IF COL_LENGTH('ContractExtensions','source_json') IS NULL ALTER TABLE ContractExtensions ADD source_json NVARCHAR(MAX) NULL",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Extensions_contract_status' AND object_id=OBJECT_ID('ContractExtensions')) CREATE INDEX IX_Extensions_contract_status ON ContractExtensions(contract_id,internal_status,is_active)",

        """IF OBJECT_ID('ContractStatementTypes','U') IS NULL CREATE TABLE ContractStatementTypes(
            id INT NOT NULL PRIMARY KEY,name NVARCHAR(160) NOT NULL)""",
        """IF OBJECT_ID('ContractStatements','U') IS NULL CREATE TABLE ContractStatements(
            id INT IDENTITY PRIMARY KEY,
            contract_id INT NOT NULL REFERENCES Contracts(id),
            legacy_statement_id INT NULL,legacy_contract_id INT NULL,
            statement_type_id INT NOT NULL,
            statement_number NVARCHAR(50) NOT NULL,
            statement_date DATE NULL,statement_date_fa NVARCHAR(50) NULL,
            letter_number NVARCHAR(50) NULL,title NVARCHAR(250) NOT NULL,
            start_date DATE NULL,start_date_fa NVARCHAR(50) NULL,
            end_date DATE NULL,end_date_fa NVARCHAR(50) NULL,
            requested_price DECIMAL(18,0) NULL,
            requested_without_vat DECIMAL(18,0) NULL,
            requested_vat DECIMAL(18,0) NULL,
            vat_factor_number NVARCHAR(50) NULL,
            confirmed_price DECIMAL(18,0) NULL,
            confirmed_without_vat DECIMAL(18,0) NULL,
            confirmed_vat DECIMAL(18,0) NULL,
            confirmed_vat_factor_number NVARCHAR(50) NULL,
            progress_percentage INT NULL,
            legacy_contract_price DECIMAL(18,0) NULL,
            legacy_remain_price DECIMAL(18,0) NULL,
            payment_method BIT NULL,vat_percentage_id INT NULL,
            period_year INT NULL,period_month TINYINT NULL,
            vat_percentage_value DECIMAL(18,2) NULL,without_vat BIT NULL,
            list_price_item_base_id INT NULL,taxpayer_status_id INT NULL,
            reference_number NVARCHAR(MAX) NULL,tax_json NVARCHAR(MAX) NULL,
            tax_sent_at DATETIME NULL,tax_sent_at_fa NVARCHAR(50) NULL,
            tax_error NVARCHAR(MAX) NULL,description NVARCHAR(MAX) NULL,
            region_id INT NULL,legacy_user_id INT NULL,
            business_status NVARCHAR(28) NOT NULL DEFAULT 'draft',
            internal_approved_by INT NULL,internal_approved_at DATETIME NULL,
            sent_at DATETIME NULL,
            employer_decision_by INT NULL,employer_decision_at DATETIME NULL,
            employer_decision_note NVARCHAR(1000) NULL,
            supersedes_statement_id INT NULL,
            is_current BIT NOT NULL DEFAULT 1,is_active BIT NOT NULL DEFAULT 1,
            legacy_attach_id INT NULL,
            source_json NVARCHAR(MAX) NULL,
            created_by INT NULL,created_at DATETIME NOT NULL DEFAULT GETDATE(),
            updated_by INT NULL,updated_at DATETIME NULL,row_version ROWVERSION)""",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='UX_Statements_legacy' AND object_id=OBJECT_ID('ContractStatements')) CREATE UNIQUE INDEX UX_Statements_legacy ON ContractStatements(legacy_statement_id) WHERE legacy_statement_id IS NOT NULL",
        "IF COL_LENGTH('ContractStatements','source_json') IS NULL ALTER TABLE ContractStatements ADD source_json NVARCHAR(MAX) NULL",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Statements_contract_status' AND object_id=OBJECT_ID('ContractStatements')) CREATE INDEX IX_Statements_contract_status ON ContractStatements(contract_id,business_status,is_current,is_active)",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Statements_financial_period' AND object_id=OBJECT_ID('ContractStatements')) CREATE INDEX IX_Statements_financial_period ON ContractStatements(period_year,period_month,business_status,is_current,is_active) INCLUDE(contract_id,requested_without_vat,requested_price,confirmed_without_vat,confirmed_price,sent_at,employer_decision_at,statement_date)",
        """IF OBJECT_ID('ContractStatementItems','U') IS NULL CREATE TABLE ContractStatementItems(
            id INT IDENTITY PRIMARY KEY,statement_id INT NOT NULL REFERENCES ContractStatements(id),
            legacy_id INT NULL,item_code NVARCHAR(80) NULL,title NVARCHAR(500) NULL,
            quantity DECIMAL(18,4) NULL,unit_price DECIMAL(18,0) NULL,total_price DECIMAL(18,0) NULL,
            source_json NVARCHAR(MAX) NULL,created_at DATETIME NOT NULL DEFAULT GETDATE())""",
        """IF OBJECT_ID('ContractStatementTaxEvents','U') IS NULL CREATE TABLE ContractStatementTaxEvents(
            id INT IDENTITY PRIMARY KEY,statement_id INT NOT NULL REFERENCES ContractStatements(id),
            event_kind NVARCHAR(30) NOT NULL,status_id INT NULL,reference_number NVARCHAR(MAX) NULL,
            payload_json NVARCHAR(MAX) NULL,error_text NVARCHAR(MAX) NULL,
            event_at DATETIME NULL,event_at_fa NVARCHAR(50) NULL,legacy_id INT NULL,
            created_at DATETIME NOT NULL DEFAULT GETDATE())""",

        """IF OBJECT_ID('FinancialPlans','U') IS NULL CREATE TABLE FinancialPlans(
            id INT IDENTITY PRIMARY KEY,jalali_year INT NOT NULL,title NVARCHAR(200) NULL,
            annual_target DECIMAL(18,0) NOT NULL DEFAULT 0,
            status NVARCHAR(20) NOT NULL DEFAULT 'draft',is_current BIT NOT NULL DEFAULT 1,
            version_no INT NOT NULL DEFAULT 1,change_reason NVARCHAR(500) NULL,
            created_by INT NULL,created_at DATETIME NOT NULL DEFAULT GETDATE(),
            updated_by INT NULL,updated_at DATETIME NULL,row_version ROWVERSION)""",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='UX_FinancialPlans_year_current' AND object_id=OBJECT_ID('FinancialPlans')) CREATE UNIQUE INDEX UX_FinancialPlans_year_current ON FinancialPlans(jalali_year) WHERE is_current=1",
        """IF OBJECT_ID('FinancialPlanRevisions','U') IS NULL CREATE TABLE FinancialPlanRevisions(
            id INT IDENTITY PRIMARY KEY,plan_id INT NOT NULL REFERENCES FinancialPlans(id),
            old_target DECIMAL(18,0) NULL,new_target DECIMAL(18,0) NULL,
            reason NVARCHAR(500) NOT NULL,changed_by INT NULL,changed_at DATETIME NOT NULL DEFAULT GETDATE())""",
        """IF OBJECT_ID('FinancialPlanPeriods','U') IS NULL CREATE TABLE FinancialPlanPeriods(
            id INT IDENTITY PRIMARY KEY,plan_id INT NOT NULL REFERENCES FinancialPlans(id),
            month_no TINYINT NOT NULL,target_amount DECIMAL(18,0) NOT NULL DEFAULT 0,
            allocation_mode NVARCHAR(10) NOT NULL DEFAULT 'amount',allocation_percent DECIMAL(9,4) NULL,
            is_locked BIT NOT NULL DEFAULT 0,change_reason NVARCHAR(500) NULL,
            updated_by INT NULL,updated_at DATETIME NULL,
            CONSTRAINT UQ_FinancialPlanPeriods UNIQUE(plan_id,month_no),
            CONSTRAINT CK_FinancialPlanPeriods_month CHECK(month_no BETWEEN 1 AND 12))""",
        """IF OBJECT_ID('PlannedStatements','U') IS NULL CREATE TABLE PlannedStatements(
            id INT IDENTITY PRIMARY KEY,plan_id INT NOT NULL REFERENCES FinancialPlans(id),
            contract_id INT NOT NULL REFERENCES Contracts(id),
            month_no TINYINT NOT NULL,planned_date DATE NULL,planned_date_fa NVARCHAR(20) NULL,
            statement_type_id INT NULL,planned_amount DECIMAL(18,0) NOT NULL,
            title NVARCHAR(250) NOT NULL,note NVARCHAR(1000) NULL,
            status NVARCHAR(20) NOT NULL DEFAULT 'planned',actual_statement_id INT NULL,
            is_active BIT NOT NULL DEFAULT 1,created_by INT NULL,created_at DATETIME NOT NULL DEFAULT GETDATE(),
            updated_by INT NULL,updated_at DATETIME NULL,row_version ROWVERSION)""",
        """IF OBJECT_ID('PlannedStatementTasks','U') IS NULL CREATE TABLE PlannedStatementTasks(
            planned_statement_id INT NOT NULL REFERENCES PlannedStatements(id),
            task_id INT NOT NULL REFERENCES Tasks(id),linked_by INT NULL,linked_at DATETIME NOT NULL DEFAULT GETDATE(),
            CONSTRAINT PK_PlannedStatementTasks PRIMARY KEY(planned_statement_id,task_id))""",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Tasks_contract' AND object_id=OBJECT_ID('Tasks')) CREATE INDEX IX_Tasks_contract ON Tasks(contract_id,status)",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Tasks_planned_statement' AND object_id=OBJECT_ID('Tasks')) CREATE INDEX IX_Tasks_planned_statement ON Tasks(planned_statement_id)",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_TaskTimeLog_report' AND object_id=OBJECT_ID('TaskTimeLog')) CREATE INDEX IX_TaskTimeLog_report ON TaskTimeLog(started_at,user_id,task_id) INCLUDE(ended_at,seconds)",
        "IF NOT EXISTS(SELECT 1 FROM sys.foreign_keys WHERE name='FK_Tasks_Contracts') ALTER TABLE Tasks ADD CONSTRAINT FK_Tasks_Contracts FOREIGN KEY(contract_id) REFERENCES Contracts(id)",
        "IF NOT EXISTS(SELECT 1 FROM sys.foreign_keys WHERE name='FK_Tasks_PlannedStatements') ALTER TABLE Tasks ADD CONSTRAINT FK_Tasks_PlannedStatements FOREIGN KEY(planned_statement_id) REFERENCES PlannedStatements(id)",

        # Content-addressed blobs keep one physical file for identical uploads.
        """IF OBJECT_ID('FileBlobs','U') IS NULL CREATE TABLE FileBlobs(
            id INT IDENTITY PRIMARY KEY,sha256 CHAR(64) NOT NULL,mime_type NVARCHAR(120) NOT NULL,
            storage_name NVARCHAR(260) NOT NULL,thumbnail_name NVARCHAR(260) NULL,
            size_bytes BIGINT NOT NULL,width INT NULL,height INT NULL,
            created_at DATETIME NOT NULL DEFAULT GETDATE())""",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='UX_FileBlobs_sha' AND object_id=OBJECT_ID('FileBlobs')) CREATE UNIQUE INDEX UX_FileBlobs_sha ON FileBlobs(sha256)",
        """IF OBJECT_ID('Attachments','U') IS NULL CREATE TABLE Attachments(
            id INT IDENTITY PRIMARY KEY,entity_type NVARCHAR(30) NOT NULL,entity_id INT NOT NULL,
            blob_id INT NOT NULL REFERENCES FileBlobs(id),original_name NVARCHAR(260) NOT NULL,
            caption NVARCHAR(500) NULL,sort_order INT NOT NULL DEFAULT 0,version_no INT NOT NULL DEFAULT 1,
            is_active BIT NOT NULL DEFAULT 1,uploaded_by INT NULL,uploaded_at DATETIME NOT NULL DEFAULT GETDATE(),
            deleted_by INT NULL,deleted_at DATETIME NULL)""",
        "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Attachments_entity' AND object_id=OBJECT_ID('Attachments')) CREATE INDEX IX_Attachments_entity ON Attachments(entity_type,entity_id,is_active)",
        """IF OBJECT_ID('FileDownloadLog','U') IS NULL CREATE TABLE FileDownloadLog(
            id BIGINT IDENTITY PRIMARY KEY,attachment_id INT NOT NULL,user_id INT NULL,
            ip NVARCHAR(45) NULL,downloaded_at DATETIME NOT NULL DEFAULT GETDATE())""",
        """IF OBJECT_ID('MigrationQuarantine','U') IS NULL CREATE TABLE MigrationQuarantine(
            id INT IDENTITY PRIMARY KEY,source_table NVARCHAR(80) NOT NULL,legacy_id NVARCHAR(80) NULL,
            reason NVARCHAR(500) NOT NULL,payload_json NVARCHAR(MAX) NULL,
            resolution_status NVARCHAR(20) NOT NULL DEFAULT 'pending',resolved_by INT NULL,resolved_at DATETIME NULL,
            created_at DATETIME NOT NULL DEFAULT GETDATE())""",
        """IF OBJECT_ID('EntityAudit','U') IS NULL CREATE TABLE EntityAudit(
            id BIGINT IDENTITY PRIMARY KEY,entity_type NVARCHAR(40) NOT NULL,entity_id INT NULL,
            action NVARCHAR(40) NOT NULL,before_json NVARCHAR(MAX) NULL,after_json NVARCHAR(MAX) NULL,
            user_id INT NULL,ip NVARCHAR(45) NULL,created_at DATETIME NOT NULL DEFAULT GETDATE())""",
        """IF OBJECT_ID('AppSettings','U') IS NULL CREATE TABLE AppSettings(
            setting_key NVARCHAR(80) NOT NULL PRIMARY KEY,setting_value NVARCHAR(500) NULL,
            updated_by INT NULL,updated_at DATETIME NULL)""",
    ]
    for sql in sqls:
        cur.execute(sql)

    seeds = [
        ("ContractTypes", 0, "نوع تعیین نشده"),
        ("ContractTypes", 1, "فروش"),
        ("ContractTypes", 2, "توسعه"),
        ("ContractTypes", 3, "پشتیبانی"),
        ("ContractStatuses", 1, "سایر"),
        ("ContractStatuses", 2, "باز"),
        ("ContractStatuses", 3, "بسته"),
        ("ContractStatuses", 4, "خاتمه یافته"),
        ("ContractExtensionTypes", 1, "الحاقیه زمانی قرارداد"),
        ("ContractExtensionTypes", 2, "الحاقیه زمانی و ریالی قرارداد"),
        ("ContractExtensionTypes", 3, "الحاقیه ریالی قرارداد"),
        ("ContractStatementTypes", 1, "صورت وضعیت موقت"),
        ("ContractStatementTypes", 2, "صورت وضعیت دائم"),
        ("ContractStatementTypes", 3, "فاکتوری"),
        ("ContractStatementTypes", 4, "صورت وضعیت تعدیل"),
        ("ContractStatementTypes", 5, "صورت وضعیت تعدیل قطعی"),
        ("ContractCategories", 1, "فاکتور"),
        ("ContractCategories", 2, "قرارداد"),
        ("ContractCategories", 4, "قراردادهای در دست انعقاد"),
        ("ContractCategories", 5, "فاکتورهای در دست انعقاد"),
        ("ProjectTypes", 1, "CMS"),
        ("ProjectTypes", 2, "CMMS"),
        ("ProjectTypes", 3, "PM"),
        ("ProjectTypes", 4, "نور"),
        ("ProjectTypes", 5, "رها"),
        ("ProjectTypes", 6, "ETL"),
        ("ProjectTypes", 7, "صندوق نصیر"),
    ]
    for table, ident, name in seeds:
        cur.execute("IF NOT EXISTS(SELECT 1 FROM %s WHERE id=?) INSERT INTO %s(id,name) VALUES(?,?)" % (table, table), ident, ident, name)
    cur.execute("IF NOT EXISTS(SELECT 1 FROM AppSettings WHERE setting_key='contract_red_days') INSERT INTO AppSettings(setting_key,setting_value) VALUES('contract_red_days','15')")
    cur.execute("IF NOT EXISTS(SELECT 1 FROM AppSettings WHERE setting_key='contract_yellow_days') INSERT INTO AppSettings(setting_key,setting_value) VALUES('contract_yellow_days','45')")
    _apply_r6_legacy_finance_repair(cur)
    _apply_r7_legacy_statement_status_repair(cur)
    conn.commit()


def _latin_digits(value):
    return str(value).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789'))


def _number(value, default=None):
    if value in (None, ''):
        return default
    try:
        return _decimal.Decimal(_latin_digits(value).replace(',', ''))
    except Exception:
        return default


def _int(value, default=None):
    try:
        return int(_latin_digits(value)) if value not in (None, '') else default
    except Exception:
        return default


def _date(value):
    if value in (None, ''):
        return None
    if isinstance(value, (_dt.date, _dt.datetime)):
        return value.date() if isinstance(value, _dt.datetime) else value
    try:
        return _dt.date.fromisoformat(str(value)[:10])
    except Exception:
        return None


def _jsonable(obj):
    if obj is None or isinstance(obj, (str, int, float, bool)):
        return obj
    if isinstance(obj, (_dt.date, _dt.datetime)):
        return obj.isoformat()
    if isinstance(obj, _decimal.Decimal):
        return str(obj)
    if isinstance(obj, (bytes, bytearray, memoryview)):
        return bytes(obj).hex()
    if isinstance(obj, dict):
        return {str(k): _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    return str(obj)


def register_v7_routes(app, get_conn, db_lock, require_auth, require_roles,
                       rows_to_list, audit, base_dir):
    """Register v7 endpoints against the existing Flask application."""

    def _ok(**kwargs):
        payload = {'ok': True}; payload.update(kwargs); return jsonify(_jsonable(payload))

    def _err(message, status=200, **kwargs):
        payload = {'ok': False, 'error': message}; payload.update(kwargs)
        return jsonify(_jsonable(payload)), status

    def _one(cur, sql, *params):
        cur.execute(sql, *params)
        row = cur.fetchone()
        if not row:
            return None
        cols = [x[0] for x in cur.description]
        return {cols[i]: _jsonable(row[i]) for i in range(len(cols))}

    def _entity_audit(cur, entity_type, entity_id, action, before=None, after=None):
        cur.execute("INSERT INTO EntityAudit(entity_type,entity_id,action,before_json,after_json,user_id,ip) VALUES(?,?,?,?,?,?,?)",
                    entity_type, entity_id, action,
                    json.dumps(_jsonable(before), ensure_ascii=False) if before is not None else None,
                    json.dumps(_jsonable(after), ensure_ascii=False) if after is not None else None,
                    g.user['id'], (request.headers.get('X-Forwarded-For') or request.remote_addr or '')[:45])

    def _team_allowed(cur, user, team_id):
        if not team_id:
            return False
        if has_company_scope(user):
            cur.execute("SELECT 1 FROM Teams WHERE id=? AND is_active=1", team_id)
        else:
            cur.execute("""SELECT 1 FROM TeamMembers tm JOIN Teams t ON t.id=tm.team_id
                WHERE tm.team_id=? AND tm.user_id=? AND tm.is_active=1 AND t.is_active=1""",
                        team_id, user['id'])
        return cur.fetchone() is not None

    def _project_allowed(cur, user, project_id):
        if has_company_scope(user):
            return True
        cur.execute("""SELECT TOP 1 1 FROM ProjectTeams pt JOIN TeamMembers tm
            ON tm.team_id=pt.team_id WHERE pt.project_id=? AND pt.is_active=1
              AND tm.user_id=? AND tm.is_active=1""", project_id, user['id'])
        return cur.fetchone() is not None

    def _task_allowed(cur, user, task_id):
        if has_company_scope(user):
            return True
        cur.execute("""SELECT TOP 1 1 FROM Tasks t JOIN ProjectTeams pt ON pt.id=t.project_team_id
            JOIN TeamMembers tm ON tm.team_id=pt.team_id
            WHERE t.id=? AND pt.is_active=1 AND tm.user_id=? AND tm.is_active=1""",
                    task_id, user['id'])
        return cur.fetchone() is not None

    def _task_where(user):
        data = request.get_json(silent=True) or {}
        requested_team = _int(data.get('_team_scope'))
        conditions, params = [], []

        # R10 authorization boundary: non-global roles are always restricted to
        # active teams, even when the administrator grants tasks.view_all.
        if not has_company_scope(user):
            conditions.append("""EXISTS(SELECT 1 FROM ProjectTeams auth_pt
                JOIN TeamMembers auth_tm ON auth_tm.team_id=auth_pt.team_id
                WHERE auth_pt.id=t.project_team_id AND auth_pt.is_active=1
                  AND auth_tm.user_id=? AND auth_tm.is_active=1)""")
            params.append(user['id'])

        if not user_has_permission(user, 'tasks.view_all'):
            if user_has_permission(user, 'tasks.self_manage'):
                conditions.append("t.project_id=?")
                params.append(user.get('project_id'))
            elif (user_has_permission(user, 'tasks.work') and
                  not user_has_permission(user, 'tasks.assign') and
                  not user_has_permission(user, 'tasks.approve') and
                  not user_has_permission(user, 'tasks.triage')):
                # R15: a group lead also sees the tasks of their sub-group, which
                # they edit, approve and delete.
                sees_group = any(user_has_permission(user, key) for key in
                                 ('tasks.edit_group', 'tasks.delete_group', 'tasks.approve_group'))
                conditions.append("(t.staff_id=? OR EXISTS(SELECT 1 FROM TaskAssignees tx "
                                  "WHERE tx.task_id=t.id AND tx.user_id=?) OR (?=1 AND EXISTS("
                                  "SELECT 1 FROM WorkGroupMembers gm JOIN WorkGroups wg ON wg.id=gm.group_id "
                                  "AND wg.is_active=1 JOIN Users gmu ON gmu.id=gm.user_id "
                                  "WHERE wg.lead_id=? AND gmu.is_active=1 AND gm.user_id=t.staff_id)))")
                params.extend([user['id'], user['id'], 1 if sees_group else 0, user['id']])
            elif user_has_permission(user, 'tasks.triage'):
                conditions.append("t.status IN ('registered','approved','rejected')")
            elif not user_has_permission(user, 'tasks.view'):
                conditions.append("1=0")
        if requested_team:
            conditions.append("""EXISTS(SELECT 1 FROM ProjectTeams scope_pt
                WHERE scope_pt.id=t.project_team_id AND scope_pt.team_id=?
                  AND scope_pt.is_active=1)""")
            params.append(requested_team)
        return (" WHERE " + " AND ".join(conditions) + " " if conditions else " ", params)

    def _team_scope_condition(user, data, entity_kind, alias):
        """Return an entity-team authorization condition.

        Global roles may optionally filter to one team. Every other role is
        always limited to active team memberships; an explicit team selector
        can only narrow that authorized set.
        """
        selected = _int((data or {}).get('_team_scope'))
        if entity_kind == 'contract':
            relation = """SELECT 1 FROM ContractProjectTeams scope_cpt
                JOIN ProjectTeams scope_pt ON scope_pt.id=scope_cpt.project_team_id"""
            relation_where = (
                "scope_cpt.contract_id=%s.id AND scope_cpt.is_active=1 "
                "AND scope_pt.is_active=1" % alias
            )
        elif entity_kind == 'extension':
            # An extension has no team link of its own; it inherits the one
            # the finance specialist set on its contract.
            relation = """SELECT 1 FROM ContractProjectTeams scope_cpt
                JOIN ProjectTeams scope_pt ON scope_pt.id=scope_cpt.project_team_id"""
            relation_where = (
                "scope_cpt.contract_id=%s.contract_id AND scope_cpt.is_active=1 "
                "AND scope_pt.is_active=1" % alias
            )
        elif entity_kind == 'statement':
            relation = """SELECT 1 FROM ContractStatementTeams scope_cst
                JOIN ProjectTeams scope_pt
                  ON scope_pt.id=scope_cst.project_team_id"""
            relation_where = (
                "scope_cst.statement_id=%s.id AND scope_cst.is_active=1 "
                "AND scope_pt.is_active=1" % alias
            )
        else:
            relation = "SELECT 1 FROM ProjectTeams scope_pt"
            relation_where = (
                "scope_pt.id=%s.project_team_id AND scope_pt.is_active=1" % alias
            )
        params = []
        if not has_company_scope(user):
            relation += " JOIN TeamMembers scope_tm ON scope_tm.team_id=scope_pt.team_id"
            relation_where += " AND scope_tm.user_id=? AND scope_tm.is_active=1"
            params.append(user['id'])
        if selected:
            relation_where += " AND scope_pt.team_id=?"
            params.append(selected)
        if has_company_scope(user) and not selected:
            return None, []
        return "EXISTS(%s WHERE %s)" % (relation, relation_where), params

    def _statement_scope_share(user, data, alias='s'):
        """Return the statement allocation percentage visible to the caller."""
        selected = _int((data or {}).get('_team_scope'))
        if has_company_scope(user) and not selected:
            return "CAST(100 AS DECIMAL(9,4))", []
        if selected:
            if has_company_scope(user):
                return (
                    "ISNULL((SELECT SUM(share_cst.allocation_percent) FROM "
                    "ContractStatementTeams share_cst JOIN ProjectTeams share_pt "
                    "ON share_pt.id=share_cst.project_team_id WHERE "
                    "share_cst.statement_id=%s.id AND share_cst.is_active=1 "
                    "AND share_pt.is_active=1 AND share_pt.team_id=?),0)" % alias,
                    [selected],
                )
            return (
                "ISNULL((SELECT SUM(share_cst.allocation_percent) FROM "
                "ContractStatementTeams share_cst JOIN ProjectTeams share_pt "
                "ON share_pt.id=share_cst.project_team_id JOIN TeamMembers share_tm "
                "ON share_tm.team_id=share_pt.team_id WHERE "
                "share_cst.statement_id=%s.id AND share_cst.is_active=1 "
                "AND share_pt.is_active=1 AND share_tm.user_id=? "
                "AND share_tm.is_active=1 AND share_pt.team_id=?),0)" % alias,
                [user['id'], selected],
            )
        return (
            "ISNULL((SELECT SUM(share_cst.allocation_percent) FROM "
            "ContractStatementTeams share_cst JOIN ProjectTeams share_pt "
            "ON share_pt.id=share_cst.project_team_id JOIN TeamMembers share_tm "
            "ON share_tm.team_id=share_pt.team_id WHERE "
            "share_cst.statement_id=%s.id AND share_cst.is_active=1 "
            "AND share_pt.is_active=1 AND share_tm.user_id=? "
            "AND share_tm.is_active=1),0)" % alias,
            [user['id']],
        )

    @app.route('/api/core_data', methods=['POST'])
    @require_auth
    def api_v7_core_data():
        """Return only the datasets permitted for the authenticated role.

        Menu visibility is a convenience; this endpoint is the real data
        boundary. Related lookup rows are included only when another allowed
        dataset needs them (for example project names for visible tasks).
        """
        try:
            see_tasks = user_has_permission(g.user, 'tasks.view')
            see_projects = user_has_permission(g.user, 'projects.view') or see_tasks
            see_cities = user_has_permission(g.user, 'cities.view') or see_projects
            see_categories = user_has_permission(g.user, 'task_categories.view') or see_tasks
            where, params = _task_where(g.user) if see_tasks else (' WHERE 1=0 ', [])
            with db_lock:
                c = get_conn(); cur = c.cursor()
                project_conditions, project_params = [], []
                if not has_company_scope(g.user):
                    project_conditions.append("""EXISTS(SELECT 1 FROM ProjectTeams auth_pt
                        JOIN TeamMembers auth_tm ON auth_tm.team_id=auth_pt.team_id
                        WHERE auth_pt.project_id=p.id AND auth_pt.is_active=1
                          AND auth_tm.user_id=? AND auth_tm.is_active=1)""")
                    project_params.append(g.user['id'])
                    # R16: the people of a group with projects see only those.
                    group_projects = sorted(group_project_ids(cur, g.user['id']))
                    if group_projects:
                        project_conditions.append("p.id IN (%s)" % ",".join("?" for _ in group_projects))
                        project_params.extend(group_projects)
                if user_has_permission(g.user, 'tasks.self_manage') and g.user.get('project_id'):
                    project_conditions.append("p.id=?")
                    project_params.append(g.user.get('project_id'))
                elif not user_has_permission(g.user, 'projects.view') and not see_tasks:
                    project_conditions.append("1=0")
                project_where = (" WHERE " + " AND ".join(project_conditions)) if project_conditions else ""

                projects, project_types = [], []
                if see_projects:
                    cur.execute("""SELECT p.id,p.city_id,p.name,p.project_type_id,
                        pt.name AS project_type_name,c.name AS cname,p.created_at,
                        (SELECT COUNT(*) FROM ProjectTeams px WHERE px.project_id=p.id
                          AND px.is_active=1) AS active_team_count,
                        STUFF((SELECT N'، '+tx.name FROM ProjectTeams ptx JOIN Teams tx
                          ON tx.id=ptx.team_id WHERE ptx.project_id=p.id AND ptx.is_active=1
                          ORDER BY tx.name FOR XML PATH(''),TYPE).value('.','NVARCHAR(MAX)'),1,2,N'') AS team_names
                        FROM Projects p JOIN Cities c ON c.id=p.city_id
                        JOIN ProjectTypes pt ON pt.id=p.project_type_id
                        """ + project_where + " ORDER BY c.name,pt.name,p.name",
                                project_params)
                    projects = rows_to_list(cur)
                    # Types are intentionally dynamic; custom rows such as
                    # «بینا» and «سپنتا» are never hard-coded or overwritten.
                    cur.execute("SELECT id,name FROM ProjectTypes WHERE id>0 AND is_active=1 ORDER BY id")
                    project_types = rows_to_list(cur)

                cities = []
                if see_cities:
                    # Cities are shared master data across the company. Team
                    # boundaries apply to projects and operational entities,
                    # not to duplicate copies of a city for each team.
                    cur.execute("SELECT id,name,created_at FROM Cities ORDER BY name")
                    cities = rows_to_list(cur)

                categories = []
                if see_categories:
                    cur.execute("SELECT id,name,description FROM TaskCategories ORDER BY name")
                    categories = rows_to_list(cur)

                tasks, helpers, contributions = [], [], []
                if see_tasks:
                    sql = """SELECT t.id,t.title,t.type,t.status,t.category_id,t.project_id,t.contact_id,t.staff_id,
                        t.date_recv,t.date_delivery,t.description,t.solution,t.test_notes,t.priority,t.created_by,
                        t.started_at,t.submitted_at,t.completed_at,t.work_seconds,t.reject_reason,t.active_actor_id,
                        t.contract_id,t.planned_statement_id,t.project_team_id,t.progress_weight,
                        p.project_type_id,pt.name AS project_type_name,
                        p.name AS pname,c.name AS cname,c.id AS city_id,emp.display_name AS cont_name,
                        emp.phone AS cont_phone,emp.phone2 AS cont_phone2,sup.display_name AS staff_name,
                        cat.name AS cat_name,act.display_name AS active_actor_name,
                        co.title AS contract_title,co.contract_number,
                        ptm.team_id,team.name AS team_name
                        FROM Tasks t LEFT JOIN Projects p ON p.id=t.project_id
                        LEFT JOIN ProjectTypes pt ON pt.id=p.project_type_id LEFT JOIN Cities c ON c.id=p.city_id
                        LEFT JOIN Users emp ON emp.id=t.contact_id LEFT JOIN Users sup ON sup.id=t.staff_id
                        LEFT JOIN TaskCategories cat ON cat.id=t.category_id LEFT JOIN Users act ON act.id=t.active_actor_id
                        LEFT JOIN Contracts co ON co.id=t.contract_id
                        LEFT JOIN ProjectTeams ptm ON ptm.id=t.project_team_id
                        LEFT JOIN Teams team ON team.id=ptm.team_id""" + where + """
                        ORDER BY CASE t.status WHEN 'registered' THEN 0 WHEN 'assigned' THEN 1 WHEN 'returned' THEN 2
                        WHEN 'doing' THEN 3 WHEN 'paused' THEN 4 WHEN 'pending_approval' THEN 5 ELSE 6 END,
                        t.priority,t.id DESC"""
                    cur.execute(sql, params)
                    tasks = rows_to_list(cur)
                    task_ids = [int(x['id']) for x in tasks]
                    if task_ids:
                        marks = ','.join('?' for _ in task_ids)
                        cur.execute("SELECT ta.task_id,ta.user_id,u.display_name,u.username FROM TaskAssignees ta JOIN Users u ON u.id=ta.user_id WHERE ta.task_id IN (%s)" % marks, task_ids)
                        helpers = rows_to_list(cur)
                        cur.execute("SELECT task_id,user_id,SUM(ISNULL(seconds,0)) AS total_seconds FROM TaskTimeLog WHERE ended_at IS NOT NULL AND task_id IN (%s) GROUP BY task_id,user_id" % marks, task_ids)
                        contributions = rows_to_list(cur)
            return _ok(cities=cities, projects=projects, project_types=project_types,
                       categories=categories, tasks=tasks, helpers=helpers,
                       contributions=contributions)
        except Exception as exc:
            return _err(str(exc), rows=[])

    @app.route('/api/master_save', methods=['POST'])
    @require_auth
    def api_v7_master_save():
        try:
            d = request.get_json() or {}; kind = d.get('kind'); ident = _int(d.get('id'))
            name = (d.get('name') or '').strip(); actor = g.user
            if kind not in ('city', 'project', 'category') or not name:
                return _err('اطلاعات ناقص است')
            with db_lock:
                c = get_conn(); cur = c.cursor()
                if kind == 'city':
                    cur.execute("""SELECT TOP 1 id,name FROM Cities
                        WHERE LTRIM(RTRIM(name))=LTRIM(RTRIM(?))
                          AND (? IS NULL OR id<>?) ORDER BY id""", name, ident, ident)
                    duplicate = cur.fetchone()
                    if duplicate:
                        if not ident:
                            return _ok(id=int(duplicate[0]), existing=True,
                                       message='این شهر از قبل وجود دارد و همان رکورد مشترک استفاده شد')
                        return _err('شهری با این نام قبلاً ثبت شده است')
                    if ident:
                        cur.execute("UPDATE Cities SET name=? WHERE id=?", name, ident)
                    else:
                        cur.execute("INSERT INTO Cities(name) OUTPUT INSERTED.id VALUES(?)", name)
                        ident = int(cur.fetchone()[0])
                elif kind == 'project':
                    city_id = _int(d.get('city_id')); project_type_id = _int(d.get('project_type_id'))
                    team_id = _int(d.get('team_id'))
                    if not city_id: return _err('شهر پروژه الزامی است')
                    if not project_type_id or project_type_id <= 0:
                        return _err('انتخاب نوع پروژه الزامی است')
                    if not _one(cur, "SELECT id FROM ProjectTypes WHERE id=? AND id>0 AND is_active=1", project_type_id):
                        return _err('نوع پروژه انتخاب‌شده معتبر نیست')
                    if ident:
                        if not _project_allowed(cur, actor, ident):
                            return _err('این پروژه خارج از محدوده تیم شماست',403)
                        cur.execute("UPDATE Projects SET city_id=?,project_type_id=?,name=? WHERE id=?",
                                    city_id, project_type_id, name, ident)
                    else:
                        if not team_id: return _err('انتخاب تیم مسئول پروژه الزامی است')
                        if not _team_allowed(cur, actor, team_id):
                            return _err('به تیم انتخاب‌شده دسترسی ندارید',403)
                        cur.execute("""INSERT INTO Projects(city_id,project_type_id,name)
                            OUTPUT INSERTED.id VALUES(?,?,?)""", city_id, project_type_id, name)
                        ident=int(cur.fetchone()[0])
                        cur.execute("""IF EXISTS(SELECT 1 FROM TeamProjectTypes WHERE team_id=? AND project_type_id=?)
                            UPDATE TeamProjectTypes SET is_active=1 WHERE team_id=? AND project_type_id=?
                            ELSE INSERT INTO TeamProjectTypes(team_id,project_type_id,assignment_role,is_active,assigned_by)
                              VALUES(?,?,'primary',1,?)""",
                            team_id,project_type_id,team_id,project_type_id,team_id,project_type_id,actor['id'])
                        cur.execute("""INSERT INTO ProjectTeams(project_id,team_id,title,is_primary,is_active,created_by)
                            VALUES(?,?,?,1,1,?)""",ident,team_id,name,actor['id'])
                else:
                    desc = (d.get('description') or '').strip() or None
                    if ident: cur.execute("UPDATE TaskCategories SET name=?,description=? WHERE id=?", name, desc, ident)
                    else: cur.execute("INSERT INTO TaskCategories(name,description) VALUES(?,?)", name, desc)
                c.commit()
            audit('master_save', '%s #%s' % (kind, ident or 'new'))
            return _ok(id=ident)
        except Exception as exc:
            message = str(exc)
            if locals().get('kind') == 'city' and (
                'UNIQUE' in message.upper() or 'duplicate' in message.lower()
            ):
                return _err('این شهر از قبل ثبت شده است؛ شهرها بین همه تیم‌ها مشترک هستند')
            return _err(message)

    @app.route('/api/master_delete', methods=['POST'])
    @require_auth
    def api_v7_master_delete():
        try:
            d = request.get_json() or {}; kind = d.get('kind'); ident = _int(d.get('id')); actor=g.user
            tables = {'city': 'Cities', 'project': 'Projects', 'category': 'TaskCategories'}
            if kind not in tables or not ident: return _err('اطلاعات ناقص است')
            with db_lock:
                c = get_conn(); cur = c.cursor()
                if kind=='project' and not _project_allowed(cur,actor,ident):
                    return _err('این پروژه خارج از محدوده تیم شماست',403)
                if kind=='city':
                    cur.execute("SELECT 1 FROM Projects WHERE city_id=?", ident)
                    if cur.fetchone():
                        return _err('این شهر دارای پروژه است و قابل حذف نیست؛ شهرها بین همه تیم‌ها مشترک هستند')
                cur.execute("DELETE FROM %s WHERE id=?" % tables[kind], ident); c.commit()
            audit('master_delete', '%s #%s' % (kind, ident))
            return _ok()
        except Exception as exc:
            return _err(str(exc))

    @app.route('/api/task_contract_options', methods=['POST'])
    @require_auth
    def api_v7_task_contract_options():
        """Minimal, team-scoped contract data for the task form.

        This endpoint deliberately excludes prices and financial details. It
        allows the contract selector to work without granting access to the
        full contracts page.
        """
        try:
            with db_lock:
                c = get_conn(); cur = c.cursor()
                sql = """SELECT co.id,co.title,co.contract_number,co.project_id,
                    p.name AS project_name,ci.id AS city_id,ci.name AS city_name,
                    ct.name AS contract_type_name,cc.name AS category_name,
                    cs.name AS contract_status_name,
                    CASE WHEN co.contract_status_id IN (3,4) THEN 1 ELSE 0 END AS is_closed
                    FROM Contracts co JOIN Projects p ON p.id=co.project_id
                    JOIN Cities ci ON ci.id=p.city_id
                    LEFT JOIN ContractTypes ct ON ct.id=co.contract_type_id
                    LEFT JOIN ContractCategories cc ON cc.id=co.category_id
                    LEFT JOIN ContractStatuses cs ON cs.id=co.contract_status_id
                    WHERE co.is_active=1"""
                params = []
                if not has_company_scope(g.user):
                    sql += """ AND EXISTS(SELECT 1 FROM ProjectTeams scope_pt
                        JOIN TeamMembers scope_tm ON scope_tm.team_id=scope_pt.team_id
                        WHERE scope_pt.project_id=p.id AND scope_pt.is_active=1
                          AND scope_tm.user_id=? AND scope_tm.is_active=1)"""
                    params.append(g.user['id'])
                sql += """ ORDER BY CASE WHEN co.contract_status_id IN (3,4) THEN 1 ELSE 0 END,
                    ci.name,p.name,co.title,co.id"""
                cur.execute(sql, params)
                rows = rows_to_list(cur)
            return _ok(rows=rows)
        except Exception as exc:
            return _err(str(exc), rows=[])

    @app.route('/api/task_save', methods=['POST'])
    @require_auth
    def api_v7_task_save():
        try:
            d = request.get_json() or {}; ident = _int(d.get('id')); u = g.user; role = u['role']
            with db_lock:
                c = get_conn(); cur = c.cursor()
                old = _one(cur, "SELECT * FROM Tasks WHERE id=?", ident) if ident else None
                if ident and not old: return _err('تسک یافت نشد')
                if ident and not _task_allowed(cur,u,ident): return _err('این تسک خارج از محدوده تیم شماست',403)
                full_task_access = user_has_permission(u, 'tasks.edit' if ident else 'tasks.create')
                if not ident and not full_task_access and user_has_permission(u, 'tasks.create_for_group'):
                    # R14 group lead: a new task for themselves or for a member of
                    # their sub-group (LeadMembers), nobody else. It enters the
                    # normal flow already assigned; weight, contract links and
                    # approval stay with the planner, so nothing here earns early.
                    allowed_staff = lead_group_ids(cur, u['id'])
                    staff_id = _int(d.get('staff_id')) or int(u['id'])
                    if staff_id not in allowed_staff:
                        return _err('سرگروه فقط برای خودش و اعضای زیرگروهش تسک ثبت می‌کند', 403)
                    d['staff_id'] = staff_id
                    d['status'] = 'assigned'
                    d['progress_weight'] = 1
                    d['contract_id'] = None
                    d['planned_statement_id'] = None
                    full_task_access = True
                if (ident and not full_task_access and user_has_permission(u, 'tasks.edit_group')
                        and _int(old.get('staff_id')) in lead_group_ids(cur, u['id'])):
                    # R15 group lead: edits a task of their own or of their
                    # sub-group and may hand it to another of those people. The
                    # status still moves only through start/forward/approve, and
                    # weight and contract links stay with the planner.
                    group = lead_group_ids(cur, u['id'])
                    staff_id = _int(d.get('staff_id')) or _int(old.get('staff_id'))
                    if staff_id not in group:
                        return _err('سرگروه تسک را فقط بین خودش و اعضای زیرگروهش جابه‌جا می‌کند', 403)
                    d['staff_id'] = staff_id
                    d['status'] = old.get('status')
                    d['progress_weight'] = str(old.get('progress_weight') or 1)
                    full_task_access = True
                if not full_task_access and user_has_permission(u, 'tasks.self_manage'):
                    if ident:
                        if old.get('created_by') != u['id'] or old.get('status') != 'registered' or old.get('project_id') != u.get('project_id'):
                            return _err('اجازه ویرایش این تسک را ندارید', 403)
                    d['project_id'] = u.get('project_id'); d['contact_id'] = u['id']; d['staff_id'] = None
                    if not ident: d['status'] = 'registered'
                elif not full_task_access and user_has_permission(u, 'tasks.work'):
                    if not ident: return _err('برای ساخت مستقیم تسک دسترسی لازم را ندارید', 403)
                    cur.execute("SELECT 1 FROM Tasks t WHERE t.id=? AND (t.staff_id=? OR EXISTS(SELECT 1 FROM TaskAssignees x WHERE x.task_id=t.id AND x.user_id=?))", ident, u['id'], u['id'])
                    if not cur.fetchone(): return _err('این تسک به شما واگذار نشده است', 403)
                    cur.execute("UPDATE Tasks SET solution=? WHERE id=?", (d.get('solution') or '').strip() or None, ident)
                    _entity_audit(cur, 'task', ident, 'worker_edit', old, {'solution': d.get('solution')})
                    c.commit(); return _ok(id=ident)
                elif not full_task_access:
                    return _err('اجازه ویرایش تسک را ندارید', 403)
                title = (d.get('title') or '').strip(); project_id = _int(d.get('project_id'))
                if not title or not project_id: return _err('عنوان و پروژه الزامی است')
                if not has_company_scope(u):
                    # R16: a group's people raise work only on the group's
                    # projects; an older task already elsewhere keeps its project.
                    group_projects = group_project_ids(cur, u['id'])
                    if (group_projects and project_id not in group_projects
                            and not (ident and project_id == _int(old.get('project_id')))):
                        return _err('این پروژه جزو پروژه‌های گروه شما نیست', 403)
                project_team_id = _int(d.get('project_team_id'))
                if not project_team_id:
                    if has_company_scope(u):
                        cur.execute("""SELECT id FROM ProjectTeams WHERE project_id=?
                            AND is_active=1 ORDER BY is_primary DESC,id""", project_id)
                    else:
                        cur.execute("""SELECT pt.id FROM ProjectTeams pt JOIN TeamMembers tm
                            ON tm.team_id=pt.team_id WHERE pt.project_id=? AND pt.is_active=1
                              AND tm.user_id=? AND tm.is_active=1 ORDER BY pt.is_primary DESC,pt.id""",
                                    project_id,u['id'])
                    candidates = [int(x[0]) for x in cur.fetchall()]
                    if len(candidates) == 1:
                        project_team_id = candidates[0]
                    elif len(candidates) > 1:
                        return _err('این پروژه چند تیم فعال دارد؛ انتخاب تیم مسئول تسک الزامی است')
                    else:
                        return _err('برای این پروژه هنوز جریان کاری تیمی تعریف نشده است')
                project_team = _one(cur, """SELECT pt.id,pt.team_id FROM ProjectTeams pt
                    WHERE pt.id=? AND pt.project_id=? AND pt.is_active=1""",
                                    project_team_id, project_id)
                if not project_team:
                    return _err('تیم انتخاب‌شده روی این پروژه فعال نیست')
                if not _team_allowed(cur,u,project_team['team_id']):
                    return _err('به تیم مسئول این تسک دسترسی ندارید', 403)
                selected_staff = _int(d.get('staff_id'))
                if selected_staff:
                    cur.execute("""SELECT 1 FROM TeamMembers WHERE team_id=? AND user_id=?
                        AND is_active=1""", project_team['team_id'], selected_staff)
                    if not cur.fetchone():
                        return _err('پشتیبان انتخاب‌شده عضو تیم مسئول پروژه نیست')
                progress_weight = _number(d.get('progress_weight'), _decimal.Decimal(1))
                if progress_weight <= 0 or progress_weight > 1000:
                    return _err('وزن پیشرفت باید بیشتر از صفر و حداکثر ۱۰۰۰ باشد')
                contract_id = _int(d.get('contract_id'))
                planned_id = _int(d.get('planned_statement_id'))
                if not user_has_permission(u, 'tasks.link_contract'):
                    if ident:
                        contract_id = _int(old.get('contract_id'))
                        planned_id = _int(old.get('planned_statement_id'))
                    elif contract_id or planned_id:
                        return _err('مجوز اتصال قرارداد به تسک برای این نقش فعال نیست', 403)
                if planned_id:
                    planned = _one(cur, """SELECT contract_id,project_team_id
                        FROM PlannedStatements WHERE id=? AND is_active=1""", planned_id)
                    if not planned: return _err('برنامه صورت‌وضعیت انتخاب‌شده معتبر نیست')
                    if contract_id and contract_id != _int(planned.get('contract_id')):
                        return _err('قرارداد تسک با قرارداد برنامه صورت‌وضعیت یکسان نیست')
                    if planned.get('project_team_id') and project_team_id != _int(planned.get('project_team_id')):
                        return _err('تیم تسک با تیم برنامه صورت‌وضعیت یکسان نیست')
                    contract_id = _int(planned.get('contract_id'))
                if contract_id:
                    contract = _one(cur, "SELECT id,project_id FROM Contracts WHERE id=? AND is_active=1", contract_id)
                    if not contract:
                        return _err('قرارداد انتخاب‌شده معتبر یا فعال نیست')
                    if _int(contract.get('project_id')) != project_id:
                        return _err('قرارداد انتخاب‌شده متعلق به پروژه این تسک نیست')
                vals = [title, d.get('type') or 'bug', d.get('status') or 'registered', _int(d.get('category_id')),
                        project_id, _int(d.get('contact_id')), _int(d.get('staff_id')), d.get('date_recv') or None,
                        d.get('date_delivery') or None, (d.get('description') or '').strip() or None,
                        (d.get('solution') or '').strip() or None, _int(d.get('priority'), 2),
                        contract_id, planned_id, project_team_id, progress_weight]
                if ident:
                    cur.execute("""UPDATE Tasks SET title=?,type=?,status=?,category_id=?,project_id=?,contact_id=?,staff_id=?,
                        date_recv=?,date_delivery=?,description=?,solution=?,priority=?,contract_id=?,planned_statement_id=?,
                        project_team_id=?,progress_weight=? WHERE id=?""", *(vals + [ident]))
                else:
                    cur.execute("""INSERT INTO Tasks(title,type,status,category_id,project_id,contact_id,staff_id,date_recv,
                        date_delivery,description,solution,priority,contract_id,planned_statement_id,project_team_id,
                        progress_weight,created_by)
                        OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", *(vals + [u['id']]))
                    ident = cur.fetchone()[0]
                new = _one(cur, "SELECT * FROM Tasks WHERE id=?", ident)
                _entity_audit(cur, 'task', ident, 'update' if old else 'create', old, new); c.commit()
            return _ok(id=ident)
        except Exception as exc:
            return _err(str(exc))

    @app.route('/api/task_delete', methods=['POST'])
    @require_auth
    def api_v7_task_delete():
        try:
            ident = _int((request.get_json() or {}).get('id')); u = g.user
            with db_lock:
                c = get_conn(); cur = c.cursor(); old = _one(cur, "SELECT * FROM Tasks WHERE id=?", ident)
                if not old: return _err('تسک یافت نشد')
                allowed = user_has_permission(u, 'tasks.delete')
                if not allowed and user_has_permission(u, 'tasks.self_manage'):
                    allowed = (old.get('created_by') == u['id'] and old.get('status') == 'registered'
                               and old.get('project_id') == u.get('project_id'))
                if not allowed and user_has_permission(u, 'tasks.delete_group'):
                    # R15 group lead: unfinished tasks of themselves and their
                    # sub-group. A done task keeps its record and its points.
                    allowed = (old.get('status') != 'done'
                               and _int(old.get('staff_id')) in lead_group_ids(cur, u['id']))
                if not allowed: return _err('اجازه حذف این تسک را ندارید', 403)
                if not _task_allowed(cur,u,ident): return _err('این تسک خارج از محدوده تیم شماست',403)
                cur.execute("DELETE FROM Tasks WHERE id=?", ident); _entity_audit(cur, 'task', ident, 'delete', old, None); c.commit()
            return _ok()
        except Exception as exc:
            return _err(str(exc))

    def _effective_end_expr():
        # Never let an erroneous historical addendum shorten the base term.
        # Imported legacy addenda are authoritative history even when an old
        # database still carries a pre-R6 workflow status.
        return """(SELECT MAX(d.effective_date) FROM
            (SELECT co.end_date AS effective_date UNION ALL
             SELECT e.end_date FROM ContractExtensions e WHERE e.contract_id=co.id
              AND e.is_active=1
              AND (e.internal_status='approved' OR e.legacy_extension_id IS NOT NULL)
              AND e.extension_type_id IN (1,2) AND e.end_date IS NOT NULL) d)"""

    def _contract_select(where=''):
        return """SELECT co.id,co.project_id,co.legacy_contract_id,co.legacy_company_id,co.title,co.employer_name,
            co.contract_number,co.base_price,co.notification_date,co.notification_date_fa,
            co.notification_letter_number,co.start_date,co.start_date_fa,co.end_date,co.end_date_fa,
            co.contract_type_id,ct.name AS contract_type_name,co.contract_status_id,cs.name AS contract_status_name,
            co.work_order_code,co.financial_code,co.register_in_sajat,co.register_date_sajat,co.register_date_sajat_fa,
            co.category_id,cc.name AS category_name,co.insurance_coefficient,co.guarantee_coefficient,
            co.commercial_code,co.scope_id,co.is_new,co.tax_info_crn,
            co.payer_rating,co.payer_note,co.payer_locked,co.is_active,
            CASE WHEN co.contract_status_id IN (3,4) THEN 1 ELSE 0 END AS is_closed,
            p.name AS project_name,p.project_type_id,pt.name AS project_type_name,
            ci.id AS city_id,ci.name AS city_name,
            STUFF((SELECT N'، '+tm.name FROM ContractProjectTeams cpt
                JOIN ProjectTeams ptx ON ptx.id=cpt.project_team_id
                JOIN Teams tm ON tm.id=ptx.team_id
                WHERE cpt.contract_id=co.id AND cpt.is_active=1
                ORDER BY tm.name FOR XML PATH(''),TYPE).value('.','NVARCHAR(MAX)'),1,2,N'') AS team_names,
            STUFF((SELECT ','+CONVERT(NVARCHAR(20),cpt.project_team_id)
                FROM ContractProjectTeams cpt
                WHERE cpt.contract_id=co.id AND cpt.is_active=1
                ORDER BY cpt.project_team_id FOR XML PATH(''),TYPE).value('.','NVARCHAR(MAX)'),1,1,'') AS project_team_ids,
            """ + _effective_end_expr() + """ AS effective_end_date,
            COALESCE((SELECT TOP 1 COALESCE(e.end_date_fa,CONVERT(NVARCHAR(10),e.end_date,23))
                FROM ContractExtensions e WHERE e.contract_id=co.id AND e.is_active=1
                AND (e.internal_status='approved' OR e.legacy_extension_id IS NOT NULL)
                AND e.extension_type_id IN (1,2) AND e.end_date IS NOT NULL
                ORDER BY e.end_date DESC,e.id DESC),co.end_date_fa,CONVERT(NVARCHAR(10),co.end_date,23)) AS effective_end_date_fa,
            ISNULL(co.base_price,0)+ISNULL((SELECT SUM(ISNULL(e.price_delta,0)) FROM ContractExtensions e
                WHERE e.contract_id=co.id AND e.is_active=1 AND e.internal_status='approved' AND e.extension_type_id IN (2,3)),0) AS effective_price,
            ISNULL((SELECT SUM(COALESCE(s.confirmed_without_vat,s.confirmed_price,s.requested_without_vat,s.requested_price,0))
                FROM ContractStatements s WHERE s.contract_id=co.id AND s.is_active=1 AND s.is_current=1
                AND (s.business_status='employer_approved' OR (s.employer_decision_at IS NOT NULL AND ISNULL(s.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(s.confirmed_without_vat,s.confirmed_price,0)>0))),0) AS approved_statement_total,
            (SELECT COUNT(*) FROM ContractExtensions e WHERE e.contract_id=co.id AND e.is_active=1) AS extension_count,
            (SELECT COUNT(*) FROM ContractStatements s WHERE s.contract_id=co.id AND s.is_active=1 AND s.is_current=1) AS statement_count,
            (SELECT MAX(s.statement_date) FROM ContractStatements s WHERE s.contract_id=co.id AND s.is_active=1 AND s.is_current=1) AS last_statement_date,
            (SELECT COUNT(*) FROM Tasks t WHERE t.contract_id=co.id AND t.status='done') AS ready_task_count,
            (SELECT COUNT(*) FROM Tasks t WHERE t.contract_id=co.id AND t.status<>'done') AS open_task_count
            FROM Contracts co LEFT JOIN Projects p ON p.id=co.project_id
            LEFT JOIN ProjectTypes pt ON pt.id=p.project_type_id LEFT JOIN Cities ci ON ci.id=p.city_id
            LEFT JOIN ContractTypes ct ON ct.id=co.contract_type_id LEFT JOIN ContractStatuses cs ON cs.id=co.contract_status_id
            LEFT JOIN ContractCategories cc ON cc.id=co.category_id """ + where

    def _has_contract_team_role(cur, contract_id, allowed_roles):
        if any(user_has_permission(g.user, key) for key in ('contracts.manage','contracts.approve','contracts.archive','extensions.manage','extensions.approve','extensions.archive','statements.manage','statements.internal_approve','statements.employer_status','statements.archive','financial_plan.manage')):
            return True
        marks = ','.join('?' for _ in allowed_roles)
        cur.execute("""SELECT 1 FROM ContractProjectTeams cpt
            JOIN ProjectTeams ptm ON ptm.id=cpt.project_team_id
            JOIN TeamMembers tm ON tm.team_id=ptm.team_id
            WHERE cpt.contract_id=? AND cpt.is_active=1 AND ptm.is_active=1
              AND tm.user_id=? AND tm.is_active=1
              AND tm.team_role IN (%s)""" % marks,
                    *([contract_id, g.user['id']] + list(allowed_roles)))
        return cur.fetchone() is not None

    def _has_project_team_role(cur, project_team_id, allowed_roles):
        if any(user_has_permission(g.user, key) for key in ('contracts.manage','contracts.approve','contracts.archive','extensions.manage','extensions.approve','extensions.archive','statements.manage','statements.internal_approve','statements.employer_status','statements.archive','financial_plan.manage')):
            return True
        marks = ','.join('?' for _ in allowed_roles)
        cur.execute("""SELECT 1 FROM ProjectTeams ptm JOIN TeamMembers tm
            ON tm.team_id=ptm.team_id
            WHERE ptm.id=? AND ptm.is_active=1 AND tm.user_id=?
              AND tm.is_active=1 AND tm.team_role IN (%s)""" % marks,
                    *([project_team_id, g.user['id']] + list(allowed_roles)))
        return cur.fetchone() is not None

    def _has_statement_team_role(cur, statement_id, allowed_roles):
        if any(user_has_permission(g.user, key) for key in ('contracts.manage','contracts.approve','contracts.archive','extensions.manage','extensions.approve','extensions.archive','statements.manage','statements.internal_approve','statements.employer_status','statements.archive','financial_plan.manage')):
            return True
        marks = ','.join('?' for _ in allowed_roles)
        cur.execute("""SELECT 1 FROM ContractStatementTeams cst
            JOIN ProjectTeams ptm ON ptm.id=cst.project_team_id
            JOIN TeamMembers tm ON tm.team_id=ptm.team_id
            WHERE cst.statement_id=? AND cst.is_active=1
              AND ptm.is_active=1 AND tm.user_id=? AND tm.is_active=1
              AND tm.team_role IN (%s)""" % marks,
                    *([statement_id, g.user['id']] + list(allowed_roles)))
        return cur.fetchone() is not None

    def _archived_view(data):
        """True when the caller asked for the archive and may see it.

        Archiving used to be a one-way door: a record set is_active=0 vanished
        from every list with no way back. The flag is silently ignored without
        the permission so a crafted request cannot reveal archived rows.
        """
        if not data.get('archived'):
            return False
        return bool(user_has_permission(g.user, 'archives.view'))

    def _restore_record(table, entity, ident, permission):
        if not user_has_permission(g.user, 'archives.restore'):
            return _err('دسترسی بازگرداندن رکورد بایگانی‌شده را ندارید', 403)
        if not user_has_permission(g.user, permission):
            return _err('دسترسی بایگانی این رکورد را ندارید', 403)
        if not ident:
            return _err('شناسه رکورد الزامی است')
        with db_lock:
            c = get_conn(); cur = c.cursor()
            row = _one(cur, "SELECT id,is_active FROM %s WHERE id=?" % table, ident)
            if not row:
                return _err('رکورد یافت نشد')
            if row.get('is_active'):
                return _err('این رکورد بایگانی نشده است')
            cur.execute("UPDATE %s SET is_active=1,updated_by=?,updated_at=GETDATE() WHERE id=?" % table,
                        g.user['id'], ident)
            if entity == 'contract':
                # Put back exactly what the archive cascade took, and nothing
                # that was archived separately before that.
                for child in ('ContractStatements', 'ContractExtensions'):
                    cur.execute("""UPDATE %s SET is_active=1,archived_with_contract=0,
                        updated_by=?,updated_at=GETDATE()
                        WHERE contract_id=? AND is_active=0 AND archived_with_contract=1""" % child,
                                g.user['id'], ident)
            _entity_audit(cur, entity, ident, 'restore', {'is_active': 0}, {'is_active': 1})
            c.commit()
        return _ok(id=ident)

    @app.route('/api/contract_restore', methods=['POST'])
    @require_auth
    def api_v7_contract_restore():
        try:
            return _restore_record('Contracts', 'contract',
                                   _int((request.get_json() or {}).get('id')),
                                   'contracts.archive')
        except Exception as exc:
            return _err(str(exc))

    @app.route('/api/extension_restore', methods=['POST'])
    @require_auth
    def api_v7_extension_restore():
        try:
            return _restore_record('ContractExtensions', 'extension',
                                   _int((request.get_json() or {}).get('id')),
                                   'extensions.archive')
        except Exception as exc:
            return _err(str(exc))

    @app.route('/api/statement_restore', methods=['POST'])
    @require_auth
    def api_v7_statement_restore():
        try:
            return _restore_record('ContractStatements', 'statement',
                                   _int((request.get_json() or {}).get('id')),
                                   'statements.archive')
        except Exception as exc:
            return _err(str(exc))

    @app.route('/api/contracts', methods=['POST'])
    @require_auth
    def api_v7_contracts():
        try:
            d = request.get_json() or {}; only_unassigned = bool(d.get('only_without_project'))
            archived = _archived_view(d)
            with db_lock:
                c = get_conn(); cur = c.cursor()
                red = int((_one(cur, "SELECT setting_value FROM AppSettings WHERE setting_key='contract_red_days'") or {'setting_value': 15})['setting_value'])
                yellow = int((_one(cur, "SELECT setting_value FROM AppSettings WHERE setting_key='contract_yellow_days'") or {'setting_value': 45})['setting_value'])
                clauses = ["co.is_active=0" if archived else "co.is_active=1"]; scope_params = []
                if only_unassigned:
                    clauses.append("co.project_id IS NULL")
                # R12: contracts follow the same team rule the statement list
                # already used. Only admin, finance and the reporter see every
                # contract; a manager sees the ones the finance specialist
                # linked to their team, and an explicit team selector can only
                # narrow that authorised set further.
                scope_clause, team_params = _team_scope_condition(
                    g.user, d, 'contract', 'co'
                )
                if scope_clause:
                    clauses.append(scope_clause)
                    scope_params.extend(team_params)
                sql = _contract_select(" WHERE " + " AND ".join(clauses))
                end_expr = _effective_end_expr()
                sql += """ ORDER BY
                    CASE WHEN co.contract_status_id IN (3,4) THEN 1 ELSE 0 END,
                    CASE WHEN co.contract_status_id IN (3,4) THEN 4 WHEN %s IS NULL THEN 3
                         WHEN DATEDIFF(day,GETDATE(),%s)<=? THEN 0
                         WHEN DATEDIFF(day,GETDATE(),%s)<=? THEN 1 ELSE 2 END,
                    CASE WHEN co.contract_status_id IN (3,4) THEN %s END DESC,
                    CASE WHEN co.contract_status_id NOT IN (3,4) OR co.contract_status_id IS NULL THEN %s END,
                    effective_price DESC""" % (end_expr,end_expr,end_expr,end_expr,end_expr)
                cur.execute(sql, *(scope_params + [red, yellow])); rows = rows_to_list(cur)
            today = _dt.date.today()
            for row in rows:
                end = _date(row.get('effective_end_date'))
                level, days = expiry_level(end, today=today, red_days=red, yellow_days=yellow)
                if row.get('is_closed'):
                    level = 'closed'
                row['days_to_end'] = days
                row['expiry_level'] = level
                row['remaining_price'] = contract_remaining(row.get('effective_price'),
                    [{'confirmed_without_vat': row.get('approved_statement_total')}])
                effective = _dec(row.get('effective_price'))
                approved = _dec(row.get('approved_statement_total'))
                row['financial_progress_percent'] = _pct(approved, effective) if effective > 0 else 0
            return _ok(rows=rows, red_days=red, yellow_days=yellow, archived=archived)
        except Exception as exc:
            return _err(str(exc), rows=[])

    @app.route('/api/contract_lookups', methods=['POST'])
    @require_auth
    def api_v7_contract_lookups():
        try:
            with db_lock:
                c = get_conn(); cur = c.cursor(); result = {}
                for key, table in [('types','ContractTypes'),('statuses','ContractStatuses'),('categories','ContractCategories'),('extension_types','ContractExtensionTypes'),('statement_types','ContractStatementTypes')]:
                    cur.execute("SELECT id,name FROM %s ORDER BY id" % table); result[key] = rows_to_list(cur)
            return _ok(**result)
        except Exception as exc: return _err(str(exc))

    @app.route('/api/contract_save', methods=['POST'])
    @require_auth
    def api_v7_contract_save():
        try:
            d = request.get_json() or {}; ident = _int(d.get('id')); title = (d.get('title') or '').strip()
            new_project_id = _int(d.get('project_id'))
            if not title: return _err('عنوان قرارداد الزامی است')
            base_price=_number(d.get('base_price'));start_date=_date(d.get('start_date'));end_date=_date(d.get('end_date'))
            if base_price is not None and base_price<0:return _err('مبلغ پایه قرارداد نمی‌تواند منفی باشد')
            if start_date and end_date and end_date<start_date:return _err('تاریخ پایان قرارداد قبل از تاریخ شروع است')
            with db_lock:
                c = get_conn(); cur = c.cursor(); old = _one(cur, "SELECT * FROM Contracts WHERE id=?", ident) if ident else None
                project_changed = bool(old) and _int(old.get('project_id')) != new_project_id
                project_history = {}
                relink_with_history = False
                if project_changed:
                    project_history = _one(cur, """SELECT
                        (SELECT COUNT(*) FROM ContractExtensions
                          WHERE contract_id=? AND is_active=1) AS extensions,
                        (SELECT COUNT(*) FROM ContractStatements
                          WHERE contract_id=? AND is_active=1) AS statements,
                        (SELECT COUNT(*) FROM PlannedStatements
                          WHERE contract_id=? AND is_active=1) AS plans,
                        (SELECT COUNT(*) FROM Tasks
                          WHERE contract_id=?) AS tasks""",
                                   ident, ident, ident, ident) or {}
                    relink_with_history = any(int(project_history.get(x) or 0)
                                              for x in ('extensions','statements','plans','tasks'))
                    if relink_with_history:
                        if not user_has_permission(g.user, 'contracts.relink_project'):
                            return _err('برای تغییر پروژه قرارداد دارای سابقه، دسترسی انتقال پروژه لازم است',403)
                        if not bool(d.get('confirm_project_relink')):
                            return _err('تغییر پروژه، سوابق قرارداد و تیم مقصد را منتقل می‌کند؛ عملیات را دوباره با تأیید انجام دهید')
                        if not new_project_id:
                            return _err('قرارداد دارای سابقه را نمی‌توان بدون پروژه مقصد ذخیره کرد')
                        if 'project_teams' not in d and 'project_team_ids' not in d:
                            return _err('برای انتقال سوابق، تیم مقصد قرارداد را مشخص کنید')
                vals = [new_project_id, title, (d.get('employer_name') or '').strip() or None,
                        (d.get('contract_number') or '').strip() or None, base_price,
                        _date(d.get('notification_date')), d.get('notification_date_fa') or None,
                        d.get('notification_letter_number') or None, start_date, d.get('start_date_fa') or None,
                        end_date, d.get('end_date_fa') or None, _int(d.get('contract_type_id')),
                        _int(d.get('contract_status_id')), d.get('work_order_code') or None,d.get('financial_code') or None,
                        bool(d.get('register_in_sajat')) if d.get('register_in_sajat') is not None else None,
                        _date(d.get('register_date_sajat')), d.get('register_date_sajat_fa') or None,
                        _int(d.get('category_id')), _number(d.get('insurance_coefficient')),
                        _number(d.get('guarantee_coefficient')), d.get('commercial_code') or None,
                        bool(d.get('is_new')) if d.get('is_new') is not None else None,
                        d.get('tax_info_crn') or None, g.user['id']]
                if ident:
                    if not old: return _err('قرارداد یافت نشد')
                    cur.execute("""UPDATE Contracts SET project_id=?,title=?,employer_name=?,contract_number=?,base_price=?,
                        notification_date=?,notification_date_fa=?,notification_letter_number=?,start_date=?,start_date_fa=?,
                        end_date=?,end_date_fa=?,contract_type_id=?,contract_status_id=?,work_order_code=?,financial_code=?,
                        register_in_sajat=?,register_date_sajat=?,register_date_sajat_fa=?,category_id=?,insurance_coefficient=?,guarantee_coefficient=?,
                        commercial_code=?,is_new=?,tax_info_crn=?,
                        updated_by=?,updated_at=GETDATE() WHERE id=?""", *(vals + [ident]))
                else:
                    cur.execute("""INSERT INTO Contracts(project_id,title,employer_name,contract_number,base_price,
                        notification_date,notification_date_fa,notification_letter_number,start_date,start_date_fa,end_date,end_date_fa,
                        contract_type_id,contract_status_id,work_order_code,financial_code,register_in_sajat,register_date_sajat,register_date_sajat_fa,
                        category_id,insurance_coefficient,guarantee_coefficient,commercial_code,is_new,tax_info_crn,created_by)
                        OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", vals)
                    ident = cur.fetchone()[0]
                # A contract may be shared by several team workstreams.  The
                # relationship lives in ContractProjectTeams, never as a
                # duplicate team id on the contract itself.
                if 'project_teams' in d or 'project_team_ids' in d or not old:
                    entries = d.get('project_teams')
                    if entries is None:
                        entries = [{'project_team_id': x} for x in (d.get('project_team_ids') or [])]
                    project_id = new_project_id
                    if not entries and project_id:
                        cur.execute("""SELECT id FROM ProjectTeams WHERE project_id=?
                            AND is_active=1 ORDER BY is_primary DESC,id""", project_id)
                        candidates = [int(x[0]) for x in cur.fetchall()]
                        if len(candidates) == 1:
                            entries = [{'project_team_id': candidates[0]}]
                        elif len(candidates) > 1:
                            c.rollback(); return _err('این پروژه چند تیم فعال دارد؛ تیم‌های قرارداد را انتخاب کنید')
                        else:
                            c.rollback(); return _err('برای پروژه قرارداد هنوز جریان کاری تیمی تعریف نشده است')
                    clean = []
                    seen_workstreams = set()
                    allocation_total = _decimal.Decimal(0)
                    for item in entries:
                        ptid = _int(item.get('project_team_id') if isinstance(item, dict) else item)
                        amount = _number(item.get('allocation_amount')) if isinstance(item, dict) else None
                        if not ptid or ptid in seen_workstreams:
                            c.rollback(); return _err('هر تیم قرارداد باید دقیقاً یک‌بار انتخاب شود')
                        seen_workstreams.add(ptid)
                        row = _one(cur, """SELECT pt.id,pt.team_id FROM ProjectTeams pt
                            WHERE pt.id=? AND pt.project_id=? AND pt.is_active=1""",
                                   ptid, project_id)
                        if not row:
                            c.rollback(); return _err('یکی از تیم‌های قرارداد روی پروژه فعال نیست')
                        if amount is not None:
                            if amount < 0:
                                c.rollback(); return _err('سهم ریالی تیم نمی‌تواند منفی باشد')
                            allocation_total += amount
                        clean.append((ptid, amount))
                    # Manager and planner edit the complete contract team list.
                    if base_price is not None and allocation_total > base_price:
                        c.rollback(); return _err('جمع سهم تیم‌ها از مبلغ پایه قرارداد بیشتر است')
                    cur.execute("""SELECT project_team_id FROM ContractProjectTeams
                        WHERE contract_id=? AND is_active=1""", ident)
                    current_workstreams = {int(x[0]) for x in cur.fetchall()}
                    removed_workstreams = current_workstreams - seen_workstreams
                    if removed_workstreams and not relink_with_history:
                        removed_marks = ','.join('?' for _ in removed_workstreams)
                        removed_params = list(removed_workstreams)
                        usage = _one(cur, """SELECT
                            (SELECT COUNT(*) FROM ContractStatementTeams cst
                              JOIN ContractStatements s ON s.id=cst.statement_id
                              WHERE s.contract_id=? AND s.is_active=1
                                AND cst.is_active=1
                                AND cst.project_team_id IN (%s)) AS statements,
                            (SELECT COUNT(*) FROM PlannedStatements ps
                              WHERE ps.contract_id=? AND ps.is_active=1
                                AND ps.project_team_id IN (%s)) AS plans,
                            (SELECT COUNT(*) FROM Tasks t
                              WHERE t.contract_id=?
                                AND t.project_team_id IN (%s)) AS tasks,
                            (SELECT COUNT(*) FROM ContractExtensionTeams cet
                              JOIN ContractExtensions e ON e.id=cet.extension_id
                              WHERE e.contract_id=? AND e.is_active=1
                                AND cet.project_team_id IN (%s)) AS extensions""" % (
                                    removed_marks, removed_marks,
                                    removed_marks, removed_marks),
                                     *([ident] + removed_params + [ident] +
                                       removed_params + [ident] + removed_params +
                                       [ident] + removed_params)) or {}
                        if any(int(usage.get(x) or 0) for x in
                               ('statements','plans','tasks','extensions')):
                            c.rollback()
                            return _err(
                                'اتصال تیمی دارای سابقه است و قابل حذف نیست؛ '
                                'ابتدا رکوردهای وابسته را اصلاح یا بایگانی کنید'
                            )
                    if relink_with_history:
                        if len(clean) != 1:
                            c.rollback()
                            return _err('برای انتقال خودکار سوابق قرارداد، دقیقاً یک تیم مقصد انتخاب کنید')
                        destination_workstream = clean[0][0]
                        # Move all dependent operational/team links atomically so
                        # no task or financial record remains attached to the old project.
                        cur.execute("""UPDATE Tasks SET project_id=?,project_team_id=?
                            WHERE contract_id=?""", new_project_id,
                                    destination_workstream, ident)
                        cur.execute("""UPDATE PlannedStatements SET project_team_id=?,
                            updated_by=?,updated_at=GETDATE()
                            WHERE contract_id=? AND is_active=1""",
                                    destination_workstream, g.user['id'], ident)
                        cur.execute("""UPDATE ContractStatements SET project_team_id=?,
                            updated_by=?,updated_at=GETDATE()
                            WHERE contract_id=? AND is_active=1""",
                                    destination_workstream, g.user['id'], ident)
                        cur.execute("""UPDATE cst SET is_active=0
                            FROM ContractStatementTeams cst
                            JOIN ContractStatements s ON s.id=cst.statement_id
                            WHERE s.contract_id=?""", ident)
                        cur.execute("""UPDATE cst SET allocation_percent=100,
                              is_active=1,linked_by=?,linked_at=GETDATE()
                            FROM ContractStatementTeams cst
                            JOIN ContractStatements s ON s.id=cst.statement_id
                            WHERE s.contract_id=? AND cst.project_team_id=?""",
                                    g.user['id'], ident, destination_workstream)
                        cur.execute("""INSERT INTO ContractStatementTeams(
                              statement_id,project_team_id,allocation_percent,
                              is_active,linked_by)
                            SELECT s.id,?,100,1,?
                            FROM ContractStatements s
                            WHERE s.contract_id=? AND s.is_active=1
                              AND NOT EXISTS(SELECT 1 FROM ContractStatementTeams x
                                WHERE x.statement_id=s.id AND x.project_team_id=?)""",
                                    destination_workstream, g.user['id'], ident,
                                    destination_workstream)
                        cur.execute("""DELETE cet FROM ContractExtensionTeams cet
                            JOIN ContractExtensions e ON e.id=cet.extension_id
                            WHERE e.contract_id=?""", ident)
                        cur.execute("""INSERT INTO ContractExtensionTeams(
                              extension_id,project_team_id)
                            SELECT e.id,? FROM ContractExtensions e
                            WHERE e.contract_id=? AND e.is_active=1""",
                                    destination_workstream, ident)

                    cur.execute("UPDATE ContractProjectTeams SET is_active=0 WHERE contract_id=?", ident)
                    for ptid, amount in clean:
                        cur.execute("""IF EXISTS(SELECT 1 FROM ContractProjectTeams
                              WHERE contract_id=? AND project_team_id=?)
                            UPDATE ContractProjectTeams SET allocation_amount=?,is_active=1,
                              linked_by=?,linked_at=GETDATE()
                              WHERE contract_id=? AND project_team_id=?
                            ELSE INSERT INTO ContractProjectTeams(contract_id,project_team_id,
                              allocation_amount,linked_by) VALUES(?,?,?,?)""",
                                    ident, ptid, amount, g.user['id'], ident, ptid,
                                    ident, ptid, amount, g.user['id'])
                    if len(clean) == 1:
                        legacy_team_id = clean[0][0]
                        cur.execute("""UPDATE s SET project_team_id=?
                            FROM ContractStatements s
                            WHERE s.contract_id=? AND s.is_active=1
                              AND NOT EXISTS(SELECT 1
                                FROM ContractStatementTeams cst
                                WHERE cst.statement_id=s.id
                                  AND cst.is_active=1)""",
                                    legacy_team_id, ident)
                        cur.execute("""UPDATE existing SET
                              allocation_percent=100,is_active=1,
                              linked_by=?,linked_at=GETDATE()
                            FROM ContractStatementTeams existing
                            JOIN ContractStatements s
                              ON s.id=existing.statement_id
                            WHERE s.contract_id=? AND s.is_active=1
                              AND existing.project_team_id=?
                              AND NOT EXISTS(SELECT 1
                                FROM ContractStatementTeams active_link
                                WHERE active_link.statement_id=s.id
                                  AND active_link.is_active=1)""",
                                    g.user['id'], ident, legacy_team_id)
                        cur.execute("""INSERT INTO ContractStatementTeams(
                              statement_id,project_team_id,
                              allocation_percent,is_active,linked_by)
                            SELECT s.id,?,100,1,?
                            FROM ContractStatements s
                            WHERE s.contract_id=? AND s.is_active=1
                              AND NOT EXISTS(SELECT 1
                                FROM ContractStatementTeams active_link
                                WHERE active_link.statement_id=s.id
                                  AND active_link.is_active=1)
                              AND NOT EXISTS(SELECT 1
                                FROM ContractStatementTeams same_link
                                WHERE same_link.statement_id=s.id
                                  AND same_link.project_team_id=?)""",
                                    legacy_team_id, g.user['id'], ident,
                                    legacy_team_id)
                new = _one(cur, "SELECT * FROM Contracts WHERE id=?", ident)
                if relink_with_history:
                    _entity_audit(cur, 'contract', ident, 'project_relink',
                                  {'project_id': old.get('project_id'),
                                   'history': project_history},
                                  {'project_id': new_project_id,
                                   'destination_project_team_id': clean[0][0]})
                _entity_audit(cur, 'contract', ident, 'update' if old else 'create', old, new); c.commit()
            return _ok(id=ident)
        except Exception as exc: return _err(str(exc))

    @app.route('/api/contract_rating', methods=['POST'])
    @require_auth
    def api_v7_contract_rating():
        try:
            d = request.get_json() or {}; ident = _int(d.get('id')); rating = _int(d.get('rating'))
            if rating is None or rating < 0 or rating > 5: return _err('امتیاز باید بین صفر تا پنج باشد')
            with db_lock:
                c = get_conn(); cur = c.cursor(); old = _one(cur, "SELECT payer_rating,payer_note,payer_locked FROM Contracts WHERE id=?", ident)
                if not old: return _err('قرارداد یافت نشد')
                if old.get('payer_locked') and not user_has_permission(g.user, 'contracts.approve'): return _err('امتیاز توسط مدیر نهایی شده است', 403)
                lock = bool(d.get('lock')) if user_has_permission(g.user, 'contracts.approve') else bool(old.get('payer_locked'))
                cur.execute("UPDATE Contracts SET payer_rating=?,payer_note=?,payer_locked=?,payer_rated_by=?,payer_rated_at=GETDATE(),updated_at=GETDATE() WHERE id=?",
                            rating, (d.get('note') or '').strip() or None, lock, g.user['id'], ident)
                new = _one(cur, "SELECT payer_rating,payer_note,payer_locked FROM Contracts WHERE id=?", ident)
                _entity_audit(cur, 'contract', ident, 'payer_rating', old, new); c.commit()
            return _ok()
        except Exception as exc: return _err(str(exc))

    @app.route('/api/contract_archive', methods=['POST'])
    @require_auth
    def api_v7_contract_archive():
        try:
            ident = _int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn(); cur=c.cursor(); old=_one(cur,"SELECT * FROM Contracts WHERE id=?",ident)
                if not old:return _err('قرارداد یافت نشد')
                cur.execute("UPDATE Contracts SET is_active=0,updated_by=?,updated_at=GETDATE() WHERE id=?",g.user['id'],ident)
                # Children follow the contract, otherwise their amounts keep
                # counting in every financial report while the contract itself
                # is archived. The flag marks what this cascade touched.
                cur.execute("""UPDATE ContractStatements SET is_active=0,archived_with_contract=1,
                    updated_by=?,updated_at=GETDATE() WHERE contract_id=? AND is_active=1""",
                            g.user['id'],ident)
                cur.execute("""UPDATE ContractExtensions SET is_active=0,archived_with_contract=1,
                    updated_by=?,updated_at=GETDATE() WHERE contract_id=? AND is_active=1""",
                            g.user['id'],ident)
                _entity_audit(cur,'contract',ident,'archive',old,None);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/extensions', methods=['POST'])
    @require_auth
    def api_v7_extensions():
        try:
            d = request.get_json() or {}; cid = _int(d.get('contract_id'))
            archived = _archived_view(d)
            with db_lock:
                c=get_conn();cur=c.cursor();params=[]
                where=" WHERE e.is_active=0" if archived else " WHERE e.is_active=1"
                if cid: where += " AND e.contract_id=?"; params.append(cid)
                # An extension belongs to whoever the contract belongs to, so
                # it reuses the contract rule against e.contract_id rather
                # than inventing a second one.
                scope_clause, scope_params = _team_scope_condition(
                    g.user, d, 'extension', 'e'
                )
                if scope_clause:
                    where += " AND " + scope_clause
                    params.extend(scope_params)
                cur.execute("""SELECT e.*,et.name AS extension_type_name,co.title AS contract_title,co.contract_number,
                    co.project_id,p.name AS project_name,ci.id AS city_id,ci.name AS city_name,
                    u.display_name AS approved_by_name,
                    STUFF((SELECT N'، '+tm.name FROM ContractProjectTeams cpt
                      JOIN ProjectTeams ptx ON ptx.id=cpt.project_team_id
                      JOIN Teams tm ON tm.id=ptx.team_id
                      WHERE cpt.contract_id=co.id AND cpt.is_active=1
                      ORDER BY tm.name FOR XML PATH(''),TYPE).value('.','NVARCHAR(MAX)'),1,2,N'') AS team_names
                    FROM ContractExtensions e JOIN Contracts co ON co.id=e.contract_id
                    LEFT JOIN ContractExtensionTypes et ON et.id=e.extension_type_id
                    LEFT JOIN Projects p ON p.id=co.project_id LEFT JOIN Cities ci ON ci.id=p.city_id
                    LEFT JOIN Users u ON u.id=e.approved_by"""+where+" ORDER BY ci.name,p.name,co.title,e.extension_date DESC,e.id DESC",params)
                rows=rows_to_list(cur)
            return _ok(rows=rows, archived=archived)
        except Exception as exc:return _err(str(exc),rows=[])

    @app.route('/api/extension_save', methods=['POST'])
    @require_auth
    def api_v7_extension_save():
        try:
            d=request.get_json() or {};ident=_int(d.get('id'));cid=_int(d.get('contract_id'));title=(d.get('title') or '').strip();typ=_int(d.get('extension_type_id'))
            if not cid or not title or typ not in (1,2,3):return _err('قرارداد، عنوان و نوع الحاقیه الزامی است')
            mode=d.get('price_mode') or 'delta'
            if mode not in ('delta','final_snapshot','none'):return _err('روش مبلغ نامعتبر است')
            price_delta=_number(d.get('price_delta')) if mode=='delta' else None
            resulting=_number(d.get('resulting_price')) if mode=='final_snapshot' else None
            start_date=_date(d.get('start_date'));end_date=_date(d.get('end_date'))
            if start_date and end_date and end_date<start_date:return _err('پایان الحاقیه قبل از شروع آن است')
            if resulting is not None and resulting<0:return _err('مبلغ نهایی قرارداد نمی‌تواند منفی باشد')
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractExtensions WHERE id=?",ident) if ident else None
                if not _one(cur,"SELECT id FROM Contracts WHERE id=? AND is_active=1",cid):return _err('قرارداد فعال یافت نشد')
                if old and old.get('internal_status') not in ('draft','rejected') and not user_has_permission(g.user, 'extensions.edit_any_status'):
                    return _err('برای ویرایش الحاقیه‌ای که وارد گردش شده، دسترسی ویرایش همه وضعیت‌ها لازم است',403)
                vals=[cid,d.get('extension_number') or None,d.get('letter_number') or None,_date(d.get('extension_date')),d.get('extension_date_fa') or None,title,
                      start_date,d.get('start_date_fa') or None,end_date,d.get('end_date_fa') or None,typ,mode,price_delta,resulting,g.user['id']]
                if ident:
                    cur.execute("""UPDATE ContractExtensions SET contract_id=?,extension_number=?,letter_number=?,extension_date=?,extension_date_fa=?,title=?,
                        start_date=?,start_date_fa=?,end_date=?,end_date_fa=?,extension_type_id=?,price_mode=?,price_delta=?,resulting_price=?,
                        internal_status='draft',approved_by=NULL,approved_at=NULL,updated_by=?,updated_at=GETDATE() WHERE id=?""",*(vals+[ident]))
                else:
                    cur.execute("""INSERT INTO ContractExtensions(contract_id,extension_number,letter_number,extension_date,extension_date_fa,title,start_date,start_date_fa,
                        end_date,end_date_fa,extension_type_id,price_mode,price_delta,resulting_price,created_by) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",vals)
                    ident=cur.fetchone()[0]
                new=_one(cur,"SELECT * FROM ContractExtensions WHERE id=?",ident);_entity_audit(cur,'extension',ident,'update' if old else 'create',old,new);c.commit()
            return _ok(id=ident)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/extension_submit', methods=['POST'])
    @require_auth
    def api_v7_extension_submit():
        try:
            ident=_int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractExtensions WHERE id=? AND is_active=1",ident)
                if not old:return _err('الحاقیه یافت نشد')
                if old.get('internal_status') not in ('draft','rejected'):
                    return _err('فقط پیش‌نویس یا الحاقیه ردشده قابل ارسال برای مدیر است')
                cur.execute("UPDATE ContractExtensions SET internal_status='pending',updated_by=?,updated_at=GETDATE() WHERE id=?",g.user['id'],ident)
                _entity_audit(cur,'extension',ident,'submit',old,{'internal_status':'pending'});c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/extension_decide', methods=['POST'])
    @require_auth
    def api_v7_extension_decide():
        try:
            d=request.get_json() or {};ident=_int(d.get('id'));decision=d.get('decision')
            if decision not in ('approved','rejected'):return _err('تصمیم نامعتبر است')
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractExtensions WHERE id=? AND is_active=1",ident)
                if not old:return _err('الحاقیه یافت نشد')
                if old.get('internal_status')!='pending':return _err('ابتدا الحاقیه باید برای تایید مدیر ارسال شود')
                if decision=='approved' and old.get('extension_type_id') in (1,2) and not old.get('end_date'):
                    return _err('الحاقیه زمانی بدون تاریخ پایان جدید قابل تایید نیست')
                if decision=='approved' and old.get('extension_type_id') in (2,3) and old.get('price_mode')=='none':
                    return _err('ابتدا مشخص کنید مبلغ قدیمی «تغییر مبلغ» است یا «مبلغ نهایی»')
                if decision=='approved' and old.get('extension_type_id') in (2,3):
                    if old.get('price_mode')=='delta' and old.get('price_delta') is None:return _err('مبلغ تغییر الحاقیه مشخص نشده است')
                    if old.get('price_mode')=='final_snapshot' and old.get('resulting_price') is None:return _err('مبلغ نهایی الحاقیه مشخص نشده است')
                if decision=='approved' and old.get('extension_type_id') in (1,2):
                    prior_end=_one(cur,"""SELECT MAX(d.end_date) AS prior_end FROM
                        (SELECT co.end_date FROM Contracts co WHERE co.id=? UNION ALL
                         SELECT e.end_date FROM ContractExtensions e WHERE e.contract_id=? AND e.id<>? AND e.is_active=1
                         AND e.internal_status='approved' AND e.extension_type_id IN (1,2)) d""",old['contract_id'],old['contract_id'],ident) or {}
                    if prior_end.get('prior_end') and _date(old.get('end_date'))<=_date(prior_end.get('prior_end')):
                        return _err('تاریخ پایان جدید باید بعد از پایان مؤثر فعلی قرارداد باشد')
                # A final-price snapshot is converted to a delta at approval,
                # preventing double counting while preserving the source value.
                delta=old.get('price_delta')
                if decision=='approved' and old.get('price_mode')=='final_snapshot' and old.get('resulting_price') is not None:
                    base=_one(cur,"""SELECT ISNULL(co.base_price,0)+ISNULL((SELECT SUM(ISNULL(e.price_delta,0)) FROM ContractExtensions e
                        WHERE e.contract_id=co.id AND e.id<>? AND e.is_active=1 AND e.internal_status='approved' AND e.extension_type_id IN (2,3)),0) AS prior_price
                        FROM Contracts co WHERE co.id=?""",ident,old['contract_id'])
                    delta=_decimal.Decimal(str(old['resulting_price']))-_decimal.Decimal(str((base or {}).get('prior_price') or 0))
                if decision=='approved' and old.get('extension_type_id') in (2,3):
                    financial=_one(cur,"""SELECT
                        ISNULL(co.base_price,0)+ISNULL((SELECT SUM(ISNULL(e.price_delta,0)) FROM ContractExtensions e
                        WHERE e.contract_id=co.id AND e.id<>? AND e.is_active=1 AND e.internal_status='approved' AND e.extension_type_id IN (2,3)),0) AS prior_price,
                        ISNULL((SELECT SUM(COALESCE(s.confirmed_without_vat,s.confirmed_price,s.requested_without_vat,s.requested_price,0))
                        FROM ContractStatements s WHERE s.contract_id=co.id AND s.is_active=1 AND s.is_current=1 AND (s.business_status='employer_approved' OR (s.employer_decision_at IS NOT NULL AND ISNULL(s.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(s.confirmed_without_vat,s.confirmed_price,0)>0))),0) AS approved_total
                        FROM Contracts co WHERE co.id=?""",ident,old['contract_id']) or {}
                    next_price=_decimal.Decimal(str(financial.get('prior_price') or 0))+_decimal.Decimal(str(delta or 0))
                    if next_price<0:return _err('اثر الحاقیه مبلغ مؤثر قرارداد را منفی می‌کند')
                    if next_price<_decimal.Decimal(str(financial.get('approved_total') or 0)):
                        return _err('مبلغ قرارداد پس از الحاقیه از صورت‌وضعیت‌های تاییدشده کمتر می‌شود')
                cur.execute("""UPDATE ContractExtensions SET internal_status=?,review_note=?,approved_by=?,approved_at=GETDATE(),
                    price_delta=COALESCE(?,price_delta),updated_by=?,updated_at=GETDATE() WHERE id=?""",
                    decision,(d.get('note') or '').strip() or None,g.user['id'],delta,g.user['id'],ident)
                new=_one(cur,"SELECT * FROM ContractExtensions WHERE id=?",ident);_entity_audit(cur,'extension',ident,'manager_'+decision,old,new);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/extension_archive', methods=['POST'])
    @require_auth
    def api_v7_extension_archive():
        return _archive_entity('ContractExtensions','extension')

    @app.route('/api/statements', methods=['POST'])
    @require_auth
    def api_v7_statements():
        try:
            d=request.get_json() or {};cid=_int(d.get('contract_id'));params=[]
            archived=_archived_view(d)
            where=" WHERE s.is_active=0" if archived else " WHERE s.is_active=1"
            if cid:where+=" AND s.contract_id=?";params.append(cid)
            scope_clause,scope_params=_team_scope_condition(
                g.user,d,'statement','s'
            )
            if scope_clause:
                where+=" AND "+scope_clause;params.extend(scope_params)
            with db_lock:
                c=get_conn();cur=c.cursor();cur.execute("""SELECT s.*,st.name AS statement_type_name,co.title AS contract_title,co.contract_number,
                    co.project_id,p.name AS project_name,ci.id AS city_id,ci.name AS city_name,
                    ia.display_name AS internal_approved_by_name,ed.display_name AS employer_decision_by_name,
                    ptm.team_id,team.name AS team_name,
                    STUFF((SELECT N'، '+team_names.name
                      FROM ContractStatementTeams cst_names
                      JOIN ProjectTeams pt_names
                        ON pt_names.id=cst_names.project_team_id
                      JOIN Teams team_names ON team_names.id=pt_names.team_id
                      WHERE cst_names.statement_id=s.id
                        AND cst_names.is_active=1
                      ORDER BY team_names.name FOR XML PATH(''),TYPE
                    ).value('.','NVARCHAR(MAX)'),1,2,N'') AS team_names,
                    COALESCE(s.confirmed_without_vat,s.confirmed_price,s.requested_without_vat,s.requested_price,0) AS effective_amount
                    FROM ContractStatements s JOIN Contracts co ON co.id=s.contract_id
                    LEFT JOIN ContractStatementTypes st ON st.id=s.statement_type_id
                    LEFT JOIN Projects p ON p.id=co.project_id LEFT JOIN Cities ci ON ci.id=p.city_id
                    LEFT JOIN Users ia ON ia.id=s.internal_approved_by LEFT JOIN Users ed ON ed.id=s.employer_decision_by
                    LEFT JOIN ProjectTeams ptm ON ptm.id=s.project_team_id
                    LEFT JOIN Teams team ON team.id=ptm.team_id"""+where+
                    " ORDER BY ci.name,p.name,co.title,s.statement_date DESC,s.id DESC",params);rows=rows_to_list(cur)
                statement_ids=[int(x['id']) for x in rows]
                allocations={}
                if statement_ids:
                    marks=','.join('?' for _ in statement_ids)
                    cur.execute("""SELECT cst.statement_id,cst.project_team_id,
                        cst.allocation_percent,ptm.team_id,t.name AS team_name
                        FROM ContractStatementTeams cst
                        JOIN ProjectTeams ptm ON ptm.id=cst.project_team_id
                        JOIN Teams t ON t.id=ptm.team_id
                        WHERE cst.is_active=1 AND cst.statement_id IN (%s)
                        ORDER BY cst.statement_id,t.name"""%marks,statement_ids)
                    for item in rows_to_list(cur):
                        allocations.setdefault(int(item['statement_id']),[]).append(item)
                for item in rows:
                    item['team_allocations']=allocations.get(int(item['id']),[])
            return _ok(rows=rows, archived=archived)
        except Exception as exc:return _err(str(exc),rows=[])

    @app.route('/api/statement_save', methods=['POST'])
    @require_auth
    def api_v7_statement_save():
        try:
            d=request.get_json() or {};ident=_int(d.get('id'));cid=_int(d.get('contract_id'));title=(d.get('title') or '').strip();num=(d.get('statement_number') or '').strip();typ=_int(d.get('statement_type_id'))
            if not cid or not title or not num or typ not in (1,2,3,4,5):return _err('قرارداد، نوع، شماره و عنوان صورت‌وضعیت الزامی است')
            month=_int(d.get('period_month'))
            if month is not None and not 1<=month<=12:return _err('ماه نامعتبر است')
            year=_int(d.get('period_year'))
            if year is not None and not 1300<=year<=1600:return _err('سال مالی نامعتبر است')
            progress=_int(d.get('progress_percentage'))
            if progress is not None and not 0<=progress<=100:return _err('درصد پیشرفت باید بین صفر تا صد باشد')
            start_date=_date(d.get('start_date'));end_date=_date(d.get('end_date'))
            if start_date and end_date and end_date<start_date:return _err('پایان دوره صورت‌وضعیت قبل از شروع آن است')
            requested_price=_number(d.get('requested_price'));requested_without=_number(d.get('requested_without_vat'));requested_vat=_number(d.get('requested_vat'))
            vat_percent=_number(d.get('vat_percentage_value'))
            if vat_percent is not None and not 0<=vat_percent<=100:return _err('درصد ارزش افزوده باید بین صفر تا صد باشد')
            # Employer-confirmed amounts enter only through
            # statement_employer_decide after an internally approved send.
            confirmed_price=confirmed_without=confirmed_vat=None
            if any(x is not None and x<0 for x in (requested_price,requested_without,requested_vat)):
                return _err('مبالغ صورت‌وضعیت نمی‌توانند منفی باشند')
            without_vat=bool(d.get('without_vat')) if d.get('without_vat') is not None else None
            if without_vat:
                requested_vat=_decimal.Decimal(0)
                if requested_without is not None:requested_price=requested_without
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractStatements WHERE id=?",ident) if ident else None
                contract=_one(cur,"SELECT id,project_id FROM Contracts WHERE id=? AND is_active=1",cid)
                if not contract:return _err('قرارداد فعال یافت نشد')
                raw_allocations=d.get('statement_teams')
                if raw_allocations is None:
                    fallback=_int(d.get('project_team_id'))
                    raw_allocations=(
                        [{'project_team_id':fallback,'allocation_percent':100}]
                        if fallback else []
                    )
                if not raw_allocations:
                    cur.execute("""SELECT project_team_id FROM ContractProjectTeams
                        WHERE contract_id=? AND is_active=1 ORDER BY project_team_id""",cid)
                    candidates=[int(x[0]) for x in cur.fetchall()]
                    if len(candidates)==1:
                        raw_allocations=[{
                            'project_team_id':candidates[0],
                            'allocation_percent':100,
                        }]
                    elif len(candidates)>1:return _err('این قرارداد بین چند تیم مشترک است؛ تیم صورت‌وضعیت الزامی است')
                    else:return _err('برای قرارداد هنوز تیم فعال تعریف نشده است')
                clean_allocations=[];seen_workstreams=set()
                for allocation in raw_allocations:
                    if not isinstance(allocation,dict):
                        return _err('تقسیم تیمی صورت‌وضعیت معتبر نیست')
                    project_team_id=_int(allocation.get('project_team_id'))
                    share=_number(allocation.get('allocation_percent'))
                    if not project_team_id or project_team_id in seen_workstreams:
                        return _err('هر تیم صورت‌وضعیت باید دقیقاً یک‌بار انتخاب شود')
                    seen_workstreams.add(project_team_id)
                    valid_team=_one(cur,"""SELECT pt.id,pt.team_id
                        FROM ContractProjectTeams cpt
                        JOIN ProjectTeams pt ON pt.id=cpt.project_team_id
                        WHERE cpt.contract_id=? AND cpt.project_team_id=?
                          AND cpt.is_active=1 AND pt.is_active=1""",
                                    cid,project_team_id)
                    if not valid_team:
                        return _err('یکی از تیم‌های صورت‌وضعیت در قرارداد فعال نیست')
                    clean_allocations.append([project_team_id,share])
                if len(clean_allocations)==1:
                    clean_allocations[0][1]=_decimal.Decimal(100)
                else:
                    if any(x[1] is None or x[1]<=0 or x[1]>100
                           for x in clean_allocations):
                        return _err('برای صورت‌وضعیت مشترک سهم درصدی مثبت هر تیم الزامی است')
                    share_total=sum((x[1] for x in clean_allocations),
                                    _decimal.Decimal(0))
                    if abs(share_total-_decimal.Decimal(100))>_decimal.Decimal('0.01'):
                        return _err('جمع سهم درصدی تیم‌های صورت‌وضعیت باید دقیقاً ۱۰۰ باشد')
                project_team_id=clean_allocations[0][0]
                cur.execute("SELECT 1 FROM ContractStatements WHERE contract_id=? AND statement_number=? AND is_active=1 AND is_current=1 AND id<>ISNULL(?,0)",cid,num,ident)
                if cur.fetchone():return _err('صورت‌وضعیت فعال دیگری با همین شماره در این قرارداد وجود دارد')
                old_status = (old or {}).get('business_status')
                if old:
                    # A statement past the employer decision used to be locked
                    # for everyone, which left no way to correct a wrong
                    # approved figure. It is now gated by the permission that
                    # already existed for exactly this, and the edit sends the
                    # document back to draft (below) so the approval cycle is
                    # repeated rather than silently inherited.
                    if old_status in ('employer_approved','employer_rejected','revised'):
                        if not user_has_permission(g.user, 'statements.edit_any_status'):
                            return _err('صورت‌وضعیت پس از تأییدیه کارفرما قفل است؛ ویرایش آن دسترسی «ویرایش در همه وضعیت‌ها» می‌خواهد',403)
                    elif old_status not in ('draft','internal_rejected') and not (
                            user_has_permission(g.user, 'statements.edit_before_employer_decision')
                            or user_has_permission(g.user, 'statements.edit_any_status')):
                        return _err('دسترسی ویرایش صورت‌وضعیت پیش از تأییدیه کارفرما را ندارید',403)
                vals=[cid,typ,num,_date(d.get('statement_date')),d.get('statement_date_fa') or None,d.get('letter_number') or None,title,
                      start_date,d.get('start_date_fa') or None,end_date,d.get('end_date_fa') or None,
                      requested_price,requested_without,requested_vat,d.get('vat_factor_number') or None,
                      confirmed_price,confirmed_without,confirmed_vat,
                      progress,year,month,vat_percent,without_vat,
                      (d.get('description') or '').strip() or None,project_team_id,g.user['id']]
                if ident:
                    preserve_status = old.get('business_status') if old and old.get('business_status') in (
                        'draft','internal_rejected','pending_internal','internal_approved','sent'
                    ) else 'draft'
                    cur.execute("""UPDATE ContractStatements SET contract_id=?,statement_type_id=?,statement_number=?,statement_date=?,statement_date_fa=?,letter_number=?,title=?,
                        start_date=?,start_date_fa=?,end_date=?,end_date_fa=?,requested_price=?,requested_without_vat=?,requested_vat=?,vat_factor_number=?,confirmed_price=?,confirmed_without_vat=?,confirmed_vat=?,
                        progress_percentage=?,period_year=?,period_month=?,vat_percentage_value=?,without_vat=?,description=?,business_status=?,employer_decision_by=NULL,
                        employer_decision_at=NULL,employer_decision_note=NULL,project_team_id=?,
                        updated_by=?,updated_at=GETDATE() WHERE id=?""",*(vals[:-2]+[preserve_status,project_team_id,g.user['id'],ident]))
                    if preserve_status == 'draft' and old_status != 'draft':
                        # Returning to draft must also drop the internal
                        # approval, otherwise the document shows a signature
                        # for figures that have since been changed. Statuses
                        # that keep their place in the flow are left alone.
                        cur.execute("""UPDATE ContractStatements SET internal_approved_by=NULL,
                            internal_approved_at=NULL WHERE id=?""", ident)
                else:
                    cur.execute("""INSERT INTO ContractStatements(contract_id,statement_type_id,statement_number,statement_date,statement_date_fa,letter_number,title,
                        start_date,start_date_fa,end_date,end_date_fa,requested_price,requested_without_vat,requested_vat,vat_factor_number,confirmed_price,confirmed_without_vat,confirmed_vat,
                        progress_percentage,period_year,period_month,vat_percentage_value,without_vat,description,project_team_id,created_by) OUTPUT INSERTED.id
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",vals);ident=cur.fetchone()[0]
                cur.execute("""UPDATE ContractStatementTeams SET is_active=0
                    WHERE statement_id=?""",ident)
                for team_id,share in clean_allocations:
                    cur.execute("""IF EXISTS(SELECT 1 FROM ContractStatementTeams
                          WHERE statement_id=? AND project_team_id=?)
                        UPDATE ContractStatementTeams SET allocation_percent=?,
                          is_active=1,linked_by=?,linked_at=GETDATE()
                          WHERE statement_id=? AND project_team_id=?
                        ELSE INSERT INTO ContractStatementTeams(statement_id,
                          project_team_id,allocation_percent,linked_by)
                          VALUES(?,?,?,?)""",
                                ident,team_id,share,g.user['id'],ident,team_id,
                                ident,team_id,share,g.user['id'])
                new=_one(cur,"SELECT * FROM ContractStatements WHERE id=?",ident);_entity_audit(cur,'statement',ident,'update' if old else 'create',old,new);c.commit()
            return _ok(id=ident)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/statement_teams_save', methods=['POST'])
    @require_auth
    def api_v8_statement_teams_save():
        """Repair or reallocate team ownership without changing financial data."""
        try:
            d=request.get_json() or {};ident=_int(d.get('id'))
            raw=d.get('statement_teams') or []
            if not ident or not raw:
                return _err('صورت‌وضعیت و حداقل یک تیم الزامی است')
            with db_lock:
                c=get_conn();cur=c.cursor()
                statement=_one(cur,"""SELECT id,contract_id,business_status
                    FROM ContractStatements
                    WHERE id=? AND is_active=1""",ident)
                if not statement:return _err('صورت‌وضعیت یافت نشد')
                if statement.get('business_status') in ('employer_approved','employer_rejected','revised'):
                    return _err('پس از ثبت تأییدیه کارفرما، تقسیم تیمی صورت‌وضعیت قابل تغییر نیست',409)
                old_rows=[]
                cur.execute("""SELECT project_team_id,allocation_percent
                    FROM ContractStatementTeams
                    WHERE statement_id=? AND is_active=1
                    ORDER BY project_team_id""",ident)
                old_rows=rows_to_list(cur)
                clean=[];seen=set()
                for item in raw:
                    if not isinstance(item,dict):
                        return _err('تقسیم تیمی صورت‌وضعیت معتبر نیست')
                    ptid=_int(item.get('project_team_id'))
                    share=_number(item.get('allocation_percent'))
                    if not ptid or ptid in seen:
                        return _err('هر تیم صورت‌وضعیت باید دقیقاً یک‌بار انتخاب شود')
                    seen.add(ptid)
                    if not _one(cur,"""SELECT 1 AS ok
                        FROM ContractProjectTeams cpt
                        JOIN ProjectTeams ptm
                          ON ptm.id=cpt.project_team_id
                        WHERE cpt.contract_id=?
                          AND cpt.project_team_id=?
                          AND cpt.is_active=1 AND ptm.is_active=1""",
                                statement['contract_id'],ptid):
                        return _err('یکی از تیم‌ها در قرارداد فعال نیست')
                    clean.append([ptid,share])
                if len(clean)==1:
                    clean[0][1]=_decimal.Decimal(100)
                else:
                    if any(x[1] is None or x[1]<=0 or x[1]>100
                           for x in clean):
                        return _err('سهم مثبت هر تیم برای صورت‌وضعیت مشترک الزامی است')
                    total=sum((x[1] for x in clean),_decimal.Decimal(0))
                    if abs(total-_decimal.Decimal(100))>_decimal.Decimal('0.01'):
                        return _err('جمع سهم درصدی تیم‌های صورت‌وضعیت باید دقیقاً ۱۰۰ باشد')
                cur.execute("""UPDATE ContractStatementTeams SET is_active=0
                    WHERE statement_id=?""",ident)
                for ptid,share in clean:
                    cur.execute("""IF EXISTS(SELECT 1
                          FROM ContractStatementTeams
                          WHERE statement_id=? AND project_team_id=?)
                        UPDATE ContractStatementTeams SET
                          allocation_percent=?,is_active=1,
                          linked_by=?,linked_at=GETDATE()
                          WHERE statement_id=? AND project_team_id=?
                        ELSE INSERT INTO ContractStatementTeams(
                          statement_id,project_team_id,
                          allocation_percent,linked_by)
                          VALUES(?,?,?,?)""",
                                ident,ptid,share,g.user['id'],ident,ptid,
                                ident,ptid,share,g.user['id'])
                cur.execute("""UPDATE ContractStatements
                    SET project_team_id=?,updated_by=?,updated_at=GETDATE()
                    WHERE id=?""",clean[0][0],g.user['id'],ident)
                _entity_audit(cur,'statement',ident,'team_allocation',
                              {'teams':old_rows},
                              {'teams':[{'project_team_id':x[0],
                                         'allocation_percent':x[1]}
                                        for x in clean]})
                c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/statement_submit', methods=['POST'])
    @require_auth
    def api_v7_statement_submit():
        try:
            ident=_int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractStatements WHERE id=? AND is_active=1",ident)
                if not old:return _err('صورت‌وضعیت یافت نشد')
                if old.get('business_status') not in ('draft','internal_rejected'):
                    return _err('فقط پیش‌نویس یا صورت‌وضعیت ردشده قابل ارسال برای مدیر است')
                cur.execute("UPDATE ContractStatements SET business_status='pending_internal',updated_by=?,updated_at=GETDATE() WHERE id=?",g.user['id'],ident)
                _entity_audit(cur,'statement',ident,'submit',old,{'business_status':'pending_internal'});c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/statement_internal_decide', methods=['POST'])
    @require_auth
    def api_v7_statement_internal_decide():
        try:
            d=request.get_json() or {};ident=_int(d.get('id'));decision=d.get('decision')
            if decision not in ('approved','rejected'):return _err('تصمیم نامعتبر است')
            target='internal_approved' if decision=='approved' else 'internal_rejected'
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractStatements WHERE id=? AND is_active=1",ident)
                if not old:return _err('صورت‌وضعیت یافت نشد')
                if old.get('business_status')!='pending_internal':return _err('ابتدا صورت‌وضعیت باید برای تایید مدیر ارسال شود')
                if decision=='approved' and not any(old.get(x) is not None and _decimal.Decimal(str(old.get(x)))>0 for x in ('requested_without_vat','requested_price')):
                    return _err('صورت‌وضعیت بدون مبلغ مثبت قابل تایید داخلی نیست')
                cur.execute("UPDATE ContractStatements SET business_status=?,internal_approved_by=?,internal_approved_at=GETDATE(),employer_decision_note=?,updated_by=?,updated_at=GETDATE() WHERE id=?",
                            target,g.user['id'],(d.get('note') or '').strip() or None,g.user['id'],ident)
                new=_one(cur,"SELECT * FROM ContractStatements WHERE id=?",ident);_entity_audit(cur,'statement',ident,'manager_'+decision,old,new);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/statement_mark_sent', methods=['POST'])
    @require_auth
    def api_v7_statement_mark_sent():
        try:
            ident=_int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractStatements WHERE id=? AND is_active=1",ident)
                if not old:return _err('صورت‌وضعیت یافت نشد')
                if old.get('business_status')!='internal_approved':return _err('ابتدا تایید داخلی مدیر لازم است')
                cur.execute("UPDATE ContractStatements SET business_status='sent',sent_at=GETDATE(),updated_by=?,updated_at=GETDATE() WHERE id=?",g.user['id'],ident)
                _entity_audit(cur,'statement',ident,'sent',old,{'business_status':'sent'});c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/statement_employer_decide', methods=['POST'])
    @require_auth
    def api_v7_statement_employer_decide():
        try:
            d=request.get_json() or {};ident=_int(d.get('id'));decision=d.get('decision')
            if decision not in ('approved','rejected'):return _err('تصمیم نامعتبر است')
            target='employer_approved' if decision=='approved' else 'employer_rejected'
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractStatements WHERE id=? AND is_active=1",ident)
                if not old:return _err('صورت‌وضعیت یافت نشد')
                if old.get('business_status')!='sent':return _err('فقط صورت‌وضعیت ارسال‌شده و تعیین‌تکلیف‌نشده قابل ثبت نتیجه کارفرماست')
                requested_vat = _dec(old.get('requested_vat'))
                requested_total = old.get('requested_price')
                requested_base = old.get('requested_without_vat')
                if requested_base is None and requested_total is not None:
                    requested_base = _dec(requested_total) - requested_vat
                requested_base = _dec(requested_base)
                if decision=='approved' and requested_base <= 0:
                    return _err('صورت‌وضعیت بدون مبلغ مثبت قابل تأیید کارفرما نیست')
                if decision=='approved':
                    confirmed_without_vat = requested_base
                    confirmed_vat = requested_vat
                    confirmed_price = (_dec(requested_total)
                                       if requested_total is not None
                                       else requested_base + requested_vat)
                    amount = confirmed_without_vat
                    totals=_one(cur,"""SELECT
                        ISNULL(co.base_price,0)+ISNULL((SELECT SUM(ISNULL(e.price_delta,0)) FROM ContractExtensions e WHERE e.contract_id=co.id AND e.is_active=1 AND e.internal_status='approved' AND e.extension_type_id IN (2,3)),0) AS effective_price,
                        ISNULL((SELECT SUM(COALESCE(s.confirmed_without_vat,s.confirmed_price,s.requested_without_vat,s.requested_price,0)) FROM ContractStatements s WHERE s.contract_id=co.id AND s.id<>? AND s.is_active=1 AND s.is_current=1 AND (s.business_status='employer_approved' OR (s.employer_decision_at IS NOT NULL AND ISNULL(s.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(s.confirmed_without_vat,s.confirmed_price,0)>0))),0) AS prior_approved
                        FROM Contracts co WHERE co.id=?""",ident,old['contract_id']) or {}
                    if _decimal.Decimal(str(totals.get('prior_approved') or 0))+amount>_decimal.Decimal(str(totals.get('effective_price') or 0)):
                        return _err('مبلغ تاییدشده باعث منفی‌شدن مانده قرارداد می‌شود؛ ابتدا مبلغ قرارداد/الحاقیه‌ها را اصلاح کنید')
                else:
                    confirmed_without_vat = None
                    confirmed_vat = None
                    confirmed_price = None
                cur.execute("""UPDATE ContractStatements SET business_status=?,confirmed_without_vat=?,confirmed_vat=?,confirmed_price=?,
                    employer_decision_by=?,employer_decision_at=GETDATE(),employer_decision_note=?,updated_by=?,updated_at=GETDATE() WHERE id=?""",
                    target,confirmed_without_vat,confirmed_vat,confirmed_price,g.user['id'],(d.get('note') or '').strip() or None,g.user['id'],ident)
                new=_one(cur,"SELECT * FROM ContractStatements WHERE id=?",ident);_entity_audit(cur,'statement',ident,'employer_'+decision,old,new);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/statement_revision', methods=['POST'])
    @require_auth
    def api_v7_statement_revision():
        try:
            source=_int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM ContractStatements WHERE id=? AND is_active=1",source)
                if not old:return _err('صورت‌وضعیت یافت نشد')
                if old.get('business_status')!='employer_rejected':
                    return _err('نسخه اصلاحی فقط بعد از رد کارفرما ساخته می‌شود')
                cols=['contract_id','statement_type_id','statement_number','statement_date','statement_date_fa','letter_number','title','start_date','start_date_fa','end_date','end_date_fa','requested_price','requested_without_vat','requested_vat','vat_factor_number','confirmed_price','confirmed_without_vat','confirmed_vat','confirmed_vat_factor_number','progress_percentage','period_year','period_month','vat_percentage_value','without_vat','description','project_team_id']
                vals=[old.get(x) for x in cols]
                cur.execute("INSERT INTO ContractStatements("+','.join(cols)+",business_status,supersedes_statement_id,created_by) OUTPUT INSERTED.id VALUES("+','.join('?' for _ in cols)+",'draft',?,?)",*(vals+[source,g.user['id']]))
                new_id=cur.fetchone()[0]
                cur.execute("""INSERT INTO ContractStatementTeams(statement_id,
                    project_team_id,allocation_percent,is_active,linked_by)
                    SELECT ?,project_team_id,allocation_percent,1,?
                    FROM ContractStatementTeams
                    WHERE statement_id=? AND is_active=1""",
                            new_id,g.user['id'],source)
                cur.execute("UPDATE ContractStatements SET is_current=0,business_status='revised',updated_by=?,updated_at=GETDATE() WHERE id=?",g.user['id'],source)
                _entity_audit(cur,'statement',source,'superseded',old,{'new_statement_id':new_id});c.commit()
            return _ok(id=new_id)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/statement_archive', methods=['POST'])
    @require_auth
    def api_v7_statement_archive():
        return _archive_entity('ContractStatements','statement')

    def _dec(value):
        return _decimal.Decimal(str(value or 0))

    def _pct(value, total):
        value, total = _dec(value), _dec(total)
        if not total:
            return 0
        return float((value * _decimal.Decimal(100) / total).quantize(
            _decimal.Decimal('0.01'), rounding=_decimal.ROUND_HALF_UP))

    def _in_date_range(value, start, end):
        """Inclusive event-date test used for sent/approved financial events."""
        day = _date(value)
        if not day:
            return start is None and end is None
        return (not start or day >= start) and (not end or day <= end)

    def _financial_report_data(cur, filters, user):
        """Build a drill-down-safe financial cube without multiplying joins.

        Sent money is always the requested amount without VAT.  Revenue is
        always the employer-approved amount without VAT.  The two event dates
        are filtered independently, so a statement sent in one period and
        approved in another is represented correctly in both periods.
        """
        year = _int(filters.get('year'))
        month = _int(filters.get('month'))
        project_type_id = _int(filters.get('project_type_id'))
        project_id = _int(filters.get('project_id'))
        city_id = _int(filters.get('city_id'))
        contract_type_id = _int(filters.get('contract_type_id'))
        start = _date(filters.get('from_date'))
        end = _date(filters.get('to_date'))
        if month is not None and not 1 <= month <= 12:
            raise ValueError('ماه گزارش نامعتبر است')
        if start and end and end < start:
            raise ValueError('تاریخ پایان گزارش قبل از تاریخ شروع است')

        contract_where = ["co.is_active=1"]
        contract_params = []
        if project_type_id is not None:
            contract_where.append("ISNULL(p.project_type_id,-1)=?")
            contract_params.append(project_type_id)
        if project_id is not None:
            contract_where.append("ISNULL(co.project_id,-1)=?")
            contract_params.append(project_id)
        if city_id is not None:
            contract_where.append("ISNULL(ci.id,-1)=?")
            contract_params.append(city_id)
        if contract_type_id is not None:
            contract_where.append("ISNULL(co.contract_type_id,0)=?")
            contract_params.append(contract_type_id)
        team_clause, team_params = _team_scope_condition(
            user, filters, 'contract', 'co'
        )
        if team_clause:
            contract_where.append(team_clause)
            contract_params.extend(team_params)
        cur.execute(_contract_select(" WHERE " + " AND ".join(contract_where)),
                    contract_params)
        contracts = rows_to_list(cur)

        statement_where = ["s.is_active=1", "s.is_current=1"]
        statement_params = []
        # Explicit date filters take precedence over the Jalali period fields.
        if not (start or end):
            if year is not None:
                statement_where.append("s.period_year=?")
                statement_params.append(year)
            if month is not None:
                statement_where.append("s.period_month=?")
                statement_params.append(month)
        if project_type_id is not None:
            statement_where.append("ISNULL(p.project_type_id,-1)=?")
            statement_params.append(project_type_id)
        if project_id is not None:
            statement_where.append("ISNULL(co.project_id,-1)=?")
            statement_params.append(project_id)
        if city_id is not None:
            statement_where.append("ISNULL(ci.id,-1)=?")
            statement_params.append(city_id)
        if contract_type_id is not None:
            statement_where.append("ISNULL(co.contract_type_id,0)=?")
            statement_params.append(contract_type_id)
        team_clause, team_params = _team_scope_condition(
            user, filters, 'statement', 's'
        )
        if team_clause:
            statement_where.append(team_clause)
            statement_params.extend(team_params)
        share_sql, share_params = _statement_scope_share(
            user, filters, 's'
        )
        cur.execute(("""SELECT s.id,s.contract_id,s.title,s.statement_number,s.statement_date,
                s.statement_date_fa,s.period_year,s.period_month,s.business_status,
                s.requested_price,s.requested_without_vat,s.confirmed_price,s.confirmed_without_vat,
                s.sent_at,s.employer_decision_at,st.name AS statement_type_name,
                co.title AS contract_title,co.contract_number,co.project_id,
                ISNULL(co.contract_type_id,0) AS contract_type_id,
                ISNULL(ct.name,N'نوع تعیین نشده') AS contract_type_name,
                p.name AS project_name,p.project_type_id,pt.name AS project_type_name,
                ci.id AS city_id,ci.name AS city_name,
                %s AS scope_percent
            FROM ContractStatements s JOIN Contracts co ON co.id=s.contract_id
            LEFT JOIN ContractStatementTypes st ON st.id=s.statement_type_id
            LEFT JOIN ContractTypes ct ON ct.id=ISNULL(co.contract_type_id,0)
            LEFT JOIN Projects p ON p.id=co.project_id
            LEFT JOIN ProjectTypes pt ON pt.id=p.project_type_id
            LEFT JOIN Cities ci ON ci.id=p.city_id
            WHERE """ + " AND ".join(statement_where) +
            " ORDER BY ci.name,pt.name,p.name,co.title,s.statement_date,s.id") %
            share_sql, share_params + statement_params)
        statements = rows_to_list(cur)

        plan_where = ["fp.is_current=1", "ps.is_active=1"]
        plan_params = []
        if year is not None:
            plan_where.append("fp.jalali_year=?")
            plan_params.append(year)
        if month is not None:
            plan_where.append("ps.month_no=?")
            plan_params.append(month)
        if project_type_id is not None:
            plan_where.append("ISNULL(p.project_type_id,-1)=?")
            plan_params.append(project_type_id)
        if project_id is not None:
            plan_where.append("ISNULL(co.project_id,-1)=?")
            plan_params.append(project_id)
        if city_id is not None:
            plan_where.append("ISNULL(p.city_id,-1)=?")
            plan_params.append(city_id)
        if contract_type_id is not None:
            plan_where.append("ISNULL(co.contract_type_id,0)=?")
            plan_params.append(contract_type_id)
        team_clause, team_params = _team_scope_condition(
            user, filters, 'planned', 'ps'
        )
        if team_clause:
            plan_where.append(team_clause)
            plan_params.extend(team_params)
        cur.execute("""SELECT ps.contract_id,SUM(ISNULL(ps.planned_amount,0)) AS planned_amount
            FROM PlannedStatements ps JOIN FinancialPlans fp ON fp.id=ps.plan_id
            JOIN Contracts co ON co.id=ps.contract_id
            LEFT JOIN Projects p ON p.id=co.project_id
            WHERE """ + " AND ".join(plan_where) + " GROUP BY ps.contract_id", plan_params)
        planned_by_contract = {int(x['contract_id']): _dec(x.get('planned_amount'))
                               for x in rows_to_list(cur)}

        annual_target = _decimal.Decimal(0)
        if year is not None:
            selected_team = _int(filters.get('_team_scope'))
            if selected_team:
                target_where = [
                    "fp.jalali_year=?", "fp.is_current=1", "tf.is_current=1"
                ]
                target_params = [year]
                if selected_team:
                    target_where.append("tf.team_id=?")
                    target_params.append(selected_team)
                plan = _one(cur, """SELECT ISNULL(SUM(tf.annual_target),0)
                        AS annual_target
                    FROM TeamFinancialTargets tf JOIN FinancialPlans fp
                      ON fp.id=tf.plan_id WHERE """ +
                            " AND ".join(target_where), *target_params)
            else:
                plan = _one(cur, """SELECT TOP 1 annual_target FROM FinancialPlans
                    WHERE jalali_year=? AND is_current=1 ORDER BY id DESC""", year)
            annual_target = _dec((plan or {}).get('annual_target'))

        statement_by_contract = {}
        visible_statements = []
        for item in statements:
            status = item.get('business_status')
            share = _dec(item.get('scope_percent')) / _decimal.Decimal(100)
            sent_amount = (_dec(item.get('requested_without_vat')
                                if item.get('requested_without_vat') is not None
                                else item.get('requested_price'))
                           if status in ('sent', 'employer_approved', 'employer_rejected') else _decimal.Decimal(0))
            approved_amount = (_dec(item.get('confirmed_without_vat')
                                    if item.get('confirmed_without_vat') is not None
                                    else (item.get('confirmed_price')
                                          if item.get('confirmed_price') is not None
                                          else (item.get('requested_without_vat')
                                                if item.get('requested_without_vat') is not None
                                                else item.get('requested_price'))))
                               if _statement_is_employer_approved(item) else _decimal.Decimal(0))
            sent_amount *= share
            approved_amount *= share
            if start or end:
                if not _in_date_range(item.get('sent_at') or item.get('statement_date'), start, end):
                    sent_amount = _decimal.Decimal(0)
                if not _in_date_range(item.get('employer_decision_at') or item.get('statement_date'), start, end):
                    approved_amount = _decimal.Decimal(0)
                # Keep only records that create a financial event in the range.
                if not sent_amount and not approved_amount:
                    continue
            item['sent_amount'] = sent_amount
            item['approved_amount'] = approved_amount
            visible_statements.append(item)
            bucket = statement_by_contract.setdefault(int(item['contract_id']), {
                'sent_amount': _decimal.Decimal(0), 'approved_amount': _decimal.Decimal(0),
                'statement_count': 0})
            bucket['sent_amount'] += sent_amount
            bucket['approved_amount'] += approved_amount
            bucket['statement_count'] += 1

        # A filtered statement can reference an archived contract. Preserve it
        # in drill-down output even though the active-contract query omits it.
        known = {int(x['id']) for x in contracts}
        for item in visible_statements:
            cid = int(item['contract_id'])
            if cid not in known:
                contracts.append({
                    'id': cid, 'title': item.get('contract_title'),
                    'contract_number': item.get('contract_number'),
                    'project_id': item.get('project_id'), 'project_name': item.get('project_name'),
                    'project_type_id': item.get('project_type_id'),
                    'project_type_name': item.get('project_type_name'),
                    'contract_type_id': item.get('contract_type_id'),
                    'contract_type_name': item.get('contract_type_name'),
                    'city_id': item.get('city_id'), 'city_name': item.get('city_name'),
                    'effective_price': 0, 'approved_statement_total': 0,
                    'is_active': False})
                known.add(cid)

        type_map, project_map, city_map, contract_type_map = {}, {}, {}, {}
        total_planned = total_sent = total_approved = _decimal.Decimal(0)
        report_contracts = []
        for item in contracts:
            cid = int(item['id'])
            metrics = statement_by_contract.get(cid, {})
            planned = planned_by_contract.get(cid, _decimal.Decimal(0))
            sent = _dec(metrics.get('sent_amount'))
            approved = _dec(metrics.get('approved_amount'))
            p_id = _int(item.get('project_id'), -1)
            pt_id = _int(item.get('project_type_id'), -1 if p_id == -1 else 0)
            p_name = item.get('project_name') or 'بدون پروژه'
            pt_name = item.get('project_type_name') or ('بدون پروژه' if p_id == -1 else 'نوع تعیین نشده')
            remaining = contract_remaining(item.get('effective_price'),
                [{'confirmed_without_vat': item.get('approved_statement_total')}])
            row = dict(item)
            row.update({'planned_amount': planned, 'sent_amount': sent,
                        'approved_amount': approved,
                        'statement_count_in_filter': int(metrics.get('statement_count') or 0),
                        'remaining_price': remaining})
            report_contracts.append(row)
            total_planned += planned; total_sent += sent; total_approved += approved
            type_row = type_map.setdefault(pt_id, {
                'project_type_id': pt_id, 'project_type_name': pt_name,
                'planned_amount': _decimal.Decimal(0), 'sent_amount': _decimal.Decimal(0),
                'approved_amount': _decimal.Decimal(0), 'contract_count': 0,
                'project_ids': set()})
            project_row = project_map.setdefault(p_id, {
                'project_id': p_id, 'project_name': p_name, 'project_type_id': pt_id,
                'project_type_name': pt_name, 'city_id': item.get('city_id'),
                'city_name': item.get('city_name') or 'بدون شهر',
                'planned_amount': _decimal.Decimal(0), 'sent_amount': _decimal.Decimal(0),
                'approved_amount': _decimal.Decimal(0), 'contract_count': 0})
            cty_id = _int(item.get('city_id'), -1)
            city_row = city_map.setdefault(cty_id, {
                'city_id': cty_id, 'city_name': item.get('city_name') or 'بدون شهر',
                'planned_amount': _decimal.Decimal(0), 'sent_amount': _decimal.Decimal(0),
                'approved_amount': _decimal.Decimal(0), 'contract_count': 0,
                'project_ids': set()})
            ct_id = _int(item.get('contract_type_id'), 0)
            contract_type_row = contract_type_map.setdefault(ct_id, {
                'contract_type_id': ct_id,
                'contract_type_name': item.get('contract_type_name') or 'نوع تعیین نشده',
                'planned_amount': _decimal.Decimal(0), 'sent_amount': _decimal.Decimal(0),
                'approved_amount': _decimal.Decimal(0), 'contract_count': 0,
                'project_ids': set()})
            for bucket in (type_row, project_row, city_row, contract_type_row):
                bucket['planned_amount'] += planned
                bucket['sent_amount'] += sent
                bucket['approved_amount'] += approved
                bucket['contract_count'] += 1
            type_row['project_ids'].add(p_id)
            city_row['project_ids'].add(p_id)
            contract_type_row['project_ids'].add(p_id)

        def finalize(row):
            row = dict(row)
            if 'project_ids' in row:
                row['project_count'] = len(row.pop('project_ids'))
            row['revenue_share_percent'] = _pct(row.get('approved_amount'), total_approved)
            row['achievement_percent'] = _pct(row.get('approved_amount'), row.get('planned_amount'))
            return row

        report_contracts = [finalize(x) for x in report_contracts]
        type_rows = [finalize(x) for x in type_map.values()]
        project_rows = [finalize(x) for x in project_map.values()]
        city_rows = [finalize(x) for x in city_map.values()]
        contract_type_rows = [finalize(x) for x in contract_type_map.values()]
        type_rows.sort(key=lambda x: (_dec(x.get('approved_amount')), _dec(x.get('sent_amount'))), reverse=True)
        project_rows.sort(key=lambda x: (_dec(x.get('approved_amount')), _dec(x.get('sent_amount'))), reverse=True)
        city_rows.sort(key=lambda x: (_dec(x.get('approved_amount')), _dec(x.get('sent_amount'))), reverse=True)
        contract_type_rows.sort(key=lambda x: (_dec(x.get('approved_amount')), _dec(x.get('sent_amount'))), reverse=True)
        report_contracts.sort(key=lambda x: (_dec(x.get('approved_amount')), _dec(x.get('sent_amount'))), reverse=True)
        return {
            'totals': {'annual_target': annual_target, 'planned_amount': total_planned,
                       'sent_amount': total_sent, 'approved_amount': total_approved,
                       'remaining_to_plan': max(_decimal.Decimal(0), total_planned - total_approved),
                       'achievement_percent': _pct(total_approved, total_planned)},
            'types': type_rows, 'projects': project_rows,
            'cities': city_rows, 'contract_types': contract_type_rows,
            'contracts': report_contracts, 'statements': visible_statements,
            'filters': {'year': year, 'month': month, 'from_date': start,
                        'to_date': end, 'project_type_id': project_type_id,
                        'project_id': project_id, 'city_id': city_id,
                        'contract_type_id': contract_type_id}}

    @app.route('/api/project_financial_report', methods=['POST'])
    @require_auth
    def api_v71_project_financial_report():
        try:
            with db_lock:
                c = get_conn(); cur = c.cursor()
                data = _financial_report_data(
                    cur, request.get_json() or {}, g.user
                )
            return _ok(**data)
        except ValueError as exc:
            return _err(str(exc))
        except Exception as exc:
            return _err(str(exc), types=[], projects=[], contracts=[], statements=[])

    def _parse_datetime(value):
        if not value:
            return None
        if isinstance(value, _dt.datetime):
            return value
        if isinstance(value, _dt.date):
            return _dt.datetime.combine(value, _dt.time.min)
        try:
            return _dt.datetime.fromisoformat(str(value).replace('Z', '+00:00')).replace(tzinfo=None)
        except Exception:
            return None

    def _work_report_data(cur, filters, user):
        start_day, end_day = _date(filters.get('from_date')), _date(filters.get('to_date'))
        if start_day and end_day and end_day < start_day:
            raise ValueError('تاریخ پایان گزارش قبل از تاریخ شروع است')
        start_dt = _dt.datetime.combine(start_day, _dt.time.min) if start_day else None
        end_dt = (_dt.datetime.combine(end_day + _dt.timedelta(days=1), _dt.time.min)
                  if end_day else None)
        project_type_id = _int(filters.get('project_type_id'))
        project_id = _int(filters.get('project_id'))
        city_id = _int(filters.get('city_id'))
        contract_type_id = _int(filters.get('contract_type_id'))
        requested_user_id = _int(filters.get('user_id'))
        # R10: team-scoped reports may include every teammate. The task/team
        # relation below is the data boundary; a broad operation permission
        # never expands the caller beyond active team memberships.
        where, params = ["1=1"], []
        if start_dt:
            where.append("COALESCE(tl.ended_at,GETDATE())>?"); params.append(start_dt)
        if end_dt:
            where.append("tl.started_at<?"); params.append(end_dt)
        if project_type_id is not None:
            where.append("ISNULL(p.project_type_id,-1)=?"); params.append(project_type_id)
        if project_id is not None:
            where.append("ISNULL(t.project_id,-1)=?"); params.append(project_id)
        # Being in a team is not permission to read every teammate's timesheet.
        # Without tasks.view_all - the same key that decides whether the task
        # lists are personal or team-wide - the work-share report is limited to
        # the caller's own entries, and an injected user_id cannot widen it.
        if not user_has_permission(user, 'tasks.view_all'):
            where.append("tl.user_id=?"); params.append(user['id'])
        elif requested_user_id is not None:
            where.append("tl.user_id=?"); params.append(requested_user_id)
        team_clause, team_params = _team_scope_condition(
            user, filters, 'task', 't'
        )
        if team_clause:
            where.append(team_clause)
            params.extend(team_params)
        cur.execute("""SELECT tl.id,tl.task_id,tl.user_id,tl.started_at,tl.ended_at,tl.seconds,
                t.title AS task_title,t.status,t.project_id,p.name AS project_name,
                p.project_type_id,pt.name AS project_type_name,ci.name AS city_name,
                u.display_name,u.username
            FROM TaskTimeLog tl JOIN Tasks t ON t.id=tl.task_id
            JOIN Users u ON u.id=tl.user_id LEFT JOIN Projects p ON p.id=t.project_id
            LEFT JOIN ProjectTypes pt ON pt.id=p.project_type_id LEFT JOIN Cities ci ON ci.id=p.city_id
            WHERE """ + " AND ".join(where) + " ORDER BY u.display_name,p.name,t.title,tl.started_at", params)
        logs = rows_to_list(cur)
        now = _dt.datetime.now()
        pair_map, task_map, user_totals, project_totals = {}, {}, {}, {}
        total_seconds = 0
        for log in logs:
            started = _parse_datetime(log.get('started_at'))
            ended = _parse_datetime(log.get('ended_at')) or now
            if not started or ended <= started:
                continue
            clipped_start = max(started, start_dt) if start_dt else started
            clipped_end = min(ended, end_dt) if end_dt else ended
            if clipped_end <= clipped_start:
                continue
            if start_dt or end_dt or not log.get('ended_at'):
                seconds = int((clipped_end - clipped_start).total_seconds())
            else:
                seconds = int(log.get('seconds') or (ended - started).total_seconds())
            if seconds <= 0:
                continue
            uid, pid, tid = int(log['user_id']), _int(log.get('project_id'), -1), int(log['task_id'])
            pair = pair_map.setdefault((uid, pid), {
                'user_id': uid, 'display_name': log.get('display_name') or log.get('username'),
                'project_id': pid, 'project_name': log.get('project_name') or 'بدون پروژه',
                'project_type_id': _int(log.get('project_type_id'), -1 if pid == -1 else 0),
                'project_type_name': log.get('project_type_name') or ('بدون پروژه' if pid == -1 else 'نوع تعیین نشده'),
                'city_name': log.get('city_name') or 'بدون شهر', 'seconds': 0,
                'task_ids': set()})
            pair['seconds'] += seconds; pair['task_ids'].add(tid)
            detail = task_map.setdefault((uid, pid, tid), {
                'user_id': uid, 'display_name': pair['display_name'], 'project_id': pid,
                'project_name': pair['project_name'], 'task_id': tid,
                'task_title': log.get('task_title'), 'status': log.get('status'), 'seconds': 0})
            detail['seconds'] += seconds
            user_totals[uid] = user_totals.get(uid, 0) + seconds
            project_totals[pid] = project_totals.get(pid, 0) + seconds
            total_seconds += seconds
        rows = []
        for pair in pair_map.values():
            pair['task_count'] = len(pair.pop('task_ids'))
            pair['project_share_percent'] = round(pair['seconds'] * 100 / project_totals[pair['project_id']], 2)
            pair['person_focus_percent'] = round(pair['seconds'] * 100 / user_totals[pair['user_id']], 2)
            pair['overall_percent'] = round(pair['seconds'] * 100 / total_seconds, 2) if total_seconds else 0
            rows.append(pair)
        rows.sort(key=lambda x: x['seconds'], reverse=True)
        details = sorted(task_map.values(), key=lambda x: x['seconds'], reverse=True)
        return {'totals': {'seconds': total_seconds, 'user_count': len(user_totals),
                           'project_count': len(project_totals)},
                'rows': rows, 'tasks': details,
                'filters': {'from_date': start_day, 'to_date': end_day,
                            'project_type_id': project_type_id, 'project_id': project_id,
                            'user_id': requested_user_id}}

    @app.route('/api/project_work_report', methods=['POST'])
    @require_auth
    def api_v71_project_work_report():
        if not user_has_permission(g.user, 'reports.view'):
            return _err('اجازه مشاهده این گزارش را ندارید', 403)
        try:
            with db_lock:
                c = get_conn(); cur = c.cursor()
                data = _work_report_data(cur, request.get_json() or {}, g.user)
            return _ok(**data)
        except ValueError as exc:
            return _err(str(exc))
        except Exception as exc:
            return _err(str(exc), rows=[], tasks=[])

    @app.route('/api/dashboard_reports', methods=['POST'])
    @require_auth
    def api_v71_dashboard_reports():
        """All dashboard charts under one consistent, all-time-by-default filter."""
        try:
            d = request.get_json() or {}
            start, end = _date(d.get('from_date')), _date(d.get('to_date'))
            start_fa = (d.get('from_jalali') or '').strip() or None
            end_fa = (d.get('to_jalali') or '').strip() or None
            if start and end and end < start:
                return _err('تاریخ پایان فیلتر قبل از تاریخ شروع است')
            task_where, task_params = [], []
            if start_fa:
                task_where.append("t.date_recv>=?"); task_params.append(start_fa)
            if end_fa:
                task_where.append("t.date_recv<=?"); task_params.append(end_fa)
            scope_clause, scope_params = _team_scope_condition(
                g.user, d, 'task', 't'
            )
            if scope_clause:
                task_where.append(scope_clause)
                task_params.extend(scope_params)
            task_clause = (" WHERE " + " AND ".join(task_where)) if task_where else ""
            done_where, done_params = ["t.status='done'", "t.work_seconds>0"], []
            if start:
                done_where.append("t.completed_at>=?")
                done_params.append(_dt.datetime.combine(start, _dt.time.min))
            if end:
                done_where.append("t.completed_at<?")
                done_params.append(_dt.datetime.combine(end + _dt.timedelta(days=1), _dt.time.min))
            if scope_clause:
                done_where.append(scope_clause)
                done_params.extend(scope_params)
            done_clause = " WHERE " + " AND ".join(done_where)
            with db_lock:
                c = get_conn(); cur = c.cursor()
                cur.execute("""SELECT ISNULL(ci.name,N'بدون شهر') AS label,COUNT(t.id) AS value
                    FROM Tasks t LEFT JOIN Projects p ON p.id=t.project_id
                    LEFT JOIN Cities ci ON ci.id=p.city_id""" + task_clause +
                    " GROUP BY ci.name ORDER BY value DESC", task_params)
                by_city = rows_to_list(cur)
                cur.execute("""SELECT TOP 12 ISNULL(p.name,N'بدون پروژه') AS label,COUNT(t.id) AS value
                    FROM Tasks t LEFT JOIN Projects p ON p.id=t.project_id""" + task_clause +
                    " GROUP BY p.name ORDER BY value DESC", task_params)
                by_project = rows_to_list(cur)
                cur.execute("""SELECT ISNULL(pt.name,N'نوع تعیین نشده') AS label,COUNT(t.id) AS value
                    FROM Tasks t LEFT JOIN Projects p ON p.id=t.project_id
                    LEFT JOIN ProjectTypes pt ON pt.id=p.project_type_id""" + task_clause +
                    " GROUP BY pt.name ORDER BY value DESC", task_params)
                by_project_type = rows_to_list(cur)
                cur.execute("SELECT t.status,COUNT(*) AS value FROM Tasks t" + task_clause +
                            " GROUP BY t.status", task_params)
                by_status = {x.get('status'): int(x.get('value') or 0) for x in rows_to_list(cur)}
                cur.execute("""SELECT ISNULL(u.display_name,N'بدون پشتیبان') AS label,
                        SUM(CASE WHEN t.status='done' THEN 1 ELSE 0 END) AS done,
                        SUM(CASE WHEN t.status NOT IN ('done','rejected') THEN 1 ELSE 0 END) AS open_count
                    FROM Tasks t LEFT JOIN Users u ON u.id=t.staff_id""" +
                    (task_clause + (" AND " if task_clause else " WHERE ") + "t.staff_id IS NOT NULL") +
                    " GROUP BY u.display_name ORDER BY open_count DESC", task_params)
                by_staff = []
                for row in rows_to_list(cur):
                    by_staff.append({'label': row.get('label'), 'done': int(row.get('done') or 0),
                                     'open': int(row.get('open_count') or 0)})
                cur.execute("""SELECT ISNULL(u.display_name,N'—') AS label,
                        AVG(CAST(t.work_seconds AS FLOAT)) AS avg_seconds,COUNT(*) AS count
                    FROM Tasks t LEFT JOIN Users u ON u.id=t.staff_id""" + done_clause +
                    " AND t.staff_id IS NOT NULL GROUP BY u.display_name ORDER BY avg_seconds ASC", done_params)
                analytics_staff = []
                for row in rows_to_list(cur):
                    analytics_staff.append({'label': row.get('label'),
                        'avg_seconds': int(row.get('avg_seconds') or 0),
                        'count': int(row.get('count') or 0)})

                comparison_year = _int(d.get('comparison_jalali_year'))
                comparison_month = _int(d.get('comparison_jalali_month'))
                if not comparison_year or not comparison_month:
                    today = _dt.date.today()
                    comparison_year, comparison_month, _ = _gregorian_to_jalali(
                        today.year, today.month, today.day
                    )
                current_start, current_end, previous_start, previous_year, previous_month = \
                    _jalali_month_window(comparison_year, comparison_month)
                comparison_where = [
                    "t.status='done'", "t.completed_at IS NOT NULL"
                ]
                comparison_params = [
                    _dt.datetime.combine(current_start, _dt.time.min),
                    _dt.datetime.combine(current_end, _dt.time.min),
                    _dt.datetime.combine(previous_start, _dt.time.min),
                    _dt.datetime.combine(current_start, _dt.time.min),
                ]
                if scope_clause:
                    comparison_where.append(scope_clause)
                    comparison_params.extend(scope_params)
                cur.execute("""SELECT
                    SUM(CASE WHEN t.completed_at>=? AND t.completed_at<? THEN 1 ELSE 0 END) AS selected_count,
                    SUM(CASE WHEN t.completed_at>=? AND t.completed_at<? THEN 1 ELSE 0 END) AS previous_count
                    FROM Tasks t WHERE """ + " AND ".join(comparison_where),
                            comparison_params)
                values = cur.fetchone()
                selected_count = int((values[0] if values else 0) or 0)
                previous_count = int((values[1] if values else 0) or 0)
                selected_label = '%s %s' % (
                    JALALI_MONTH_NAMES[comparison_month - 1], comparison_year
                )
                previous_label = '%s %s' % (
                    JALALI_MONTH_NAMES[previous_month - 1], previous_year
                )
            return _ok(by_city=by_city, by_project=by_project,
                       by_project_type=by_project_type, by_status=by_status,
                       by_staff=by_staff, analytics_staff=analytics_staff,
                       selected_count=selected_count, previous_count=previous_count,
                       selected_label=selected_label, previous_label=previous_label,
                       comparison_jalali_year=comparison_year,
                       comparison_jalali_month=comparison_month,
                       previous_jalali_year=previous_year,
                       previous_jalali_month=previous_month,
                       filters={'from_date': start, 'to_date': end,
                                'from_jalali': start_fa, 'to_jalali': end_fa})
        except Exception as exc:
            return _err(str(exc))

    def _archive_entity(table, entity):
        try:
            ident=_int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM %s WHERE id=?"%table,ident)
                if not old:return _err('رکورد یافت نشد')
                cur.execute("UPDATE %s SET is_active=0,updated_by=?,updated_at=GETDATE() WHERE id=?"%table,g.user['id'],ident)
                _entity_audit(cur,entity,ident,'archive',old,None);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/financial_plan', methods=['POST'])
    @require_auth
    def api_v7_financial_plan():
        try:
            d=request.get_json() or {}
            year=_int(d.get('year')) or 1405
            with db_lock:
                c=get_conn();cur=c.cursor();plan=_one(cur,"SELECT TOP 1 * FROM FinancialPlans WHERE jalali_year=? AND is_current=1 ORDER BY id DESC",year)
                periods=[];planned=[]
                if plan:
                    cur.execute("SELECT * FROM FinancialPlanPeriods WHERE plan_id=? ORDER BY month_no",plan['id']);periods=rows_to_list(cur)
                    planned_scope, planned_scope_params = _team_scope_condition(
                        g.user, d, 'planned', 'ps'
                    )
                    planned_where = ["ps.plan_id=?", "ps.is_active=1"]
                    planned_params = [plan['id']]
                    if planned_scope:
                        planned_where.append(planned_scope)
                        planned_params.extend(planned_scope_params)
                    cur.execute("""SELECT ps.*,co.title AS contract_title,co.contract_number,p.name AS project_name,ci.name AS city_name,
                        team.name AS team_name,ptm.team_id,
                        (SELECT COUNT(*) FROM PlannedStatementTasks pt WHERE pt.planned_statement_id=ps.id) AS task_count
                        FROM PlannedStatements ps JOIN Contracts co ON co.id=ps.contract_id LEFT JOIN Projects p ON p.id=co.project_id
                        LEFT JOIN Cities ci ON ci.id=p.city_id
                        LEFT JOIN ProjectTeams ptm ON ptm.id=ps.project_team_id
                        LEFT JOIN Teams team ON team.id=ptm.team_id
                        WHERE """+" AND ".join(planned_where)+
                        " ORDER BY ps.month_no,ps.planned_date,ps.id",
                                planned_params);planned=rows_to_list(cur)
                    cur.execute("""SELECT pt.planned_statement_id,pt.task_id FROM PlannedStatementTasks pt
                        JOIN PlannedStatements ps ON ps.id=pt.planned_statement_id
                        WHERE """+" AND ".join(planned_where)+
                        " ORDER BY pt.planned_statement_id,pt.task_id",
                                planned_params)
                    links=rows_to_list(cur);by_planned={}
                    for link in links:by_planned.setdefault(link['planned_statement_id'],[]).append(link['task_id'])
                    for item in planned:item['task_ids']=by_planned.get(item['id'],[])
                actual_scope, actual_scope_params = _team_scope_condition(
                    g.user, d, 'statement', 's'
                )
                actual_where = [
                    "s.is_active=1", "s.is_current=1", "s.period_year=?"
                ]
                actual_params = [year]
                if actual_scope:
                    actual_where.append(actual_scope)
                    actual_params.extend(actual_scope_params)
                actual_share_sql,actual_share_params=_statement_scope_share(
                    g.user,d,'s'
                )
                cur.execute(("""SELECT scoped.period_month AS month_no,
                    SUM(CASE WHEN scoped.business_status IN
                      ('sent','employer_approved','employer_rejected')
                      THEN COALESCE(scoped.requested_without_vat,
                        scoped.requested_price,0)*scoped.scope_percent/100
                      ELSE 0 END) AS sent_amount,
                    SUM(CASE WHEN (scoped.business_status='employer_approved' OR (scoped.employer_decision_at IS NOT NULL AND ISNULL(scoped.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(scoped.confirmed_without_vat,scoped.confirmed_price,0)>0))
                      THEN COALESCE(scoped.confirmed_without_vat,
                        scoped.confirmed_price,scoped.requested_without_vat,
                        scoped.requested_price,0)*scoped.scope_percent/100
                      ELSE 0 END) AS approved_amount
                    FROM (SELECT s.*,%s AS scope_percent
                      FROM ContractStatements s WHERE """+
                      " AND ".join(actual_where)+
                    """) scoped GROUP BY scoped.period_month""")%
                    actual_share_sql,
                    actual_share_params+actual_params);actuals=rows_to_list(cur)
            return _ok(plan=plan,periods=periods,planned=planned,actuals=actuals,year=year)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/financial_dashboard', methods=['POST'])
    @require_auth
    def api_v7_financial_dashboard():
        """Return annual/monthly Jalali target and live statement progress."""
        try:
            d = request.get_json() or {}
            selected_month = _int(d.get('jalali_month'))
            if selected_month is not None and not 1 <= selected_month <= 12:
                return _err('ماه شمسی انتخاب‌شده نامعتبر است', plan=None)
            with db_lock:
                c = get_conn(); cur = c.cursor()
                plan = _one(cur, """SELECT TOP 1 id,jalali_year,title,annual_target,version_no
                    FROM FinancialPlans WHERE is_current=1 ORDER BY jalali_year DESC,id DESC""")
                if not plan:
                    return _ok(plan=None, target_amount=0, sent_amount=0,
                               approved_amount=0, sent_count=0, approved_count=0)

                selected_team = _int(d.get('_team_scope'))
                target_amount = _dec(plan.get('annual_target'))
                if not has_company_scope(g.user):
                    if selected_team and not _team_allowed(cur, g.user, selected_team):
                        return _err('به تیم انتخاب‌شده دسترسی ندارید', plan=None)
                    team_filter = 'AND t.team_id=?' if selected_team else """AND EXISTS(
                        SELECT 1 FROM TeamMembers tm WHERE tm.team_id=t.team_id
                          AND tm.user_id=? AND tm.is_active=1)"""
                    team_param = selected_team if selected_team else g.user['id']
                    if selected_month:
                        row = _one(cur, """SELECT SUM(ISNULL(p.target_amount,0)) AS target_amount
                            FROM TeamFinancialTargets t
                            JOIN TeamFinancialTargetPeriods p ON p.target_id=t.id
                            WHERE t.plan_id=? AND t.is_current=1 AND p.month_no=? """ + team_filter,
                                   plan['id'], selected_month, team_param)
                    else:
                        row = _one(cur, """SELECT SUM(ISNULL(t.annual_target,0)) AS target_amount
                            FROM TeamFinancialTargets t
                            WHERE t.plan_id=? AND t.is_current=1 """ + team_filter,
                                   plan['id'], team_param)
                    target_amount = _dec((row or {}).get('target_amount'))
                    plan['annual_target'] = target_amount
                elif selected_team:
                    team_target = _one(cur, """SELECT TOP 1 id,annual_target
                        FROM TeamFinancialTargets WHERE plan_id=? AND team_id=?
                          AND is_current=1 ORDER BY id DESC""",
                                       plan['id'], selected_team)
                    target_amount = _dec((team_target or {}).get('annual_target'))
                    plan['annual_target'] = target_amount
                    if selected_month:
                        period = _one(cur, """SELECT p.target_amount
                            FROM TeamFinancialTargetPeriods p
                            WHERE p.target_id=? AND p.month_no=?""",
                                      (team_target or {}).get('id'), selected_month) if team_target else None
                        target_amount = _dec((period or {}).get('target_amount'))
                elif selected_month:
                    period = _one(cur, """SELECT target_amount
                        FROM FinancialPlanPeriods WHERE plan_id=? AND month_no=?""",
                                  plan['id'], selected_month)
                    target_amount = _dec((period or {}).get('target_amount'))

                share_sql, share_params = _statement_scope_share(g.user, d, 's')
                sql = ("""SELECT business_status,employer_decision_at,taxpayer_status_id,
                    requested_without_vat,requested_price,confirmed_without_vat,confirmed_price,
                    %s AS scope_percent
                    FROM ContractStatements s
                    WHERE s.is_active=1 AND s.is_current=1 AND s.period_year=?""" % share_sql)
                params = [plan['jalali_year']]
                if selected_month:
                    sql += " AND s.period_month=?"
                    params.append(selected_month)
                statement_scope, statement_scope_params = _team_scope_condition(
                    g.user, d, 'statement', 's'
                )
                if statement_scope:
                    sql += " AND " + statement_scope
                    params.extend(statement_scope_params)
                cur.execute(sql, share_params + params)
                statement_rows = rows_to_list(cur)

            sent_amount = approved_amount = _decimal.Decimal(0)
            sent_count = approved_count = 0
            for row in statement_rows:
                status = row.get('business_status')
                share = _dec(row.get('scope_percent')) / _decimal.Decimal(100)
                if status in ('sent', 'employer_approved', 'employer_rejected'):
                    sent_amount += (_dec(row.get('requested_without_vat')
                                         if row.get('requested_without_vat') is not None
                                         else row.get('requested_price')) * share)
                    sent_count += 1
                # New records use business_status. Migrated legacy records may
                # already contain the employer decision date and confirmed amount
                # while their textual status is still ``sent``. Treat the persisted
                # employer decision as authoritative so the dashboard percentage
                # never falls back to zero despite valid approved money.
                is_employer_approved = _statement_is_employer_approved(row)
                if is_employer_approved:
                    approved_amount += (_dec(row.get('confirmed_without_vat')
                                             if row.get('confirmed_without_vat') is not None
                                             else (row.get('confirmed_price')
                                                   if row.get('confirmed_price') is not None
                                                   else (row.get('requested_without_vat')
                                                         if row.get('requested_without_vat') is not None
                                                         else row.get('requested_price')))) * share)
                    approved_count += 1
            target_label = ('هدف ' + JALALI_MONTH_NAMES[selected_month - 1]) \
                if selected_month else 'هدف سالانه'
            return _ok(plan=plan, selected_month=selected_month,
                       selected_month_name=(JALALI_MONTH_NAMES[selected_month - 1]
                                            if selected_month else None),
                       target_label=target_label, target_amount=target_amount,
                       sent_amount=sent_amount, approved_amount=approved_amount,
                       sent_count=sent_count, approved_count=approved_count)
        except Exception as exc:
            return _err(str(exc), plan=None)

    @app.route('/api/financial_plan_save' , methods=['POST'])
    @require_auth
    def api_v7_financial_plan_save():
        try:
            if not user_has_permission(g.user, 'financial_plan.manage'):
                return _err('دسترسی تنظیم هدف کل شرکت را ندارید',403)
            # R12: annual_target is a roll-up of the team goals and is written
            # only by recompute_company_plan. What an administrator types here
            # is the board-approved figure, kept purely for comparison - it
            # never caps a team and never blocks a goal from being recorded.
            d=request.get_json() or {};year=_int(d.get('year'));target=_number(d.get('approved_target'))
            reason=(d.get('reason') or '').strip()
            if not year:return _err('سال معتبر الزامی است')
            if target is not None and target<0:return _err('هدف مصوب باید عددی نامنفی باشد')
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT TOP 1 * FROM FinancialPlans WHERE jalali_year=? AND is_current=1 ORDER BY id DESC",year)
                if old and _decimal.Decimal(str(old.get('approved_target') or 0))!=_decimal.Decimal(str(target or 0)) and not reason:return _err('برای تغییر هدف مصوب، ثبت دلیل الزامی است')
                if old:
                    cur.execute("INSERT INTO FinancialPlanRevisions(plan_id,old_target,new_target,reason,changed_by) VALUES(?,?,?,?,?)",old['id'],old.get('approved_target'),target,reason or 'ویرایش عنوان',g.user['id'])
                    cur.execute("UPDATE FinancialPlans SET title=?,approved_target=?,version_no=version_no+1,change_reason=?,updated_by=?,updated_at=GETDATE() WHERE id=?",d.get('title') or None,target,reason or None,g.user['id'],old['id']);pid=old['id']
                else:
                    cur.execute("INSERT INTO FinancialPlans(jalali_year,title,annual_target,approved_target,created_by) OUTPUT INSERTED.id VALUES(?,?,0,?,?)",year,d.get('title') or None,target,g.user['id']);pid=cur.fetchone()[0]
                    for m in range(1,13):cur.execute("INSERT INTO FinancialPlanPeriods(plan_id,month_no,target_amount) VALUES(?,?,0)",pid,m)
                _entity_audit(cur,'financial_plan',pid,'update' if old else 'create',old,{'approved_target':str(target),'reason':reason});c.commit()
            return _ok(id=pid)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/financial_periods_save', methods=['POST'])
    @require_auth
    def api_v7_financial_periods_save():
        # R12 inverted the goal model: the company monthly split is the sum of
        # the team monthly splits, written only by recompute_company_plan. A
        # hand-entered value here would be silently overwritten by the next
        # team goal save, so this refuses rather than pretending it worked.
        # The route stays registered so an out-of-date browser tab gets a clear
        # message instead of a 404.
        return _err('تقسیم ماهانه شرکت از جمع تقسیم ماهانه تیم‌ها ساخته می‌شود. ماه‌ها را در هدف همان تیم ویرایش کنید.')

    @app.route('/api/financial_period_lock', methods=['POST'])
    @require_auth
    def api_v7_financial_period_lock():
        try:
            d=request.get_json() or {};pid=_int(d.get('plan_id'));month=_int(d.get('month_no'));locked=bool(d.get('locked'))
            with db_lock:
                c=get_conn();cur=c.cursor();cur.execute("UPDATE FinancialPlanPeriods SET is_locked=?,updated_by=?,updated_at=GETDATE() WHERE plan_id=? AND month_no=?",locked,g.user['id'],pid,month);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/planned_statement_save', methods=['POST'])
    @require_auth
    def api_v7_planned_statement_save():
        try:
            d=request.get_json() or {};ident=_int(d.get('id'));pid=_int(d.get('plan_id'));cid=_int(d.get('contract_id'));month=_int(d.get('month_no'));amount=_number(d.get('planned_amount'));title=(d.get('title') or '').strip()
            if not pid or not cid or not month or not 1<=month<=12 or amount is None or amount<=0 or not title:return _err('برنامه، قرارداد، ماه، عنوان و مبلغ معتبر الزامی است')
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM PlannedStatements WHERE id=?",ident) if ident else None
                project_team_id=_int(d.get('project_team_id'))
                if not project_team_id:
                    cur.execute("""SELECT project_team_id FROM ContractProjectTeams
                        WHERE contract_id=? AND is_active=1 ORDER BY project_team_id""",cid)
                    candidates=[int(x[0]) for x in cur.fetchall()]
                    if len(candidates)==1:project_team_id=candidates[0]
                    elif len(candidates)>1:return _err('قرارداد چندتیمی است؛ تیم برنامه صورت‌وضعیت را انتخاب کنید')
                    else:return _err('برای این قرارداد تیم فعالی تعریف نشده است')
                valid_team=_one(cur,"""SELECT pt.team_id FROM ContractProjectTeams cpt
                    JOIN ProjectTeams pt ON pt.id=cpt.project_team_id
                    WHERE cpt.contract_id=? AND cpt.project_team_id=? AND cpt.is_active=1
                      AND pt.is_active=1""",cid,project_team_id)
                if not valid_team:return _err('تیم برنامه در این قرارداد فعال نیست')
                period=_one(cur,"SELECT target_amount FROM FinancialPlanPeriods WHERE plan_id=? AND month_no=?",pid,month)
                if not period:return _err('برای این ماه سهم مالی ثبت نشده است')
                used=_one(cur,"SELECT ISNULL(SUM(planned_amount),0) AS total FROM PlannedStatements WHERE plan_id=? AND month_no=? AND is_active=1 AND id<>ISNULL(?,0)",pid,month,ident) or {}
                if _decimal.Decimal(str(used.get('total') or 0))+amount>_decimal.Decimal(str(period.get('target_amount') or 0)):
                    return _err('جمع برنامه‌های این ماه از هدف همان ماه بیشتر می‌شود')
                vals=[pid,cid,month,_date(d.get('planned_date')),d.get('planned_date_fa') or None,_int(d.get('statement_type_id')),amount,title,(d.get('note') or '').strip() or None,d.get('status') or 'planned',project_team_id,g.user['id']]
                if ident:
                    cur.execute("UPDATE PlannedStatements SET plan_id=?,contract_id=?,month_no=?,planned_date=?,planned_date_fa=?,statement_type_id=?,planned_amount=?,title=?,note=?,status=?,project_team_id=?,updated_by=?,updated_at=GETDATE() WHERE id=?",*(vals+[ident]))
                else:
                    cur.execute("INSERT INTO PlannedStatements(plan_id,contract_id,month_no,planned_date,planned_date_fa,statement_type_id,planned_amount,title,note,status,project_team_id,created_by) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",vals);ident=cur.fetchone()[0]
                _entity_audit(cur,'planned_statement',ident,'update' if old else 'create',old,d);c.commit()
            return _ok(id=ident)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/planned_statement_archive', methods=['POST'])
    @require_auth
    def api_v7_planned_statement_archive():
        try:
            ident=_int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn();cur=c.cursor();old=_one(cur,"SELECT * FROM PlannedStatements WHERE id=? AND is_active=1",ident)
                if not old:return _err('برنامه صورت‌وضعیت یافت نشد')
                cur.execute("UPDATE Tasks SET planned_statement_id=NULL WHERE planned_statement_id=?",ident)
                cur.execute("DELETE FROM PlannedStatementTasks WHERE planned_statement_id=?",ident)
                cur.execute("UPDATE PlannedStatements SET is_active=0,updated_by=?,updated_at=GETDATE() WHERE id=?",g.user['id'],ident)
                _entity_audit(cur,'planned_statement',ident,'archive',old,None);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/planned_statement_link_tasks', methods=['POST'])
    @require_auth
    def api_v7_planned_link_tasks():
        try:
            d=request.get_json() or {};pid=_int(d.get('planned_statement_id'));task_ids=list(dict.fromkeys(_int(x) for x in (d.get('task_ids') or []) if _int(x)))
            with db_lock:
                c=get_conn();cur=c.cursor();ps=_one(cur,"SELECT contract_id,project_team_id FROM PlannedStatements WHERE id=? AND is_active=1",pid)
                if not ps:return _err('برنامه صورت‌وضعیت یافت نشد')
                if task_ids:
                    marks=','.join('?' for _ in task_ids)
                    cur.execute("SELECT id,contract_id,project_team_id FROM Tasks WHERE id IN (%s)"%marks,task_ids);task_rows=rows_to_list(cur)
                    if len(task_rows)!=len(task_ids):return _err('یکی از تسک‌های انتخاب‌شده یافت نشد')
                    if any(x.get('contract_id') is not None and _int(x.get('contract_id'))!=_int(ps['contract_id']) for x in task_rows):
                        return _err('یکی از تسک‌ها به قرارداد دیگری متصل است')
                    if any(x.get('project_team_id') is not None and _int(x.get('project_team_id'))!=_int(ps.get('project_team_id')) for x in task_rows):
                        return _err('همه تسک‌های برنامه باید متعلق به همان تیم صورت‌وضعیت باشند')
                cur.execute("UPDATE Tasks SET planned_statement_id=NULL WHERE planned_statement_id=?",pid)
                cur.execute("DELETE FROM PlannedStatementTasks WHERE planned_statement_id=?",pid)
                for tid in task_ids:
                    cur.execute("INSERT INTO PlannedStatementTasks(planned_statement_id,task_id,linked_by) VALUES(?,?,?)",pid,tid,g.user['id'])
                    cur.execute("""UPDATE Tasks SET planned_statement_id=?,
                        contract_id=COALESCE(contract_id,?),
                        project_team_id=COALESCE(project_team_id,?) WHERE id=?""",
                                pid,ps['contract_id'],ps.get('project_team_id'),tid)
                _entity_audit(cur,'planned_statement',pid,'link_tasks',None,{'task_ids':task_ids});c.commit()
            return _ok(count=len(task_ids))
        except Exception as exc:return _err(str(exc))

    @app.route('/api/contract_recommendations', methods=['POST'])
    @require_auth
    def api_v7_recommendations():
        try:
            d=request.get_json() or {}
            with db_lock:
                c=get_conn();cur=c.cursor()
                where=["co.is_active=1"];params=[]
                scope_clause,scope_params=_team_scope_condition(
                    g.user,d,'contract','co'
                )
                if scope_clause:
                    where.append(scope_clause);params.extend(scope_params)
                cur.execute(_contract_select(" WHERE "+" AND ".join(where)),params)
                rows=rows_to_list(cur)
                scope_metrics={}
                selected_team=_int(d.get('_team_scope'))
                scope_team_ids=None
                if selected_team:
                    scope_team_ids=[selected_team]
                if scope_team_ids and rows:
                    contract_ids=[int(x['id']) for x in rows]
                    team_marks=','.join('?' for _ in scope_team_ids)
                    contract_marks=','.join('?' for _ in contract_ids)
                    cur.execute("""SELECT cpt.contract_id,
                        SUM(cpt.allocation_amount) AS allocated_amount,
                        COUNT(*) AS link_count,
                        SUM(CASE WHEN cpt.allocation_amount IS NOT NULL
                          THEN 1 ELSE 0 END) AS allocated_count
                        FROM ContractProjectTeams cpt
                        JOIN ProjectTeams ptm ON ptm.id=cpt.project_team_id
                        WHERE cpt.is_active=1 AND ptm.is_active=1
                          AND ptm.team_id IN (%s)
                          AND cpt.contract_id IN (%s)
                        GROUP BY cpt.contract_id"""%(
                            team_marks,contract_marks),
                            *(scope_team_ids+contract_ids))
                    for metric in rows_to_list(cur):
                        scope_metrics[int(metric['contract_id'])]=metric
                    cur.execute("""SELECT s.contract_id,
                        SUM(COALESCE(s.confirmed_without_vat,
                          s.confirmed_price,s.requested_without_vat,
                          s.requested_price,0)
                          *cst.allocation_percent/100) AS approved_amount
                        FROM ContractStatementTeams cst
                        JOIN ContractStatements s ON s.id=cst.statement_id
                        JOIN ProjectTeams ptm ON ptm.id=cst.project_team_id
                        WHERE cst.is_active=1 AND s.is_active=1
                          AND s.is_current=1
                          AND (s.business_status='employer_approved' OR (s.employer_decision_at IS NOT NULL AND ISNULL(s.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(s.confirmed_without_vat,s.confirmed_price,0)>0))
                          AND ptm.team_id IN (%s)
                          AND s.contract_id IN (%s)
                        GROUP BY s.contract_id"""%(
                            team_marks,contract_marks),
                            *(scope_team_ids+contract_ids))
                    for metric in rows_to_list(cur):
                        scope_metrics.setdefault(
                            int(metric['contract_id']),{}
                        )['approved_amount']=metric.get('approved_amount')
            today=_dt.date.today();recommendations=[];alerts=[]
            for row in rows:
                if row.get('is_closed'):
                    continue
                company_effective=_decimal.Decimal(str(row.get('effective_price') or 0))
                company_approved=_decimal.Decimal(str(row.get('approved_statement_total') or 0))
                company_remaining=max(_decimal.Decimal(0),company_effective-company_approved)
                metric=scope_metrics.get(int(row['id'])) if scope_team_ids else None
                allocation_complete=bool(
                    metric and int(metric.get('link_count') or 0)>0 and
                    int(metric.get('link_count') or 0)==
                    int(metric.get('allocated_count') or 0)
                )
                if scope_team_ids and allocation_complete:
                    effective=_decimal.Decimal(str(metric.get('allocated_amount') or 0))
                    approved=_decimal.Decimal(str(metric.get('approved_amount') or 0))
                    remaining=max(_decimal.Decimal(0),effective-approved)
                elif scope_team_ids:
                    remaining=None
                else:
                    remaining=company_remaining
                end=_date(row.get('effective_end_date'));days=(end-today).days if end else None
                months=max(_decimal.Decimal(1),_decimal.Decimal(str(max(days or 30,1)))/_decimal.Decimal(30))
                burn=(remaining/months) if remaining is not None else None
                item=dict(row);item['remaining_price']=remaining
                item['company_remaining_price']=company_remaining
                item['requires_team_allocation']=bool(
                    scope_team_ids and not allocation_complete
                )
                item['days_to_end']=days;item['required_monthly_burn']=burn
                if days is not None and days<=45:alerts.append(item)
                if remaining is None or remaining>0:recommendations.append(item)
            recommendations.sort(key=lambda x:(-(x.get('payer_rating') or 0),-(x.get('ready_task_count') or 0),x.get('days_to_end') if x.get('days_to_end') is not None else 10**9,bool(x.get('requires_team_allocation')),-float(x.get('remaining_price') or 0)))
            alerts.sort(key=lambda x:(x.get('days_to_end') if x.get('days_to_end') is not None else 10**9,bool(x.get('requires_team_allocation')),-float(x.get('remaining_price') or 0)))
            return _ok(rows=recommendations[:30],critical_alerts=alerts[:30])
        except Exception as exc:return _err(str(exc),rows=[],critical_alerts=[])

    # ── Attachments ────────────────────────────────────────────────
    storage_dir=os.path.join(base_dir,'data','files');thumb_dir=os.path.join(storage_dir,'thumbs')
    os.makedirs(storage_dir,exist_ok=True);os.makedirs(thumb_dir,exist_ok=True)
    allowed={'.jpg','image/jpeg','.jpeg','.png','image/png','.webp','image/webp','.pdf','application/pdf','.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document','.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','.xls','application/vnd.ms-excel','.doc','application/msword'}

    def _can_access_entity(cur, entity_type, entity_id, required_permission='attachments.view'):
        required = required_permission
        if not user_has_permission(g.user, required):
            return False
        if has_company_scope(g.user):
            return True
        uid=g.user['id']
        if entity_type=='task':
            cur.execute("""SELECT TOP 1 1 FROM Tasks t JOIN ProjectTeams pt
                ON pt.id=t.project_team_id JOIN TeamMembers tm ON tm.team_id=pt.team_id
                WHERE t.id=? AND pt.is_active=1 AND tm.user_id=? AND tm.is_active=1""",
                        entity_id,uid)
            return cur.fetchone() is not None
        if entity_type=='extension':
            if not user_has_permission(g.user,'extensions.view'):return False
            cur.execute("""SELECT TOP 1 1 FROM ContractExtensions e
                JOIN ContractProjectTeams cpt ON cpt.contract_id=e.contract_id AND cpt.is_active=1
                JOIN ProjectTeams pt ON pt.id=cpt.project_team_id AND pt.is_active=1
                JOIN TeamMembers tm ON tm.team_id=pt.team_id
                WHERE e.id=? AND tm.user_id=? AND tm.is_active=1""",entity_id,uid)
            return cur.fetchone() is not None
        if entity_type=='statement':
            if not user_has_permission(g.user,'statements.view'):return False
            cur.execute("""SELECT TOP 1 1 FROM ContractStatementTeams cst
                JOIN ProjectTeams pt ON pt.id=cst.project_team_id AND pt.is_active=1
                JOIN TeamMembers tm ON tm.team_id=pt.team_id
                WHERE cst.statement_id=? AND cst.is_active=1
                  AND tm.user_id=? AND tm.is_active=1""",entity_id,uid)
            return cur.fetchone() is not None
        return False

    @app.route('/api/attachments', methods=['POST'])
    @require_auth
    def api_v7_attachments():
        try:
            d=request.get_json() or {};et=d.get('entity_type');eid=_int(d.get('entity_id'))
            with db_lock:
                c=get_conn();cur=c.cursor()
                if not _can_access_entity(cur,et,eid,'attachments.view'):return _err('دسترسی به فایل‌ها مجاز نیست',403)
                cur.execute("""SELECT a.id,a.entity_type,a.entity_id,a.original_name,a.caption,a.sort_order,a.version_no,a.uploaded_at,
                    b.mime_type,b.size_bytes,b.width,b.height,CASE WHEN b.thumbnail_name IS NULL THEN 0 ELSE 1 END AS has_thumbnail
                    FROM Attachments a JOIN FileBlobs b ON b.id=a.blob_id WHERE a.entity_type=? AND a.entity_id=? AND a.is_active=1
                    ORDER BY a.sort_order,a.id""",et,eid);rows=rows_to_list(cur)
            return _ok(rows=rows)
        except Exception as exc:return _err(str(exc),rows=[])

    @app.route('/api/attachment_upload', methods=['POST'])
    @require_auth
    def api_v7_attachment_upload():
        try:
            if os.path.exists(os.path.join(base_dir, 'data', 'backup.lock')):
                return _err('پشتیبان‌گیری در حال انجام است؛ چند دقیقه دیگر دوباره تلاش کنید', 423)
            et=request.form.get('entity_type');eid=_int(request.form.get('entity_id'));caption=(request.form.get('caption') or '').strip() or None;f=request.files.get('file')
            if not f or not et or not eid:return _err('فایل و مقصد الزامی است')
            raw=f.read(30*1024*1024+1)
            if len(raw)>30*1024*1024:return _err('حداکثر حجم هر فایل ۳۰ مگابایت است')
            ext=os.path.splitext(f.filename or '')[1].lower();mime=(f.mimetype or mimetypes.guess_type(f.filename or '')[0] or 'application/octet-stream').lower()
            if ext not in allowed or mime not in allowed:return _err('نوع فایل مجاز نیست')
            width=height=None;thumb_bytes=None;storage_ext=ext
            if ext in ('.jpg','.jpeg','.png','.webp'):
                from PIL import Image,ImageOps
                if not mime.startswith('image/'):return _err('پسوند و نوع تصویر با هم سازگار نیستند')
                try:
                    source=Image.open(io.BytesIO(raw))
                    if source.width*source.height>40_000_000:return _err('ابعاد تصویر بیش از حد مجاز است')
                    image=ImageOps.exif_transpose(source).convert('RGB')
                except Exception:
                    return _err('ساختار فایل تصویر معتبر نیست')
                image.thumbnail((1600,1600));width,height=image.size
                out=io.BytesIO();image.save(out,'WEBP',quality=80,method=6);raw=out.getvalue();mime='image/webp';storage_ext='.webp'
                th=image.copy();th.thumbnail((360,240));tout=io.BytesIO();th.save(tout,'WEBP',quality=72,method=6);thumb_bytes=tout.getvalue()
            elif ext=='.pdf' and not raw.startswith(b'%PDF'):
                return _err('ساختار فایل PDF معتبر نیست')
            elif ext in ('.docx','.xlsx') and not raw.startswith(b'PK'):
                return _err('ساختار فایل Office معتبر نیست')
            elif ext in ('.doc','.xls') and not raw.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
                return _err('ساختار فایل Office قدیمی معتبر نیست')
            sha=hashlib.sha256(raw).hexdigest();storage_name=sha+storage_ext;thumb_name=(sha+'_thumb.webp') if thumb_bytes else None
            with db_lock:
                c=get_conn();cur=c.cursor()
                if not _can_access_entity(cur,et,eid,'attachments.upload'):return _err('اجازه بارگذاری فایل ندارید',403)
                blob=_one(cur,"SELECT id FROM FileBlobs WHERE sha256=?",sha)
                if blob:bid=blob['id']
                else:
                    with open(os.path.join(storage_dir,storage_name),'wb') as out:out.write(raw)
                    if thumb_bytes:
                        with open(os.path.join(thumb_dir,thumb_name),'wb') as out:out.write(thumb_bytes)
                    cur.execute("INSERT INTO FileBlobs(sha256,mime_type,storage_name,thumbnail_name,size_bytes,width,height) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?)",sha,mime,storage_name,thumb_name,len(raw),width,height);bid=cur.fetchone()[0]
                cur.execute("SELECT ISNULL(MAX(version_no),0)+1 FROM Attachments WHERE entity_type=? AND entity_id=? AND original_name=?",et,eid,os.path.basename(f.filename or 'file'));ver=cur.fetchone()[0]
                cur.execute("INSERT INTO Attachments(entity_type,entity_id,blob_id,original_name,caption,version_no,uploaded_by) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?)",et,eid,bid,os.path.basename(f.filename or 'file')[:260],caption,ver,g.user['id']);aid=cur.fetchone()[0]
                _entity_audit(cur,'attachment',aid,'upload',None,{'entity_type':et,'entity_id':eid,'sha256':sha});c.commit()
            return _ok(id=aid,size=len(raw),deduplicated=bool(blob))
        except Exception as exc:return _err(str(exc))

    @app.route('/api/attachment_download/<int:attachment_id>', methods=['GET'])
    @require_auth
    def api_v7_attachment_download(attachment_id):
        try:
            thumb=request.args.get('thumb')=='1'
            with db_lock:
                c=get_conn();cur=c.cursor();row=_one(cur,"""SELECT a.entity_type,a.entity_id,a.original_name,b.storage_name,b.thumbnail_name,b.mime_type
                    FROM Attachments a JOIN FileBlobs b ON b.id=a.blob_id WHERE a.id=? AND a.is_active=1""",attachment_id)
                if not row:return _err('فایل یافت نشد',404)
                if not _can_access_entity(cur,row['entity_type'],row['entity_id'],'attachments.view'):return _err('دسترسی به فایل مجاز نیست',403)
                cur.execute("INSERT INTO FileDownloadLog(attachment_id,user_id,ip) VALUES(?,?,?)",attachment_id,g.user['id'],(request.remote_addr or '')[:45]);c.commit()
            name=row.get('thumbnail_name') if thumb and row.get('thumbnail_name') else row.get('storage_name');path=os.path.join(thumb_dir if thumb and row.get('thumbnail_name') else storage_dir,name)
            if not os.path.isfile(path):return _err('فایل فیزیکی یافت نشد',404)
            return send_file(path,mimetype='image/webp' if thumb else row.get('mime_type'),as_attachment=not thumb,download_name=row.get('original_name'))
        except Exception as exc:return _err(str(exc))

    @app.route('/api/attachment_archive', methods=['POST'])
    @require_auth
    def api_v7_attachment_archive():
        try:
            aid=_int((request.get_json() or {}).get('id'))
            with db_lock:
                c=get_conn();cur=c.cursor();row=_one(cur,"SELECT * FROM Attachments WHERE id=? AND is_active=1",aid)
                if not row:return _err('فایل یافت نشد')
                if not _can_access_entity(cur,row['entity_type'],row['entity_id'],'attachments.archive'):return _err('اجازه بایگانی فایل ندارید',403)
                cur.execute("UPDATE Attachments SET is_active=0,deleted_by=?,deleted_at=GETDATE() WHERE id=?",g.user['id'],aid);_entity_audit(cur,'attachment',aid,'archive',row,None);c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/attachment_update', methods=['POST'])
    @require_auth
    def api_v7_attachment_update():
        try:
            d=request.get_json() or {};aid=_int(d.get('id'));move=_int(d.get('move'),0)
            with db_lock:
                c=get_conn();cur=c.cursor();row=_one(cur,"SELECT * FROM Attachments WHERE id=? AND is_active=1",aid)
                if not row:return _err('فایل یافت نشد')
                if not _can_access_entity(cur,row['entity_type'],row['entity_id'],'attachments.edit'):return _err('اجازه ویرایش فایل ندارید',403)
                caption=(d.get('caption') if 'caption' in d else row.get('caption'))
                sort_order=max(0,_int(row.get('sort_order'),0)+move)
                cur.execute("UPDATE Attachments SET caption=?,sort_order=? WHERE id=?",(caption or '').strip() or None,sort_order,aid)
                _entity_audit(cur,'attachment',aid,'metadata_update',row,{'caption':caption,'sort_order':sort_order});c.commit()
            return _ok()
        except Exception as exc:return _err(str(exc))

    @app.route('/api/storage_stats', methods=['POST'])
    @require_roles('admin')
    def api_v7_storage_stats():
        try:
            with db_lock:
                c=get_conn();cur=c.cursor();stats=_one(cur,"SELECT COUNT(*) AS blob_count,ISNULL(SUM(size_bytes),0) AS physical_bytes FROM FileBlobs") or {}
                links=_one(cur,"SELECT COUNT(*) AS attachment_count FROM Attachments WHERE is_active=1") or {};stats.update(links)
                chat=_one(cur,"""SELECT COUNT(*) AS encrypted_chat_file_count,
                    ISNULL(SUM(size_bytes),0) AS encrypted_chat_bytes
                    FROM ChatEncryptedFiles""") or {}
                stats.update(chat)
                stats['total_file_bytes']=(
                    _int(stats.get('physical_bytes'),0)+
                    _int(stats.get('encrypted_chat_bytes'),0)
                )
            return _ok(**stats)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/backups', methods=['POST'])
    @require_roles('admin')
    def api_v7_backups():
        try:
            import backup as backup_module
            rows=[]
            for path in reversed(backup_module._archives()):
                try:
                    info=backup_module.manifest(path) if hasattr(backup_module,'manifest') else None
                except Exception:
                    info=None
                # backup.py intentionally keeps manifest parsing private; read
                # the tiny JSON entry without extracting any archive content.
                if info is None:
                    import zipfile
                    with zipfile.ZipFile(path) as z:
                        info=json.loads(z.read('manifest.json').decode('utf-8'))
                rows.append({'name':os.path.basename(path),'size_bytes':os.path.getsize(path),
                             'kind':info.get('kind'),'created_at':info.get('created_at'),
                             'parent_full':info.get('parent_full')})
            return _ok(rows=rows)
        except Exception as exc:return _err(str(exc),rows=[])

    @app.route('/api/backup_run', methods=['POST'])
    @require_roles('admin')
    def api_v7_backup_run():
        try:
            import backup as backup_module
            full=bool((request.get_json() or {}).get('full'))
            path=backup_module.run_backup(force_full=full)
            if not path:return _err('ساخت پشتیبان ناموفق بود؛ فایل backup.log را بررسی کنید')
            audit('backup_create',os.path.basename(path))
            return _ok(name=os.path.basename(path))
        except Exception as exc:return _err(str(exc))

    @app.route('/api/backup_verify', methods=['POST'])
    @require_roles('admin')
    def api_v7_backup_verify():
        try:
            import backup as backup_module
            name=os.path.basename((request.get_json() or {}).get('name') or '')
            path=os.path.join(backup_module.BACKUP_DIR,name)
            if not name or not os.path.isfile(path):return _err('فایل پشتیبان یافت نشد')
            ok,errors=backup_module.verify_archive(path)
            audit('backup_verify','%s: %s'%(name,'ok' if ok else 'failed'))
            return _ok(valid=ok,errors=errors)
        except Exception as exc:return _err(str(exc))

    @app.route('/api/backup_download/<path:filename>', methods=['GET'])
    @require_roles('admin')
    def api_v7_backup_download(filename):
        try:
            import backup as backup_module
            name=os.path.basename(filename)
            if name!=filename or not name.lower().endswith('.taskhubbackup'):return _err('نام فایل نامعتبر است',400)
            path=os.path.join(backup_module.BACKUP_DIR,name)
            if not os.path.isfile(path):return _err('فایل یافت نشد',404)
            audit('backup_download',name)
            return send_file(path,as_attachment=True,download_name=name,mimetype='application/zip')
        except Exception as exc:return _err(str(exc))

    @app.route('/api/v7_export', methods=['POST'])
    @require_auth
    def api_v7_export():
        """Role-aware Excel/PDF export for the new financial domain."""
        try:
            d=request.get_json() or {};kind=d.get('kind') or 'contracts';fmt=d.get('format') or 'xlsx';cid=_int(d.get('contract_id'))
            if not user_has_permission(g.user, 'reports.export'):
                return _err('اجازه دریافت خروجی را ندارید',403)
            if kind != 'work_share' and not user_has_permission(g.user, 'reports.financial'):
                return _err('اجازه دریافت گزارش مالی را ندارید',403)
            with db_lock:
                c=get_conn();cur=c.cursor()
                if kind=='contracts':
                    where=["co.is_active=1"];params=[]
                    scope,scope_params=_team_scope_condition(
                        g.user,d,'contract','co'
                    )
                    if scope:where.append(scope);params.extend(scope_params)
                    cur.execute(_contract_select(" WHERE "+" AND ".join(where)),params)
                    rows=rows_to_list(cur)
                elif kind=='extensions':
                    sql="""SELECT e.*,co.title AS contract_title,co.contract_number,et.name AS extension_type_name,
                        p.name AS project_name,ci.name AS city_name
                        FROM ContractExtensions e JOIN Contracts co ON co.id=e.contract_id
                        LEFT JOIN ContractExtensionTypes et ON et.id=e.extension_type_id
                        LEFT JOIN Projects p ON p.id=co.project_id LEFT JOIN Cities ci ON ci.id=p.city_id
                        WHERE e.is_active=1"""
                    params=[]
                    if cid:sql+=" AND e.contract_id=?";params.append(cid)
                    scope,scope_params=_team_scope_condition(
                        g.user,d,'contract','co'
                    )
                    if scope:sql+=" AND "+scope;params.extend(scope_params)
                    cur.execute(sql,params)
                    rows=rows_to_list(cur)
                elif kind=='statements':
                    sql="""SELECT s.*,co.title AS contract_title,co.contract_number,st.name AS statement_type_name,
                        p.name AS project_name,ci.name AS city_name,
                        COALESCE(s.confirmed_without_vat,s.confirmed_price,s.requested_without_vat,s.requested_price,0) AS effective_amount
                        FROM ContractStatements s JOIN Contracts co ON co.id=s.contract_id
                        LEFT JOIN ContractStatementTypes st ON st.id=s.statement_type_id
                        LEFT JOIN Projects p ON p.id=co.project_id LEFT JOIN Cities ci ON ci.id=p.city_id
                        WHERE s.is_active=1"""
                    params=[]
                    if cid:sql+=" AND s.contract_id=?";params.append(cid)
                    scope,scope_params=_team_scope_condition(
                        g.user,d,'statement','s'
                    )
                    if scope:sql+=" AND "+scope;params.extend(scope_params)
                    cur.execute(sql,params)
                    rows=rows_to_list(cur)
                elif kind=='plans':
                    selected_team=_int(d.get('_team_scope'))
                    team_view=bool(selected_team)
                    if team_view:
                        where=["fp.is_current=1","tf.is_current=1","t.is_active=1"]
                        params=[]
                        if selected_team:
                            where.append("t.id=?");params.append(selected_team)
                        cur.execute("""SELECT fp.jalali_year,fp.title,fp.approved_target,
                            tf.annual_target,tf.version_no,fp.status,t.name AS team_name,
                            pp.month_no,pp.target_amount,pp.allocation_mode,
                            pp.allocation_percent,pp.is_locked,
                            ISNULL(a.sent_amount,0) AS sent_amount,
                            ISNULL(a.approved_amount,0) AS approved_amount
                            FROM FinancialPlans fp JOIN TeamFinancialTargets tf
                              ON tf.plan_id=fp.id JOIN Teams t ON t.id=tf.team_id
                            JOIN TeamFinancialTargetPeriods pp ON pp.target_id=tf.id
                            OUTER APPLY(SELECT
                              SUM(CASE WHEN s.business_status IN
                                ('sent','employer_approved','employer_rejected')
                                THEN COALESCE(s.requested_without_vat,s.requested_price,0)
                                  *cst.allocation_percent/100
                                ELSE 0 END) AS sent_amount,
                              SUM(CASE WHEN (s.business_status='employer_approved' OR (s.employer_decision_at IS NOT NULL AND ISNULL(s.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(s.confirmed_without_vat,s.confirmed_price,0)>0))
                                THEN COALESCE(s.confirmed_without_vat,s.confirmed_price,
                                  s.requested_without_vat,s.requested_price,0)
                                  *cst.allocation_percent/100
                                ELSE 0 END) AS approved_amount
                              FROM ContractStatementTeams cst
                              JOIN ContractStatements s ON s.id=cst.statement_id
                              JOIN ProjectTeams ptm
                                ON ptm.id=cst.project_team_id
                              WHERE ptm.team_id=t.id AND s.is_active=1
                                AND cst.is_active=1 AND s.is_current=1
                                AND s.period_year=fp.jalali_year
                                AND s.period_month=pp.month_no) a
                            WHERE """+" AND ".join(where)+
                            " ORDER BY fp.jalali_year,t.name,pp.month_no",params)
                    else:
                        cur.execute("""SELECT fp.jalali_year,fp.title,fp.approved_target,
                            fp.annual_target,fp.version_no,fp.status,
                            CAST(NULL AS NVARCHAR(160)) AS team_name,
                            pp.month_no,pp.target_amount,pp.allocation_mode,
                            pp.allocation_percent,pp.is_locked,
                            ISNULL(a.sent_amount,0) AS sent_amount,
                            ISNULL(a.approved_amount,0) AS approved_amount
                            FROM FinancialPlans fp JOIN FinancialPlanPeriods pp
                              ON pp.plan_id=fp.id
                            OUTER APPLY(SELECT
                              SUM(CASE WHEN s.business_status IN
                                ('sent','employer_approved','employer_rejected')
                                THEN COALESCE(s.requested_without_vat,s.requested_price,0)
                                ELSE 0 END) AS sent_amount,
                              SUM(CASE WHEN (s.business_status='employer_approved' OR (s.employer_decision_at IS NOT NULL AND ISNULL(s.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(s.confirmed_without_vat,s.confirmed_price,0)>0))
                                THEN COALESCE(s.confirmed_without_vat,s.confirmed_price,
                                  s.requested_without_vat,s.requested_price,0)
                                ELSE 0 END) AS approved_amount
                              FROM ContractStatements s WHERE s.is_active=1
                                AND s.is_current=1 AND s.period_year=fp.jalali_year
                                AND s.period_month=pp.month_no) a
                            WHERE fp.is_current=1
                            ORDER BY fp.jalali_year,pp.month_no""")
                    rows=rows_to_list(cur)
                elif kind=='attachments':
                    sql="""SELECT a.entity_type,a.entity_id,a.original_name,a.caption,a.version_no,a.uploaded_at,
                        b.mime_type,b.size_bytes,b.sha256,u.display_name AS uploaded_by_name
                        FROM Attachments a JOIN FileBlobs b ON b.id=a.blob_id LEFT JOIN Users u ON u.id=a.uploaded_by
                        WHERE a.is_active=1"""
                    params=[]
                    task_scope,task_params=_team_scope_condition(
                        g.user,d,'task','scope_task'
                    )
                    statement_scope,statement_params=_team_scope_condition(
                        g.user,d,'statement','scope_statement'
                    )
                    contract_scope,contract_params=_team_scope_condition(
                        g.user,d,'contract','scope_contract'
                    )
                    if task_scope or statement_scope or contract_scope:
                        sql+=""" AND (
                          (a.entity_type='task' AND EXISTS(
                            SELECT 1 FROM Tasks scope_task WHERE scope_task.id=a.entity_id
                              AND %s))
                          OR (a.entity_type='statement' AND EXISTS(
                            SELECT 1 FROM ContractStatements scope_statement
                            WHERE scope_statement.id=a.entity_id AND %s))
                          OR (a.entity_type='extension' AND EXISTS(
                            SELECT 1 FROM ContractExtensions extension
                            JOIN Contracts scope_contract
                              ON scope_contract.id=extension.contract_id
                            WHERE extension.id=a.entity_id AND %s))
                        )""" % (task_scope or "1=1",
                                 statement_scope or "1=1",
                                 contract_scope or "1=1")
                        params.extend(task_params)
                        params.extend(statement_params)
                        params.extend(contract_params)
                    sql+=" ORDER BY a.entity_type,a.entity_id,a.id"
                    cur.execute(sql,params);rows=rows_to_list(cur)
                elif kind=='project_revenue':
                    report_data=_financial_report_data(cur,d,g.user)
                    dimension=d.get('dimension') or 'contracts'
                    rows=report_data.get({'city':'cities','contract_type':'contract_types','project_type':'types','project':'projects'}.get(dimension,'contracts'),[])
                elif kind=='work_share':
                    report_data=_work_report_data(cur,d,g.user);rows=report_data['rows']
                else:return _err('نوع گزارش نامعتبر است')
            status_fa={'draft':'پیش‌نویس','pending':'در انتظار مدیر','approved':'تایید مدیر','rejected':'رد مدیر',
                       'pending_internal':'در انتظار مدیر','internal_approved':'تایید داخلی','internal_rejected':'رد داخلی',
                       'sent':'ارسال به کارفرما','employer_approved':'تایید کارفرما','employer_rejected':'رد کارفرما',
                       'revised':'نسخه قبلی','void':'باطل','planned':'برنامه‌ریزی‌شده'}
            report_rows=[]
            if kind=='contracts':
                for r in rows:
                    remaining=contract_remaining(r.get('effective_price'),[{'confirmed_without_vat':r.get('approved_statement_total')}])
                    report_rows.append({'شهر':r.get('city_name') or 'بدون شهر','پروژه':r.get('project_name') or 'بدون پروژه',
                        'عنوان قرارداد':r.get('title'),'شماره قرارداد':r.get('contract_number'),'کارفرما':r.get('employer_name'),
                        'تاریخ پایان مؤثر':r.get('effective_end_date_fa') or r.get('effective_end_date'),
                        'مبلغ مؤثر':r.get('effective_price'),'صورت‌وضعیت تاییدشده':r.get('approved_statement_total'),
                        'مانده قرارداد':remaining,'امتیاز خوش‌حسابی':r.get('payer_rating'),'وضعیت':r.get('contract_status_name')})
            elif kind=='extensions':
                for r in rows:report_rows.append({'شهر':r.get('city_name') or 'بدون شهر','پروژه':r.get('project_name') or 'بدون پروژه',
                    'قرارداد':r.get('contract_title'),'شماره قرارداد':r.get('contract_number'),'عنوان الحاقیه':r.get('title'),
                    'شماره الحاقیه':r.get('extension_number'),'نوع':r.get('extension_type_name'),'تاریخ':r.get('extension_date_fa') or r.get('extension_date'),
                    'پایان جدید':r.get('end_date_fa') or r.get('end_date'),'اثر ریالی':r.get('price_delta'),
                    'مبلغ نهایی ثبت‌شده':r.get('resulting_price'),'وضعیت':status_fa.get(r.get('internal_status'),r.get('internal_status'))})
            elif kind=='statements':
                for r in rows:report_rows.append({'شهر':r.get('city_name') or 'بدون شهر','پروژه':r.get('project_name') or 'بدون پروژه',
                    'قرارداد':r.get('contract_title'),'شماره قرارداد':r.get('contract_number'),'عنوان صورت‌وضعیت':r.get('title'),
                    'نوع':r.get('statement_type_name'),'شماره':r.get('statement_number'),'تاریخ':r.get('statement_date_fa') or r.get('statement_date'),
                    'سال':r.get('period_year'),'ماه':r.get('period_month'),'مبلغ ارسالی بدون مالیات':r.get('requested_without_vat'),
                    'مبلغ تاییدشده بدون مالیات':r.get('confirmed_without_vat'),'وضعیت':status_fa.get(r.get('business_status'),r.get('business_status'))})
            elif kind=='plans':
                for r in rows:
                    target=_dec(r.get('target_amount'));approved=_dec(r.get('approved_amount'))
                    report_rows.append({'سال':r.get('jalali_year'),'تیم':r.get('team_name') or 'کل شرکت','عنوان برنامه':r.get('title'),'هدف سالانه':r.get('annual_target'),
                        'هدف مصوب شرکت':r.get('approved_target'),
                        'نسخه':r.get('version_no'),'ماه':r.get('month_no'),'هدف ماه':r.get('target_amount'),
                        'روش تخصیص':r.get('allocation_mode'),'درصد':r.get('allocation_percent'),
                        'ارسال‌شده':r.get('sent_amount'),'تایید کارفرما':r.get('approved_amount'),
                        'درصد تحقق ماه':(int(approved*100/target) if target else 0),
                        'فاصله تا هدف ماه':target-approved,'ماه بسته':r.get('is_locked')})
            elif kind=='project_revenue':
                dimension=d.get('dimension') or 'contracts'
                if dimension=='city':
                    for r in rows:report_rows.append({'شهر':r.get('city_name') or 'بدون شهر','تعداد پروژه':r.get('project_count'),
                        'تعداد قرارداد':r.get('contract_count'),'مبلغ برنامه':r.get('planned_amount'),
                        'مبلغ ارسال‌شده بدون مالیات':r.get('sent_amount'),'درآمد تاییدشده بدون مالیات':r.get('approved_amount'),
                        'سهم از کل درآمد (درصد)':r.get('revenue_share_percent'),'تحقق برنامه (درصد)':r.get('achievement_percent')})
                elif dimension=='contract_type':
                    for r in rows:report_rows.append({'نوع قرارداد':r.get('contract_type_name') or 'نوع تعیین نشده','تعداد پروژه':r.get('project_count'),
                        'تعداد قرارداد':r.get('contract_count'),'مبلغ برنامه':r.get('planned_amount'),
                        'مبلغ ارسال‌شده بدون مالیات':r.get('sent_amount'),'درآمد تاییدشده بدون مالیات':r.get('approved_amount'),
                        'سهم از کل درآمد (درصد)':r.get('revenue_share_percent'),'تحقق برنامه (درصد)':r.get('achievement_percent')})
                elif dimension=='project_type':
                    for r in rows:report_rows.append({'نوع پروژه':r.get('project_type_name'),'تعداد پروژه':r.get('project_count'),
                        'تعداد قرارداد':r.get('contract_count'),'مبلغ برنامه':r.get('planned_amount'),
                        'مبلغ ارسال‌شده بدون مالیات':r.get('sent_amount'),'درآمد تاییدشده بدون مالیات':r.get('approved_amount'),
                        'سهم از کل درآمد (درصد)':r.get('revenue_share_percent'),'تحقق برنامه (درصد)':r.get('achievement_percent')})
                else:
                    for r in rows:report_rows.append({'نوع قرارداد':r.get('contract_type_name') or 'نوع تعیین نشده','نوع پروژه':r.get('project_type_name'),'شهر':r.get('city_name') or 'بدون شهر',
                        'پروژه':r.get('project_name') or 'بدون پروژه','قرارداد':r.get('title'),
                        'شماره قرارداد':r.get('contract_number'),'مبلغ مؤثر قرارداد':r.get('effective_price'),
                        'مبلغ برنامه':r.get('planned_amount'),'مبلغ ارسال‌شده بدون مالیات':r.get('sent_amount'),
                        'درآمد تاییدشده بدون مالیات':r.get('approved_amount'),'مانده قرارداد':r.get('remaining_price'),
                        'سهم از کل درآمد (درصد)':r.get('revenue_share_percent'),
                        'تحقق برنامه (درصد)':r.get('achievement_percent')})
            elif kind=='work_share':
                for r in rows:report_rows.append({'همکار':r.get('display_name'),'نوع پروژه':r.get('project_type_name'),
                    'شهر':r.get('city_name'),'پروژه':r.get('project_name'),'تعداد تسک':r.get('task_count'),
                    'زمان کارکرد (ثانیه)':r.get('seconds'),'سهم فرد از کار پروژه (درصد)':r.get('project_share_percent'),
                    'تمرکز کار فرد روی پروژه (درصد)':r.get('person_focus_percent'),
                    'سهم از کل کارکرد (درصد)':r.get('overall_percent')})
            else:
                for r in rows:report_rows.append({'نوع رکورد':r.get('entity_type'),'شناسه رکورد':r.get('entity_id'),
                    'نام فایل':r.get('original_name'),'توضیح':r.get('caption'),'نسخه':r.get('version_no'),
                    'تاریخ بارگذاری':r.get('uploaded_at'),'نوع فایل':r.get('mime_type'),'حجم بایت':r.get('size_bytes'),
                    'SHA-256':r.get('sha256'),'بارگذار':r.get('uploaded_by_name')})
            rows=report_rows
            if fmt=='pdf':
                from reportlab.lib import colors
                from reportlab.lib.pagesizes import A4,landscape
                from reportlab.lib.styles import ParagraphStyle
                from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph
                import arabic_reshaper
                from bidi.algorithm import get_display
                fd,path=tempfile.mkstemp(suffix='.pdf');os.close(fd);font_name=fa_font.register()
                def fa(x):return get_display(arabic_reshaper.reshape(str(x if x is not None else '')))
                cols=list(rows[0].keys())[:10] if rows else ['گزارش'];data=[[fa(x) for x in cols]]+[[fa(r.get(x,'')) for x in cols] for r in rows]
                doc=SimpleDocTemplate(path,pagesize=landscape(A4),rightMargin=18,leftMargin=18,topMargin=18,bottomMargin=18);tbl=Table(data,repeatRows=1)
                tbl.setStyle(TableStyle([('FONT',(0,0),(-1,-1),font_name,7),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#274690')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.25,colors.grey),('ALIGN',(0,0),(-1,-1),'RIGHT')]))
                title_style=ParagraphStyle('fa-title',fontName=font_name,fontSize=12,leading=18,alignment=2)
                doc.build([Paragraph(fa('گزارش تسک هاب نسخه ۸'),title_style),tbl]);mime='application/pdf';name='TaskHub_%s.pdf'%kind
            else:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment,Font,PatternFill
                wb=Workbook();ws=wb.active;ws.title={'contracts':'قراردادها','extensions':'الحاقیه‌ها','statements':'صورت وضعیت‌ها','plans':'برنامه مالی','attachments':'فایل‌ها','project_revenue':'درآمد پروژه‌ها','work_share':'سهم کارکرد'}[kind];ws.sheet_view.rightToLeft=True
                cols=list(rows[0].keys()) if rows else ['نتیجه'];ws.append(cols)
                for r in rows:ws.append([_jsonable(r.get(x)) for x in cols])
                for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='274690');cell.alignment=Alignment(horizontal='center')
                ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
                for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(35,max(12,max(len(str(x.value or '')) for x in col)+2))
                fd,path=tempfile.mkstemp(suffix='.xlsx');os.close(fd);wb.save(path);mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet';name='TaskHub_%s.xlsx'%kind
            @after_this_request
            def cleanup(response):
                try:os.remove(path)
                except Exception:pass
                return response
            return send_file(path,mimetype=mime,as_attachment=True,download_name=name)
        except Exception as exc:return _err(str(exc))
