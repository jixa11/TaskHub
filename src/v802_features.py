# -*- coding: utf-8 -*-
"""TaskHub 8.0.2 reporting and deployment hardening.

Adds contract-type/task-category weighting, analytical financial attribution,
PDF/Excel exports, and a one-time safe contract-type migration.  The module is
additive and receives the already configured database connection factory.
"""
from __future__ import annotations

import datetime as _dt
import decimal as _decimal
import io
import json
import os

from flask import g, jsonify, request, send_file
from api_errors import public_error

import fa_font
from rbac import user_has_permission
from team_scope import has_company_scope

from v802_domain import (allocate_approved_amount, split_person_task_score,
                         weighted_task_score)


FINANCIAL_REPORT_ROLES = ("admin", "manager", "planner", "finance", "reporter")
GLOBAL_FINANCIAL_ROLES = ("admin", "finance", "reporter")


def _int(value, default=None):
    if value in (None, ""):
        return default
    try:
        return int(str(value).translate(str.maketrans(
            "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"
        )))
    except Exception:
        return default


def _dec(value, default=None):
    if value in (None, ""):
        return default
    try:
        return _decimal.Decimal(str(value).replace(",", ""))
    except Exception:
        return default


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    try:
        return _dt.date.fromisoformat(str(value)[:10])
    except Exception:
        return None


def _jsonable(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, _decimal.Decimal):
        return str(value)
    if isinstance(value, (_dt.date, _dt.datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(x) for x in value]
    return str(value)


def _pct(value, total):
    value = _dec(value, _decimal.Decimal(0))
    total = _dec(total, _decimal.Decimal(0))
    if not total:
        return 0.0
    return float((value * _decimal.Decimal(100) / total).quantize(
        _decimal.Decimal("0.01"), rounding=_decimal.ROUND_HALF_UP
    ))


def init_v802_tables(conn):
    """Apply idempotent 8.0.2 schema and the one-time contract-type migration.

    The whole upgrade is transactional. A failed index, seed, or migration
    cannot leave the database marked as upgraded while only half of the schema
    exists.
    """
    cur = conn.cursor()
    try:
        sqls = [
            """IF OBJECT_ID('ContractTaskCategoryWeights','U') IS NULL CREATE TABLE ContractTaskCategoryWeights(
                contract_type_id INT NOT NULL REFERENCES ContractTypes(id),
                task_category_id INT NOT NULL REFERENCES TaskCategories(id),
                weight DECIMAL(9,4) NOT NULL CONSTRAINT DF_ContractTaskCategoryWeights_weight DEFAULT 1,
                is_active BIT NOT NULL CONSTRAINT DF_ContractTaskCategoryWeights_active DEFAULT 1,
                updated_by INT NULL,updated_at DATETIME NULL,
                CONSTRAINT PK_ContractTaskCategoryWeights PRIMARY KEY(contract_type_id,task_category_id),
                CONSTRAINT CK_ContractTaskCategoryWeights_weight CHECK(weight>0 AND weight<=1000))""",
            """IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_ContractTaskWeights_category'
                  AND object_id=OBJECT_ID('ContractTaskCategoryWeights'))
                  CREATE INDEX IX_ContractTaskWeights_category
                  ON ContractTaskCategoryWeights(task_category_id,is_active,contract_type_id) INCLUDE(weight)""",
            """IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE name='IX_Tasks_financial_attribution'
                  AND object_id=OBJECT_ID('Tasks'))
                  CREATE INDEX IX_Tasks_financial_attribution
                  ON Tasks(contract_id,status,category_id,project_team_id)
                  INCLUDE(progress_weight,staff_id,completed_at,created_at)""",
            """IF NOT EXISTS(SELECT 1 FROM ContractTypes WHERE id=0)
                  INSERT INTO ContractTypes(id,name,is_active) VALUES(0,N'نوع تعیین نشده',1)""",
            "UPDATE ContractTypes SET name=N'نوع تعیین نشده',is_active=1 WHERE id=0",
            """IF NOT EXISTS(SELECT 1 FROM ContractTypes WHERE id=1)
                  INSERT INTO ContractTypes(id,name,is_active) VALUES(1,N'فروش',1)
                ELSE UPDATE ContractTypes SET name=N'فروش',is_active=1 WHERE id=1""",
            """IF NOT EXISTS(SELECT 1 FROM ContractTypes WHERE id=2)
                  INSERT INTO ContractTypes(id,name,is_active) VALUES(2,N'توسعه',1)
                ELSE UPDATE ContractTypes SET name=N'توسعه',is_active=1 WHERE id=2""",
            """IF NOT EXISTS(SELECT 1 FROM ContractTypes WHERE id=3)
                  INSERT INTO ContractTypes(id,name,is_active) VALUES(3,N'پشتیبانی',1)
                ELSE UPDATE ContractTypes SET name=N'پشتیبانی',is_active=1 WHERE id=3""",
            "UPDATE ContractTypes SET is_active=0 WHERE id NOT IN (0,1,2,3)",
        ]
        for sql in sqls:
            cur.execute(sql)

        # One-time migration only. Existing government/private values are moved
        # to undefined exactly once; later edits and new contracts are never reset.
        cur.execute("SELECT setting_value FROM AppSettings WHERE setting_key='v802_contract_types_migrated'")
        migrated = cur.fetchone()
        if not migrated:
            cur.execute("UPDATE Contracts SET contract_type_id=0")
            cur.execute("""INSERT INTO AppSettings(setting_key,setting_value,updated_at)
                           VALUES('v802_contract_types_migrated','1',GETDATE())""")

        # Every current category receives a neutral weight for each supported
        # type, so the report works immediately and admin tuning is optional.
        cur.execute("""INSERT INTO ContractTaskCategoryWeights(contract_type_id,task_category_id,weight,is_active)
            SELECT ct.id,tc.id,CAST(1 AS DECIMAL(9,4)),1
            FROM ContractTypes ct CROSS JOIN TaskCategories tc
            WHERE ct.id IN (0,1,2,3)
              AND NOT EXISTS(SELECT 1 FROM ContractTaskCategoryWeights w
                WHERE w.contract_type_id=ct.id AND w.task_category_id=tc.id)""")
        conn.commit()
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass
        raise


def register_v802_routes(app, get_conn, db_lock, require_auth, require_roles,
                          rows_to_list, audit, base_dir):
    def ok(**kwargs):
        payload = {"ok": True}; payload.update(kwargs)
        return jsonify(_jsonable(payload))

    def err(message, status=200, **kwargs):
        payload = {"ok": False, "error": str(message)}; payload.update(kwargs)
        return jsonify(_jsonable(payload)), status

    def _team_ids(cur, user, data):
        requested = _int(data.get("_team_scope", data.get("team_id")))
        if requested is not None:
            if not has_company_scope(user):
                cur.execute("""SELECT 1 FROM TeamMembers WHERE team_id=? AND user_id=? AND is_active=1""",
                            requested, user["id"])
                if not cur.fetchone():
                    raise PermissionError("به این تیم دسترسی ندارید")
            return [requested]
        if has_company_scope(user):
            cur.execute("SELECT id FROM Teams WHERE is_active=1")
        else:
            cur.execute("""SELECT tm.team_id FROM TeamMembers tm JOIN Teams t ON t.id=tm.team_id
                WHERE tm.user_id=? AND tm.is_active=1 AND t.is_active=1""", user["id"])
        return [int(x[0]) for x in cur.fetchall()]

    def _load_attribution(cur, data, user):
        start, end = _date(data.get("from_date")), _date(data.get("to_date"))
        if start and end and end < start:
            raise ValueError("تاریخ پایان گزارش قبل از تاریخ شروع است")
        city_id = _int(data.get("city_id"))
        contract_type_id = _int(data.get("contract_type_id"))
        contract_id = _int(data.get("contract_id"))
        user_id_filter = _int(data.get("user_id"))
        tids = _team_ids(cur, user, data)
        scope_limited = not has_company_scope(user)

        # Team-limited roles with no active membership must never fall through
        # to an unrestricted query. Manager/planner are company-wide roles.
        if scope_limited and not tids:
            return {
                "totals": {"approved_amount": 0, "team_attributed_amount": 0,
                           "person_attributed_amount": 0, "unallocated_team_amount": 0,
                           "unallocated_person_amount": 0, "completed_task_count": 0},
                "contracts": [], "teams": [], "people": [], "cities": [],
                "contract_types": [], "quality": {}, "filters": data,
                "access_scope": {"limited": True, "team_ids": []},
            }

        statement_join = ["s.is_active=1", "s.is_current=1", "(s.business_status='employer_approved' OR (s.employer_decision_at IS NOT NULL AND ISNULL(s.business_status,'') NOT IN ('employer_rejected','revised','void') AND COALESCE(s.confirmed_without_vat,s.confirmed_price,0)>0))"]
        statement_params = []
        if start:
            statement_join.append("CAST(COALESCE(s.employer_decision_at,s.statement_date) AS DATE)>=?")
            statement_params.append(start)
        if end:
            statement_join.append("CAST(COALESCE(s.employer_decision_at,s.statement_date) AS DATE)<=?")
            statement_params.append(end)

        where = ["co.is_active=1"]
        params = list(statement_params)
        if city_id is not None:
            where.append("ISNULL(ci.id,-1)=?"); params.append(city_id)
        if contract_type_id is not None:
            where.append("ISNULL(co.contract_type_id,0)=?"); params.append(contract_type_id)
        if contract_id is not None:
            where.append("co.id=?"); params.append(contract_id)
        if tids:
            marks = ",".join("?" for _ in tids)
            where.append("""EXISTS(SELECT 1 FROM ContractProjectTeams vis_cpt
                JOIN ProjectTeams vis_pt ON vis_pt.id=vis_cpt.project_team_id
                WHERE vis_cpt.contract_id=co.id AND vis_cpt.is_active=1
                  AND vis_pt.is_active=1 AND vis_pt.team_id IN (%s))""" % marks)
            params.extend(tids)

        cur.execute("""SELECT co.id AS contract_id,co.title AS contract_title,
                co.contract_number,ISNULL(co.contract_type_id,0) AS contract_type_id,
                ISNULL(ct.name,N'نوع تعیین نشده') AS contract_type_name,
                p.id AS project_id,p.name AS project_name,ci.id AS city_id,
                ISNULL(ci.name,N'بدون شهر') AS city_name,
                ISNULL(SUM(COALESCE(s.confirmed_without_vat,s.confirmed_price,
                    s.requested_without_vat,s.requested_price,0)),0) AS approved_amount
            FROM Contracts co
            LEFT JOIN Projects p ON p.id=co.project_id
            LEFT JOIN Cities ci ON ci.id=p.city_id
            LEFT JOIN ContractTypes ct ON ct.id=ISNULL(co.contract_type_id,0)
            LEFT JOIN ContractStatements s ON s.contract_id=co.id AND %s
            WHERE %s
            GROUP BY co.id,co.title,co.contract_number,co.contract_type_id,ct.name,
                     p.id,p.name,ci.id,ci.name
            ORDER BY ci.name,ct.name,co.title""" % (
                " AND ".join(statement_join), " AND ".join(where)
            ), params)
        contracts = rows_to_list(cur)
        contract_ids = [int(x["contract_id"]) for x in contracts]
        if not contract_ids:
            return {
                "totals": {"approved_amount": 0, "team_attributed_amount": 0,
                           "person_attributed_amount": 0, "unallocated_team_amount": 0,
                           "unallocated_person_amount": 0, "completed_task_count": 0},
                "contracts": [], "teams": [], "people": [], "cities": [],
                "contract_types": [], "quality": {}, "filters": data,
                "access_scope": {"limited": scope_limited, "team_ids": tids},
            }

        marks = ",".join("?" for _ in contract_ids)
        task_params = list(contract_ids)
        task_date = ""
        # Financial attribution is cumulative up to the report end date. This
        # prevents a payment approved this month from losing tasks completed in
        # an earlier month of the same contract.
        if end:
            task_date = " AND CAST(COALESCE(t.completed_at,t.created_at) AS DATE)<=?"
            task_params.append(end)
        cur.execute("""SELECT t.id AS task_id,t.contract_id,t.category_id,
                t.project_team_id,pt.team_id,tm.name AS team_name,t.staff_id,
                ISNULL(t.progress_weight,1) AS progress_weight,
                ISNULL(w.weight,1) AS category_weight,
                CASE WHEN w.task_category_id IS NULL THEN 1 ELSE 0 END AS mapping_missing
            FROM Tasks t
            JOIN Contracts co ON co.id=t.contract_id
            LEFT JOIN ProjectTeams pt ON pt.id=t.project_team_id
            LEFT JOIN Teams tm ON tm.id=pt.team_id
            LEFT JOIN ContractTaskCategoryWeights w
              ON w.contract_type_id=ISNULL(co.contract_type_id,0)
             AND w.task_category_id=t.category_id AND w.is_active=1
            WHERE t.status='done' AND t.contract_id IN (%s)%s""" % (marks, task_date), task_params)
        tasks = rows_to_list(cur)
        task_ids = [int(x["task_id"]) for x in tasks]

        time_by_task = {}
        helpers_by_task = {}
        if task_ids:
            tmarks = ",".join("?" for _ in task_ids)
            cur.execute("""SELECT task_id,user_id,
                SUM(ISNULL(seconds,DATEDIFF(second,started_at,ISNULL(ended_at,GETDATE())))) AS seconds
                FROM TaskTimeLog WHERE task_id IN (%s) GROUP BY task_id,user_id""" % tmarks,
                        task_ids)
            for row in rows_to_list(cur):
                tid = int(row["task_id"]); uid = int(row["user_id"])
                secs = max(0, int(row.get("seconds") or 0))
                if secs:
                    time_by_task.setdefault(tid, {})[uid] = secs
            cur.execute("SELECT task_id,user_id FROM TaskAssignees WHERE task_id IN (%s)" % tmarks,
                        task_ids)
            for row in rows_to_list(cur):
                helpers_by_task.setdefault(int(row["task_id"]), set()).add(int(row["user_id"]))

        contracts_by_id = {int(x["contract_id"]): x for x in contracts}
        total_score = {}
        team_scores = {}
        person_scores = {}
        contract_task_ids = {}
        team_task_ids = {}
        person_task_ids = {}
        quality = {
            "completed_without_contract": 0,
            "completed_without_team": 0,
            "multi_assignee_without_time": 0,
            "mapping_missing": 0,
            "contracts_without_completed_task": 0,
        }

        for task in tasks:
            tid = int(task["task_id"]); cid = int(task["contract_id"])
            score = weighted_task_score(
                task.get("progress_weight"), task.get("category_weight")
            )
            if score <= 0:
                continue
            total_score[cid] = total_score.get(cid, _decimal.Decimal(0)) + score
            contract_task_ids.setdefault(cid, set()).add(tid)
            if int(task.get("mapping_missing") or 0):
                quality["mapping_missing"] += 1

            team_id = _int(task.get("team_id"))
            if team_id is None:
                quality["completed_without_team"] += 1
            else:
                key = (cid, team_id)
                team_scores[key] = team_scores.get(key, _decimal.Decimal(0)) + score
                team_task_ids.setdefault(key, set()).add(tid)

            assignees = set(helpers_by_task.get(tid, set()))
            if task.get("staff_id") is not None:
                assignees.add(int(task["staff_id"]))
            allocations, unresolved = split_person_task_score(
                score, time_by_task.get(tid, {}), assignees
            )
            for uid, part in allocations.items():
                pkey = (cid, team_id, uid)
                person_scores[pkey] = person_scores.get(pkey, _decimal.Decimal(0)) + part
                person_task_ids.setdefault(pkey, set()).add(tid)
            if unresolved:
                quality["multi_assignee_without_time"] += 1

        # Count completed tasks that cannot participate in any financial report.
        scope_clause = ""
        scope_params = []
        if scope_limited:
            smarks = ",".join("?" for _ in tids)
            scope_clause = """ AND EXISTS(
                SELECT 1 FROM ProjectTeams qp WHERE qp.id=t.project_team_id
                  AND qp.team_id IN (%s))""" % smarks
            scope_params = tids
        cur.execute("""SELECT COUNT(*) FROM Tasks t
            WHERE t.status='done' AND t.contract_id IS NULL%s""" % scope_clause, scope_params)
        quality["completed_without_contract"] = int((cur.fetchone() or [0])[0] or 0)

        user_ids = sorted({int(k[2]) for k in person_scores})
        users = {}
        if user_ids:
            umarks = ",".join("?" for _ in user_ids)
            cur.execute("SELECT id,display_name,username FROM Users WHERE id IN (%s)" % umarks,
                        user_ids)
            users = {int(x["id"]): x for x in rows_to_list(cur)}

        visible_tids = set(tids)
        team_rows = []
        person_rows = []
        contract_rows = []
        total_approved = total_team_attr = total_person_attr = _decimal.Decimal(0)
        total_tasks = 0
        for cid, co in contracts_by_id.items():
            approved = _dec(co.get("approved_amount"), _decimal.Decimal(0))
            score_total = total_score.get(cid, _decimal.Decimal(0))
            task_count = len(contract_task_ids.get(cid, set()))
            if not task_count:
                quality["contracts_without_completed_task"] += 1
            total_approved += approved; total_tasks += task_count

            contract_team_amount = _decimal.Decimal(0)
            contract_person_amount = _decimal.Decimal(0)
            for (ccid, team_id), score in team_scores.items():
                if ccid != cid or (visible_tids and team_id not in visible_tids):
                    continue
                amount = allocate_approved_amount(approved, score, score_total)
                contract_team_amount += amount
                team_name = next((t.get("team_name") for t in tasks
                                  if int(t["contract_id"]) == cid and _int(t.get("team_id")) == team_id), None)
                team_rows.append({
                    **co, "team_id": team_id, "team_name": team_name or "تیم نامشخص",
                    "completed_task_count": len(team_task_ids.get((cid, team_id), set())),
                    "weighted_score": score, "share_percent": _pct(score, score_total),
                    "financial_share_amount": amount,
                })
            for (ccid, team_id, uid), score in person_scores.items():
                if ccid != cid or (visible_tids and team_id not in visible_tids):
                    continue
                if user_id_filter is not None and uid != user_id_filter:
                    continue
                amount = allocate_approved_amount(approved, score, score_total)
                contract_person_amount += amount
                u = users.get(uid, {})
                team_name = next((t.get("team_name") for t in tasks
                                  if int(t["contract_id"]) == cid and _int(t.get("team_id")) == team_id), None)
                person_rows.append({
                    **co, "team_id": team_id, "team_name": team_name or "تیم نامشخص",
                    "user_id": uid,
                    "display_name": u.get("display_name") or u.get("username") or str(uid),
                    "completed_task_count": len(person_task_ids.get((cid, team_id, uid), set())),
                    "weighted_score": score, "share_percent": _pct(score, score_total),
                    "financial_share_amount": amount,
                })
            total_team_attr += contract_team_amount
            total_person_attr += contract_person_amount
            contract_rows.append({
                **co, "completed_task_count": task_count, "weighted_score": score_total,
                "team_attributed_amount": contract_team_amount,
                "person_attributed_amount": contract_person_amount,
                "unallocated_team_amount": max(_decimal.Decimal(0), approved-contract_team_amount),
                "unallocated_person_amount": max(_decimal.Decimal(0), approved-contract_person_amount),
            })

        def aggregate(rows, id_key, name_key):
            result = {}
            for row in rows:
                key = row.get(id_key)
                out = result.setdefault(key, {
                    id_key: key, name_key: row.get(name_key) or "تعیین نشده",
                    "approved_amount": _decimal.Decimal(0),
                    "team_attributed_amount": _decimal.Decimal(0),
                    "person_attributed_amount": _decimal.Decimal(0),
                    "contract_count": 0, "completed_task_count": 0,
                })
                out["approved_amount"] += _dec(row.get("approved_amount"), _decimal.Decimal(0))
                out["team_attributed_amount"] += _dec(row.get("team_attributed_amount"), _decimal.Decimal(0))
                out["person_attributed_amount"] += _dec(row.get("person_attributed_amount"), _decimal.Decimal(0))
                out["contract_count"] += 1
                out["completed_task_count"] += int(row.get("completed_task_count") or 0)
            values = list(result.values())
            for row in values:
                row["revenue_share_percent"] = _pct(row["approved_amount"], total_approved)
            values.sort(key=lambda x: _dec(x["approved_amount"], _decimal.Decimal(0)), reverse=True)
            return values

        team_rows.sort(key=lambda x: _dec(x.get("financial_share_amount"), _decimal.Decimal(0)), reverse=True)
        person_rows.sort(key=lambda x: _dec(x.get("financial_share_amount"), _decimal.Decimal(0)), reverse=True)
        contract_rows.sort(key=lambda x: _dec(x.get("approved_amount"), _decimal.Decimal(0)), reverse=True)
        return {
            "totals": {
                "approved_amount": total_approved,
                "team_attributed_amount": total_team_attr,
                "person_attributed_amount": total_person_attr,
                "unallocated_team_amount": max(_decimal.Decimal(0), total_approved-total_team_attr),
                "unallocated_person_amount": max(_decimal.Decimal(0), total_approved-total_person_attr),
                "completed_task_count": total_tasks,
            },
            "contracts": contract_rows,
            "teams": team_rows,
            "people": person_rows,
            "cities": aggregate(contract_rows, "city_id", "city_name"),
            "contract_types": aggregate(contract_rows, "contract_type_id", "contract_type_name"),
            "quality": quality,
            "filters": {"from_date": start, "to_date": end, "city_id": city_id,
                        "contract_type_id": contract_type_id, "contract_id": contract_id,
                        "user_id": user_id_filter, "team_ids": tids},
            "access_scope": {"limited": scope_limited, "team_ids": tids},
        }

    @app.route("/api/v8/contract_task_weights", methods=["POST"])
    @require_auth
    def api_contract_task_weights():
        if not user_has_permission(g.user, "contracts.view"):
            return err("اجازه مشاهده تنظیمات وزن را ندارید", 403)
        data = request.get_json() or {}; action = data.get("action") or "read"
        try:
            with db_lock:
                c = get_conn(); cur = c.cursor()
                if action == "save":
                    if not user_has_permission(g.user, "contract_weights.manage"):
                        return err("دسترسی تغییر وزن‌ها را ندارید", 403)
                    rows = data.get("rows") or []
                    for row in rows:
                        ctid = _int(row.get("contract_type_id")); catid = _int(row.get("task_category_id"))
                        weight = _dec(row.get("weight"))
                        if ctid not in (0, 1, 2, 3) or catid is None or weight is None or weight <= 0 or weight > 1000:
                            raise ValueError("وزن یا نوع انتخاب‌شده نامعتبر است")
                        cur.execute("""IF EXISTS(SELECT 1 FROM ContractTaskCategoryWeights
                              WHERE contract_type_id=? AND task_category_id=?)
                            UPDATE ContractTaskCategoryWeights SET weight=?,is_active=1,updated_by=?,updated_at=GETDATE()
                              WHERE contract_type_id=? AND task_category_id=?
                            ELSE INSERT INTO ContractTaskCategoryWeights(contract_type_id,task_category_id,weight,is_active,updated_by,updated_at)
                              VALUES(?,?,?,1,?,GETDATE())""",
                            ctid, catid, weight, g.user["id"], ctid, catid,
                            ctid, catid, weight, g.user["id"])
                    c.commit()
                cur.execute("""SELECT ct.id AS contract_type_id,ct.name AS contract_type_name,
                        tc.id AS task_category_id,tc.name AS task_category_name,
                        ISNULL(w.weight,1) AS weight
                    FROM ContractTypes ct CROSS JOIN TaskCategories tc
                    LEFT JOIN ContractTaskCategoryWeights w
                      ON w.contract_type_id=ct.id AND w.task_category_id=tc.id AND w.is_active=1
                    WHERE ct.id IN (0,1,2,3)
                    ORDER BY ct.id,tc.name""")
                rows = rows_to_list(cur)
            return ok(rows=rows, can_edit=user_has_permission(g.user, "contract_weights.manage"))
        except Exception as exc:
            return err(public_error(exc))

    @app.route("/api/v8/financial_attribution", methods=["POST"])
    @require_auth
    def api_financial_attribution():
        try:
            with db_lock:
                c = get_conn(); cur = c.cursor()
                data = _load_attribution(cur, request.get_json() or {}, g.user)
            return ok(**data)
        except PermissionError as exc:
            return err(exc, 403)
        except ValueError as exc:
            return err(exc)
        except Exception as exc:
            return err(public_error(exc))

    def _fa_pdf(value):
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(str(value if value is not None else "")))

    def _report_value(key, value):
        if value in (None, ""):
            return "—"
        if key.endswith("_amount"):
            try:
                return f"{_decimal.Decimal(str(value)):,.0f}"
            except Exception:
                return str(value)
        if key.endswith("_percent") or key in ("share_percent", "revenue_share_percent"):
            try:
                return f"{float(value):,.2f}%"
            except Exception:
                return str(value)
        if key == "weighted_score":
            try:
                return f"{float(value):,.2f}"
            except Exception:
                return str(value)
        return str(value)

    def _pdf_export(data):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
        buf = io.BytesIO()
        font_name = fa_font.register()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=18, leftMargin=18,
                                topMargin=18, bottomMargin=18)
        title = ParagraphStyle("title", fontName=font_name, fontSize=12, leading=18, alignment=2)
        note = ParagraphStyle("note", fontName=font_name, fontSize=8, leading=13, alignment=2,
                              textColor=colors.HexColor("#374151"))
        story = [Paragraph(_fa_pdf("گزارش سهم مالی تحلیلی قراردادها"), title), Spacer(1, 5)]
        filters = data.get("filters") or {}
        range_text = "بازه: %s تا %s" % (
            filters.get("from_date") or "ابتدای اطلاعات",
            filters.get("to_date") or "انتهای اطلاعات",
        )
        story.append(Paragraph(_fa_pdf(range_text), note))
        story.append(Paragraph(_fa_pdf(
            "مبالغ این گزارش سهم تحلیلی از درآمد تاییدشده هستند و حقوق، پاداش یا بدهی شرکت محسوب نمی‌شوند. "
            "سهم تیم از امتیاز تسک تکمیل‌شده و سهم فرد ابتدا از زمان واقعی ثبت‌شده محاسبه می‌شود."
        ), note))
        if (data.get("access_scope") or {}).get("limited"):
            story.append(Paragraph(_fa_pdf(
                "گزارش در محدوده تیم‌های مجاز کاربر تهیه شده است؛ بخش تخصیص‌نیافته ممکن است شامل سهم تیم‌های خارج از این محدوده باشد."
            ), note))
        quality = data.get("quality") or {}
        warnings = []
        labels = {
            "completed_without_contract": "تسک تکمیل‌شده بدون قرارداد",
            "completed_without_team": "تسک تکمیل‌شده بدون تیم",
            "multi_assignee_without_time": "تسک چندمسئوله بدون ثبت زمان",
            "mapping_missing": "تسک با وزن خنثی به دلیل نبود نگاشت",
            "contracts_without_completed_task": "قرارداد دارای درآمد بدون تسک تکمیل‌شده",
        }
        for key, label in labels.items():
            if int(quality.get(key) or 0):
                warnings.append(f"{label}: {int(quality[key])}")
        if warnings:
            story.append(Paragraph(_fa_pdf("کیفیت داده — " + " | ".join(warnings)), note))
        story.append(Spacer(1, 8))

        sections = [
            ("خلاصه شهرها", data.get("cities") or [], [
                ("city_name", "شهر"), ("contract_count", "قرارداد"),
                ("completed_task_count", "تسک تکمیل‌شده"), ("approved_amount", "درآمد تاییدشده"),
                ("team_attributed_amount", "سهم تیم‌ها"), ("person_attributed_amount", "سهم افراد"),
                ("revenue_share_percent", "سهم از درآمد")]),
            ("خلاصه انواع قرارداد", data.get("contract_types") or [], [
                ("contract_type_name", "نوع قرارداد"), ("contract_count", "قرارداد"),
                ("completed_task_count", "تسک تکمیل‌شده"), ("approved_amount", "درآمد تاییدشده"),
                ("team_attributed_amount", "سهم تیم‌ها"), ("person_attributed_amount", "سهم افراد"),
                ("revenue_share_percent", "سهم از درآمد")]),
            ("خلاصه قراردادها", data.get("contracts") or [], [
                ("city_name", "شهر"), ("contract_type_name", "نوع قرارداد"),
                ("contract_title", "قرارداد"), ("completed_task_count", "تسک"),
                ("approved_amount", "درآمد تاییدشده"), ("team_attributed_amount", "سهم تیم‌ها"),
                ("person_attributed_amount", "سهم افراد"), ("unallocated_person_amount", "تخصیص‌نیافته فردی")]),
            ("سهم تیمی", data.get("teams") or [], [
                ("city_name", "شهر"), ("contract_type_name", "نوع قرارداد"),
                ("contract_title", "قرارداد"), ("team_name", "تیم"),
                ("completed_task_count", "تسک"), ("weighted_score", "امتیاز وزنی"),
                ("share_percent", "درصد سهم"), ("financial_share_amount", "سهم مالی")]),
            ("سهم فردی", data.get("people") or [], [
                ("city_name", "شهر"), ("contract_type_name", "نوع قرارداد"),
                ("contract_title", "قرارداد"), ("team_name", "تیم"),
                ("display_name", "همکار"), ("completed_task_count", "تسک"),
                ("share_percent", "درصد سهم"), ("financial_share_amount", "سهم تحلیلی مالی")]),
        ]
        for index, (name, rows, cols) in enumerate(sections):
            if index:
                story.append(PageBreak())
            story.append(Paragraph(_fa_pdf(name), title)); story.append(Spacer(1, 6))
            table_data = [[_fa_pdf(label) for _, label in cols]]
            for row in rows:
                table_data.append([_fa_pdf(_report_value(key, row.get(key))) for key, _ in cols])
            if len(table_data) == 1:
                table_data.append([_fa_pdf("داده‌ای وجود ندارد")] + [""] * (len(cols)-1))
            table = Table(table_data, repeatRows=1)
            table.setStyle(TableStyle([
                ("FONT", (0, 0), (-1, -1), font_name, 7),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#274690")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("GRID", (0, 0), (-1, -1), .25, colors.grey),
                ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(table)
        doc.build(story); buf.seek(0)
        return buf

    def _xlsx_export(data):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook(); wb.remove(wb.active)
        sheets = [
            ("شهرها", data.get("cities") or [], [
                ("city_name", "شهر"), ("contract_count", "تعداد قرارداد"),
                ("completed_task_count", "تسک تکمیل‌شده"), ("approved_amount", "درآمد تاییدشده"),
                ("team_attributed_amount", "سهم تیم‌ها"), ("person_attributed_amount", "سهم افراد"),
                ("revenue_share_percent", "درصد سهم درآمد")]),
            ("انواع قرارداد", data.get("contract_types") or [], [
                ("contract_type_name", "نوع قرارداد"), ("contract_count", "تعداد قرارداد"),
                ("completed_task_count", "تسک تکمیل‌شده"), ("approved_amount", "درآمد تاییدشده"),
                ("team_attributed_amount", "سهم تیم‌ها"), ("person_attributed_amount", "سهم افراد"),
                ("revenue_share_percent", "درصد سهم درآمد")]),
            ("قراردادها", data.get("contracts") or [], [
                ("city_name", "شهر"), ("project_name", "پروژه"),
                ("contract_type_name", "نوع قرارداد"), ("contract_number", "شماره قرارداد"),
                ("contract_title", "عنوان قرارداد"), ("completed_task_count", "تسک تکمیل‌شده"),
                ("weighted_score", "امتیاز وزنی"), ("approved_amount", "درآمد تاییدشده"),
                ("team_attributed_amount", "سهم تیم‌ها"), ("person_attributed_amount", "سهم افراد"),
                ("unallocated_team_amount", "تخصیص‌نیافته تیمی"),
                ("unallocated_person_amount", "تخصیص‌نیافته فردی")]),
            ("سهم تیمی", data.get("teams") or [], [
                ("city_name", "شهر"), ("contract_type_name", "نوع قرارداد"),
                ("contract_title", "قرارداد"), ("team_name", "تیم"),
                ("completed_task_count", "تسک"), ("weighted_score", "امتیاز وزنی"),
                ("share_percent", "درصد سهم"), ("financial_share_amount", "سهم مالی")]),
            ("سهم فردی", data.get("people") or [], [
                ("city_name", "شهر"), ("contract_type_name", "نوع قرارداد"),
                ("contract_title", "قرارداد"), ("team_name", "تیم"),
                ("display_name", "همکار"), ("completed_task_count", "تسک"),
                ("weighted_score", "امتیاز وزنی"), ("share_percent", "درصد سهم"),
                ("financial_share_amount", "سهم تحلیلی مالی")]),
        ]
        for title, rows, columns in sheets:
            ws = wb.create_sheet(title); ws.sheet_view.rightToLeft = True
            ws.append([label for _, label in columns])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="274690")
                cell.alignment = Alignment(horizontal="center")
            for row in rows:
                ws.append([_jsonable(row.get(key)) for key, _ in columns])
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                width = max((len(str(c.value or "")) for c in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(45, max(12, width + 2))
        meta = wb.create_sheet("راهنما", 0); meta.sheet_view.rightToLeft = True
        meta.append(["عنوان", "مقدار"])
        meta.append(["ماهیت مبلغ", "سهم تحلیلی؛ نه حقوق، پاداش یا بدهی شرکت"])
        meta.append(["مبنای تیم", "امتیاز تسک‌های تکمیل‌شده متصل به قرارداد"])
        meta.append(["مبنای فرد", "زمان واقعی؛ و فقط در تسک تک‌مسئوله بدون زمان، مسئول اصلی"])
        meta.append(["محدوده دسترسی", "تیمی" if (data.get("access_scope") or {}).get("limited") else "سراسری"])
        for cell in meta[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="274690")
        meta.column_dimensions["A"].width = 24; meta.column_dimensions["B"].width = 80
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return buf

    @app.route("/api/v8/financial_attribution_export", methods=["POST"])
    @require_auth
    def api_financial_attribution_export():
        try:
            payload = request.get_json() or {}; fmt = (payload.get("format") or "xlsx").lower()
            with db_lock:
                c = get_conn(); cur = c.cursor(); data = _load_attribution(cur, payload, g.user)
            if fmt == "pdf":
                buf = _pdf_export(data)
                return send_file(buf, as_attachment=True,
                    download_name="TaskHub_financial_attribution.pdf", mimetype="application/pdf")
            buf = _xlsx_export(data)
            return send_file(buf, as_attachment=True,
                download_name="TaskHub_financial_attribution.xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        except PermissionError as exc:
            return err(exc, 403)
        except Exception as exc:
            return err(public_error(exc))
